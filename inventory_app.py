import os
import glob
import sys
import unicodedata
import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl
from openpyxl.utils import get_column_letter
import socket
import json
import threading
import time
import urllib.request
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

def get_base_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def make_hoverable(button, hover_bg, hover_fg=None, normal_bg=None, normal_fg=None):
    # Determine defaults dynamically if not provided
    try:
        if not normal_bg:
            normal_bg = button.cget("bg")
        if not normal_fg:
            normal_fg = button.cget("fg")
    except Exception:
        pass
        
    def on_enter(e):
        button.config(bg=hover_bg)
        if hover_fg:
            button.config(fg=hover_fg)
            
    def on_leave(e):
        if normal_bg:
            button.config(bg=normal_bg)
        if normal_fg:
            button.config(fg=normal_fg)
            
    button.bind("<Enter>", on_enter)
    button.bind("<Leave>", on_leave)

# Define UI Colors (Modern Dark Theme)
BG_COLOR = "#0F172A"       # Slate 900 (Deep background)
CARD_BG = "#1E293B"        # Slate 800 (Card background)
TEXT_COLOR = "#F8FAFC"     # Slate 50 (Bright text)
TEXT_MUTED = "#94A3B8"     # Slate 400 (Muted text)
ACCENT_COLOR = "#6366F1"   # Indigo 500 (Primary accent button)
ACCENT_HOVER = "#4F46E5"   # Indigo 600
SUCCESS_COLOR = "#10B981"  # Emerald 500 (Positive/Surplus)
DANGER_COLOR = "#EF4444"   # Red 500 (Negative/Shortage)
WARNING_COLOR = "#FACC15"  # Yellow 400 (Exact count / yellow)
ENTRY_BG = "#334155"       # Slate 700 (Input background)
BORDER_COLOR = "#475569"   # Slate 600 (Card border)

# Global Variables
filepath = ""
wb = None
sheet = None
products = {}
operator_counts = {}
count_order = 0
operator_name = "Mərkəz"

# Network Variables
is_network_mode = False
server_ip = ""
server_port = 5000
server_thread = None
httpd = None
db_lock = threading.Lock()
active_operators = {}
is_server_active = False

# Allowed Operators from config file
allowed_operators = {}

def load_operators():
    global allowed_operators
    current_dir = get_base_dir()
    file_path = os.path.join(current_dir, "operatorlar.txt")
    if not os.path.exists(file_path):
        with open(file_path, "w", encoding="utf-8") as f:
            f.write("Orxan:1234\nElmir:2222\nPerviz:3333\nNamiq:4444\nVuqar:5555\n")
            
    allowed_operators.clear()
    with open(file_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and ":" in line:
                parts = line.split(":", 1)
                allowed_operators[parts[0].strip()] = parts[1].strip()

load_operators()

def save_output_excel(dest_path):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    try:
        temp_wb = openpyxl.load_workbook(dest_path)
        sheet = temp_wb.active
        
        # 1. Enable AutoFilter for all columns
        last_col_letter = get_column_letter(sheet.max_column)
        sheet.auto_filter.ref = f"A1:{last_col_letter}{sheet.max_row}"
        
        # 2. Design system tokens
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Slate Gray
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        border_side = Side(border_style="thin", color="CBD5E1") # Soft border
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        
        # Style headers
        sheet.row_dimensions[1].height = 28
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = cell_border
            
        # Style data cells
        for row_idx in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_idx].height = 20
            for col_idx in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                cell.border = cell_border
                cell.font = Font(name="Segoe UI", size=10)
                
                is_text_col = False
                if 'adi_col' in globals() and col_idx == adi_col:
                    is_text_col = True
                elif 'brend_col' in globals() and col_idx == brend_col:
                    is_text_col = True
                
                if is_text_col:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
        # 3. Add YEKUN CƏMİ row at the bottom
        total_row = sheet.max_row + 1
        label_col = brend_col or 2
        sheet.cell(row=total_row, column=label_col, value="YEKUN CƏMİ:")
        
        if yeni_sayim_col:
            yeni_let = get_column_letter(yeni_sayim_col)
            sheet.cell(row=total_row, column=yeni_sayim_col, value=f"=SUM({yeni_let}2:{yeni_let}{total_row-1})")
            
        if say_ferqi_col:
            ferq_let = get_column_letter(say_ferqi_col)
            sheet.cell(row=total_row, column=say_ferqi_col, value=f"=SUM({ferq_let}2:{ferq_let}{total_row-1})")
            
        if qiymet_ferqi_col:
            qiymet_ferqi_letter = get_column_letter(qiymet_ferqi_col)
            sheet.cell(row=total_row, column=qiymet_ferqi_col, value=f"=SUM({qiymet_ferqi_letter}2:{qiymet_ferqi_letter}{total_row-1})")
            
        # Style the Total row
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Soft Gray highlight
        sheet.row_dimensions[total_row].height = 24
        
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=total_row, column=col_idx)
            cell.fill = total_fill
            cell.border = cell_border
            if col_idx == label_col:
                cell.font = Font(name="Segoe UI", size=10, bold=True)
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in [yeni_sayim_col, say_ferqi_col, qiymet_ferqi_col]:
                cell.font = Font(name="Segoe UI", size=10, bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
        # Adjust column widths automatically
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if val_str.startswith('='):
                    val_str = "100.00 AZN"
                if len(val_str) > max_len:
                    max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        sheet.views.sheetView[0].showGridLines = True
        temp_wb.save(dest_path)
        temp_wb.close()
    except Exception as e:
        print("Excel yadda saxlama və bəzəmə xətası:", e)

autosave_pending = False
def autosave_worker():
    global autosave_pending
    import shutil
    while True:
        time.sleep(5)  # Check every 5 seconds
        if autosave_pending:
            try:
                # Save main file under lock (very fast, less than 100ms)
                with db_lock:
                    wb.save(filepath)
                    autosave_pending = False
                
                # Copy and style outside the lock!
                dir_name = os.path.dirname(filepath)
                new_file_path = os.path.join(dir_name, "yeni sayim neticesi.xlsx")
                shutil.copy2(filepath, new_file_path)
                save_output_excel(new_file_path)
            except Exception as e:
                print("Autosave xətası:", e)

threading.Thread(target=autosave_worker, daemon=True).start()

def trigger_autosave():
    global autosave_pending
    if is_network_mode:
        return
    autosave_pending = True

def backup_manager_worker():
    import shutil
    import datetime
    while True:
        time.sleep(1200) # Run every 20 minutes (1200 seconds)
        if filepath and os.path.exists(filepath):
            try:
                # Create backup folder
                dir_name = os.path.dirname(filepath)
                backup_dir = os.path.join(dir_name, "backups")
                if not os.path.exists(backup_dir):
                    os.makedirs(backup_dir)
                
                # Make backup file path
                timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                base_name = os.path.basename(filepath)
                name, ext = os.path.splitext(base_name)
                backup_file = os.path.join(backup_dir, f"{name}_{timestamp}{ext}")
                
                # Copy currently saved file to backup path (no lock needed because it's on disk)
                shutil.copy2(filepath, backup_file)
                
                # Style backup file outside the lock!
                save_output_excel(backup_file)
                
                # Keep only last 20 backups
                existing_backups = glob.glob(os.path.join(backup_dir, f"{name}_*{ext}"))
                if len(existing_backups) > 20:
                    # Sort by modification time (oldest first)
                    existing_backups.sort(key=os.path.getmtime)
                    # Delete the oldest ones
                    for old_file in existing_backups[:-20]:
                        try:
                            os.remove(old_file)
                        except Exception:
                            pass
            except Exception as e:
                print("Ehtiyat nüsxə yaradılarkən xətə baş verdi:", e)

threading.Thread(target=backup_manager_worker, daemon=True).start()

# Column mappings
brend_col = None
adi_col = None
anbar_qaligi_col = None
barkod_col = None
yeni_sayim_col = None
qiymet_col = None
say_ferqi_col = None
qiymet_ferqi_col = None
operator_col = None

yeni_letter = ""
qaliq_letter = ""
say_ferqi_letter = ""
qiymet_letter = ""

def normalize_header(text):
    if not text:
        return ""
    text = str(text).strip()
    replacements = {
        'İ': 'i', 'I': 'i', 'İ': 'i', 'ı': 'i',
        'Ə': 'e', 'ə': 'e',
        'Ö': 'o', 'ö': 'o',
        'Ü': 'u', 'ü': 'u',
        'Ğ': 'g', 'ğ': 'g',
        'Ç': 'c', 'ç': 'c',
        'Ş': 's', 'ş': 's'
    }
def is_summary_row(barkod_val, kod_val, brend_val, adi_val):
    combined = (str(barkod_val) + " " + str(kod_val) + " " + str(brend_val) + " " + str(adi_val)).lower()
    keywords = ['yekun', 'cemi', 'cem', 'total', 'итого', 'всего', 'summary', 'subtotal', 'cəmi', 'cəm']
    for kw in keywords:
        if kw in combined:
            return True
    return False


def merge_excel_files_data(files_list):
    """
    files_list: list of tuples (filename, file_bytes)
    Returns openpyxl.Workbook matching Picture 3 structure
    """
    merged_products = {}
    
    barkod_synonyms = ['barkod', 'barkkod', 'barcode', 'strixkod', 'strix-kod', 'strix_kod', 'ean', 'shtrihkod', 'штрихкод', 'bar_code', 'sh_kod', 'bar_kod', 'barkod_no', 'barkkod_no']
    kod_synonyms = ['kod', 'kodu', 'product_code', 'mal_kodu', 'mehsul_kodu', 'kod_mehsul', 'kod_mal', 'код', 'артикул', 'artikul', 'item_code', 'item_id']
    brend_synonyms = ['brend', 'brand', 'marka', 'istehsalci', 'производитель', 'фирма', 'firma', 'vendor', 'manufacturer', 'malin_markasi', 'taminatci', 'supplier']
    adi_synonyms = ['mehsulun adi', 'mehsulun_adi', 'malin adi', 'malin_adi', 'adi', 'name', 'mehsul', 'description', 'nomenklatura', 'наименование', 'название', 'товар', 'urun_adi', 'tam_adi', 'mehsul_adi', 'item_name', 'aciglama']
    qaliq_synonyms = ['sistem qaligi', 'sistem_qaligi', 'anbar qaligi', 'anbar_qaligi', 'qaliq', 'stock', 'qaliq_miqdari', 'остаток', 'stok_miktari', 'stok', 'balance', 'qty', 'quantity', 'son_qaliq', 'tek_qaliq', 'mevcut']
    sayim_synonyms = ['yeni sayim', 'yeni_sayim', 'sayim', 'sayilan', 'sayim miqdari', 'sayim_miqdari', 'sayilan_miqdar', 'sayim_miqdar', 'faktiki sayim', 'faktiki_sayim', 'miqdar', 'sayi', 'say', 'количество', 'fakt', 'yeni_sayim_miqdari', 'real say']
    qiymet_synonyms = ['qiymet', 'qiymeti', 'mehsulun qiymeti', 'mehsulun_qiymeti', 'price', 'satis_qiymeti', 'satis qiymeti', 'satis_qiymet', 'satis qiymet', 'satiş qiyməti', 'maya_qiymeti', 'maya qiymeti', 'цена', 'стоимость', 'fiyat', 'cost', 'unit_price', 'qiymat', 'mebleg', 'perakende']
    operator_synonyms = ['operator', 'sayimci', 'user', 'istifadəçi', 'operatorlar', 'sayımçı']

    for filename, file_bytes in files_list:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb.active
        except Exception:
            continue
            
        headers = {}
        for col in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col).value
            if val is not None:
                headers[normalize_header(val)] = col

        def get_col_idx(synonyms):
            for syn in synonyms:
                norm_syn = normalize_header(syn)
                for h_key, col_idx in headers.items():
                    if norm_syn == h_key or norm_syn in h_key or h_key in norm_syn:
                        return col_idx
            return None

        barkod_col = get_col_idx(barkod_synonyms)
        kod_col = get_col_idx(kod_synonyms)
        brend_col = get_col_idx(brend_synonyms)
        adi_col = get_col_idx(adi_synonyms)
        qaliq_col = get_col_idx(qaliq_synonyms)
        sayim_col = get_col_idx(sayim_synonyms)
        qiymet_col = get_col_idx(qiymet_synonyms)
        operator_col = get_col_idx(operator_synonyms)

        if not sayim_col:
            sayim_col = qaliq_col

        for row in range(2, sheet.max_row + 1):
            barkod_val = str(sheet.cell(row=row, column=barkod_col).value or "").strip() if barkod_col else ""
            if barkod_val.endswith(".0"): barkod_val = barkod_val[:-2]

            kod_val = str(sheet.cell(row=row, column=kod_col).value or "").strip() if kod_col else ""
            if kod_val.endswith(".0"): kod_val = kod_val[:-2]

            brend_val = str(sheet.cell(row=row, column=brend_col).value or "").strip() if brend_col else ""
            adi_val = str(sheet.cell(row=row, column=adi_col).value or "").strip() if adi_col else ""
            operator_val = str(sheet.cell(row=row, column=operator_col).value or "").strip() if operator_col else ""

            if is_summary_row(barkod_val, kod_val, brend_val, adi_val):
                continue

            raw_count = sheet.cell(row=row, column=sayim_col).value if sayim_col else 0
            try:
                count_num = float(raw_count) if raw_count is not None else 0.0
            except (ValueError, TypeError):
                count_num = 0.0

            raw_qaliq = sheet.cell(row=row, column=qaliq_col).value if qaliq_col else 0
            try:
                qaliq_num = float(raw_qaliq) if raw_qaliq is not None else 0.0
            except (ValueError, TypeError):
                qaliq_num = 0.0

            raw_qiymet = sheet.cell(row=row, column=qiymet_col).value if qiymet_col else 0
            try:
                qiymet_num = float(raw_qiymet) if raw_qiymet is not None else 0.0
            except (ValueError, TypeError):
                qiymet_num = 0.0

            product_key = barkod_val or kod_val or normalize_header(adi_val)
            if not product_key:
                continue

            clean_filename = os.path.basename(filename)

            if product_key not in merged_products:
                merged_products[product_key] = {
                    'barkod': barkod_val,
                    'kod': kod_val,
                    'brend': brend_val,
                    'adi': adi_val,
                    'anbar_qaligi': qaliq_num,
                    'yeni_sayim': count_num,
                    'qiymet': qiymet_num,
                    'operators': [operator_val] if operator_val else [],
                    'sources': [clean_filename]
                }
            else:
                existing = merged_products[product_key]
                existing['yeni_sayim'] += count_num
                if not existing['anbar_qaligi'] and qaliq_num:
                    existing['anbar_qaligi'] = qaliq_num
                if not existing['qiymet'] and qiymet_num:
                    existing['qiymet'] = qiymet_num
                if not existing['brend'] and brend_val:
                    existing['brend'] = brend_val
                if not existing['adi'] and adi_val:
                    existing['adi'] = adi_val
                if operator_val and operator_val not in existing['operators']:
                    existing['operators'].append(operator_val)
                if clean_filename not in existing['sources']:
                    existing['sources'].append(clean_filename)

    out_wb = openpyxl.Workbook()
    out_sheet = out_wb.active
    out_sheet.title = "Yekun Sayım"

    out_sheet.views.sheetView[0].showGridLines = True

    # Header matching Picture 3
    headers_list = [
        "No", "Kod", "Barkod", "Brend", "Məhsulun Adı",
        "Qiymət", "Sistem Qalıq", "Yeni Sayım", "Say Fərqi",
        "Qiymət Fərqi", "Operator", "Keçdiyi Fayllar"
    ]
    out_sheet.append(headers_list)
    out_sheet.row_dimensions[1].height = 28

    sorted_products = sorted(
        merged_products.values(),
        key=lambda x: (x['brend'].lower(), x['adi'].lower(), x['barkod'])
    )

    total_qaliq = 0
    total_sayim = 0
    total_ferq = 0
    total_mebleg_ferqi = 0.0

    for idx, prod in enumerate(sorted_products, 1):
        r = idx + 1
        qaliq = prod['anbar_qaligi']
        sayim = prod['yeni_sayim']
        qiymet = prod['qiymet']
        operators_str = ", ".join(prod['operators']) if prod['operators'] else ""
        sources_str = ", ".join(prod['sources'])

        q_disp = int(qaliq) if qaliq.is_integer() else round(qaliq, 2)
        s_disp = int(sayim) if sayim.is_integer() else round(sayim, 2)

        # Dynamic Excel Formulas as requested by user:
        # Say Fərqi = Yeni Sayım - Sistem Qalıq (=H{r}-G{r})
        # Qiymət Fərqi = Say Fərqi * Qiymət (=I{r}*F{r})
        say_ferqi_formula = f"=H{r}-G{r}"
        qiymet_ferqi_formula = f"=I{r}*F{r}"

        out_sheet.append([
            idx,
            prod['kod'],
            prod['barkod'],
            prod['brend'],
            prod['adi'],
            round(qiymet, 2),
            q_disp,
            s_disp,
            say_ferqi_formula,
            qiymet_ferqi_formula,
            operators_str,
            sources_str
        ])

    last_row = len(sorted_products) + 2
    data_last_row = last_row - 1

    total_row = [
        "YEKUN CƏM", "", "", "", "",
        "",
        f"=SUM(G2:G{data_last_row})",
        f"=SUM(H2:H{data_last_row})",
        f"=SUM(I2:I{data_last_row})",
        f"=SUM(J2:J{data_last_row})",
        "",
        f"Cəmi {len(sorted_products)} çeşit məhsul"
    ]
    out_sheet.append(total_row)

    # Enable AutoFilter for header row
    last_col_letter = get_column_letter(len(headers_list))
    out_sheet.auto_filter.ref = f"A1:{last_col_letter}{out_sheet.max_row}"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_wrap_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col in range(1, len(headers_list) + 1):
        cell = out_sheet.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    
    green_font = Font(name="Segoe UI", size=10, bold=True, color="15803D")
    red_font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
    data_font = Font(name="Segoe UI", size=10, bold=False, color="1E293B")
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for r in range(2, last_row):
        out_sheet.row_dimensions[r].height = 22
        is_even = (r % 2 == 0)
        row_bg = zebra_fill if is_even else white_fill

        for c in range(1, len(headers_list) + 1):
            cell = out_sheet.cell(row=r, column=c)
            cell.font = data_font
            cell.border = thin_border
            cell.fill = row_bg
            
            if c == 1:
                cell.alignment = center_align
                cell.number_format = '0'
            elif c in [2, 3]:
                cell.alignment = center_align
                cell.number_format = '@'
            elif c in [4, 5, 11]:
                cell.alignment = left_wrap_align
            elif c in [6, 10]:
                cell.alignment = right_align
                cell.number_format = '#,##0.00'
            elif c in [7, 8, 9]:
                cell.alignment = right_align
                cell.number_format = '#,##0' if isinstance(cell.value, int) else '#,##0.00'
            elif c == 12:
                cell.alignment = left_wrap_align

        idx_prod = r - 2
        prod_data = sorted_products[idx_prod]
        ferq_calc = prod_data['yeni_sayim'] - prod_data['anbar_qaligi']
        cell_f = out_sheet.cell(row=r, column=9)
        cell_m = out_sheet.cell(row=r, column=10)
        
        if ferq_calc > 0:
            cell_f.fill = green_fill
            cell_f.font = green_font
            cell_m.fill = green_fill
            cell_m.font = green_font
        elif ferq_calc < 0:
            cell_f.fill = red_fill
            cell_f.font = red_font
            cell_m.fill = red_fill
            cell_m.font = red_font

    out_sheet.row_dimensions[last_row].height = 26
    summary_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    summary_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    
    for c in range(1, len(headers_list) + 1):
        cell = out_sheet.cell(row=last_row, column=c)
        cell.fill = summary_fill
        cell.font = summary_font
        cell.border = thin_border
        if c in [7, 8, 9]:
            cell.alignment = right_align
            cell.number_format = '#,##0'
        elif c in [6, 10]:
            cell.alignment = right_align
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = center_align if c == 1 else left_align

    min_widths = {
        1: 10,   # No
        2: 16,   # Kod
        3: 20,   # Barkod
        4: 22,   # Brend
        5: 45,   # Məhsulun Adı
        6: 14,   # Qiymət
        7: 18,   # Sistem Qalıq
        8: 18,   # Yeni Sayım
        9: 18,   # Say Fərqi
        10: 20,  # Qiymət Fərqi
        11: 20,  # Operator
        12: 40   # Keçdiyi Fayllar
    }

    for col_idx in range(1, len(headers_list) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in out_sheet[col_letter]:
            val_str = str(cell.value or '')
            if cell.row == 1:
                max_len = max(max_len, len(val_str) + 6)
            else:
                max_len = max(max_len, len(val_str) + 3)
        
        calculated_width = max(max_len, min_widths.get(col_idx, 15))
        out_sheet.column_dimensions[col_letter].width = calculated_width

    return out_wb

def scan_excel_files():
    current_dir = get_base_dir()
    xlsx_files = glob.glob(os.path.join(current_dir, "*.xlsx"))
    # Filter out temporary Excel files and backups
    xlsx_files = [f for f in xlsx_files if not os.path.basename(f).startswith("~$") and not f.endswith(".bak")]
    return xlsx_files

def select_file_gui(files):
    # Selection window if multiple files exist
    select_win = tk.Tk()
    select_win.title("Excel Faylı Seçin - Anbar Sayımı")
    select_win.geometry("500x400")
    select_win.configure(bg=BG_COLOR)
    select_win.resizable(False, False)
    
    # Center Window
    select_win.update_idletasks()
    x = (select_win.winfo_screenwidth() - select_win.winfo_width()) // 2
    y = (select_win.winfo_screenheight() - select_win.winfo_height()) // 2
    select_win.geometry(f"+{x}+{y}")
    
    title_lbl = tk.Label(
        select_win, 
        text="ANBAR SAYIMI BOTU", 
        font=('Segoe UI', 16, 'bold'), 
        bg=BG_COLOR, 
        fg=ACCENT_COLOR
    )
    title_lbl.pack(pady=(20, 10))
    
    desc_lbl = tk.Label(
        select_win, 
        text="Qovluqda birdən çox Excel faylı tapıldı.\nZəhmət olmasa işləyəcəyiniz faylı seçin:", 
        font=('Segoe UI', 10), 
        bg=BG_COLOR, 
        fg=TEXT_MUTED, 
        justify="center"
    )
    desc_lbl.pack(pady=(0, 20))
    
    list_frame = tk.Frame(select_win, bg=BG_COLOR)
    list_frame.pack(padx=30, fill="both", expand=True)
    
    scrollbar = tk.Scrollbar(list_frame)
    scrollbar.pack(side="right", fill="y")
    
    listbox = tk.Listbox(
        list_frame, 
        bg=CARD_BG, 
        fg=TEXT_COLOR, 
        selectbackground=ACCENT_COLOR, 
        selectforeground=TEXT_COLOR,
        font=('Segoe UI', 10), 
        bd=1, 
        relief="flat",
        highlightbackground=BORDER_COLOR,
        highlightcolor=ACCENT_COLOR,
        yscrollcommand=scrollbar.set
    )
    listbox.pack(side="left", fill="both", expand=True)
    scrollbar.config(command=listbox.yview)
    
    for f in files:
        listbox.insert(tk.END, os.path.basename(f))
        
    # Select first by default
    if files:
        listbox.selection_set(0)
        
    selected_file = [None]
    
    def on_select():
        sel = listbox.curselection()
        if sel:
            selected_file[0] = files[sel[0]]
            select_win.destroy()
        else:
            messagebox.showwarning("Seçim Edin", "Zəhmət olmasa siyahıdan bir fayl seçin.", parent=select_win)
            
    btn = tk.Button(
        select_win, 
        text="Faylı Yüklə", 
        command=on_select, 
        bg=ACCENT_COLOR, 
        fg=TEXT_COLOR, 
        activebackground=ACCENT_HOVER,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 11, 'bold'), 
        relief="flat", 
        bd=0,
        padx=20, 
        pady=8,
        cursor="hand2"
    )
    btn.pack(pady=25)
    make_hoverable(btn, ACCENT_HOVER)
    
    select_win.mainloop()
    return selected_file[0]

def show_no_files_gui():
    no_win = tk.Tk()
    no_win.title("Excel Faylı Tapılmadı")
    no_win.geometry("500x320")
    no_win.configure(bg=BG_COLOR)
    no_win.resizable(False, False)
    
    no_win.update_idletasks()
    x = (no_win.winfo_screenwidth() - no_win.winfo_width()) // 2
    y = (no_win.winfo_screenheight() - no_win.winfo_height()) // 2
    no_win.geometry(f"+{x}+{y}")
    
    title_lbl = tk.Label(
        no_win, 
        text="Excel Faylı Tapılmadı!", 
        font=('Segoe UI', 16, 'bold'), 
        bg=BG_COLOR, 
        fg=DANGER_COLOR
    )
    title_lbl.pack(pady=(30, 20))
    
    desc_lbl = tk.Label(
        no_win, 
        text="Qovluqda heç bir Excel (.xlsx) faylı tapılmadı.\n\nZəhmət olmasa cari qovluğa son anbar qalığı\nolan Excel faylını kopyalayın və aşağıdakı düyməni sıxın.", 
        font=('Segoe UI', 11), 
        bg=BG_COLOR, 
        fg=TEXT_COLOR, 
        justify="center"
    )
    desc_lbl.pack(pady=(0, 30))
    
    action = [False]
    def on_retry():
        action[0] = True
        no_win.destroy()
        
    btn = tk.Button(
        no_win, 
        text="Yenidən Yoxla", 
        command=on_retry, 
        bg=ACCENT_COLOR, 
        fg=TEXT_COLOR, 
        activebackground=ACCENT_HOVER,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 11, 'bold'), 
        relief="flat", 
        bd=0,
        padx=25, 
        pady=8,
        cursor="hand2"
    )
    btn.pack()
    make_hoverable(btn, ACCENT_HOVER)
    
    no_win.mainloop()
    return action[0]
    
def get_active_operators():
    now = time.time()
    with db_lock:
        inactive = [op for op, t in active_operators.items() if now - t > 60]
        for op in inactive:
            del active_operators[op]
        return list(active_operators.keys())

def get_templates_dir():
    base_dir = get_base_dir()
    
    path1 = os.path.join(base_dir, "templates")
    if os.path.exists(path1):
        return path1
        
    path2 = os.path.join(os.getcwd(), "templates")
    if os.path.exists(path2):
        return path2
        
    return None

class WarehouseHTTPRequestHandler(BaseHTTPRequestHandler):
    def setup(self):
        BaseHTTPRequestHandler.setup(self)
        self.connection.settimeout(10)  # Prevent hung threads from poor Wi-Fi signals

    def log_message(self, format, *args):
        pass
        
    def send_response_data(self, content_type, body_bytes, status_code=200):
        self.send_response(status_code)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body_bytes)))
        self.send_header("Connection", "close")
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        self.end_headers()
        self.wfile.write(body_bytes)
        
    def do_GET(self):
        parsed_url = urllib.parse.urlparse(self.path)
        path = parsed_url.path
        query = urllib.parse.parse_qs(parsed_url.query)
        
        # Track operator connection
        operator = query.get("operator", [""])[0].strip()
        if operator:
            global active_operators
            with db_lock:
                active_operators[operator] = time.time()
                
        from http.cookies import SimpleCookie
        cookie_header = self.headers.get('Cookie', '')
        cookie = SimpleCookie(cookie_header)
        is_logged_in = cookie.get('admin_logged_in') and cookie['admin_logged_in'].value == 'true'
        
        if path == "/":
            self.send_response_data("text/html; charset=utf-8", MOBILE_HTML_CONTENT.encode("utf-8"))
            
        elif path in ["/admin", "/admin.html"]:
            if not is_logged_in:
                self.send_response(302)
                self.send_header("Location", "/login")
                self.end_headers()
                return
            templates_dir = get_templates_dir()
            if templates_dir:
                admin_path = os.path.join(templates_dir, "admin.html")
                if os.path.exists(admin_path):
                    with open(admin_path, "r", encoding="utf-8") as f:
                        content = f.read()
                    self.send_response_data("text/html; charset=utf-8", content.encode("utf-8"))
                    return
            self.send_response_data("text/plain; charset=utf-8", b"admin.html tapilmadi. templates qovlugunun oldugundan emin olun.", 404)
            
        elif path in ["/login", "/login.html"]:
            if is_logged_in:
                self.send_response(302)
                self.send_header("Location", "/admin")
                self.end_headers()
                return
            templates_dir = get_templates_dir()
            if templates_dir:
                login_path = os.path.join(templates_dir, "login.html")
                if os.path.exists(login_path):
                    with open(login_path, "r", encoding="utf-8") as f:
                        content = f.read()
                    self.send_response_data("text/html; charset=utf-8", content.encode("utf-8"))
                    return
            self.send_response_data("text/plain; charset=utf-8", b"login.html tapilmadi. templates qovlugunun oldugundan emin olun.", 404)
            
        elif path == "/logout":
            self.send_response(302)
            self.send_header("Location", "/login")
            self.send_header("Set-Cookie", "admin_logged_in=; Path=/; Expires=Thu, 01 Jan 1970 00:00:00 GMT; HttpOnly")
            self.end_headers()
            
        elif path == "/api/operators":
            if not is_logged_in:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Giris icazesi yoxdur!"}).encode("utf-8"), 401)
                return
            load_operators()
            ops_list = [{"username": op, "pin": pin} for op, pin in allowed_operators.items()]
            self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "success", "operators": ops_list}).encode("utf-8"))
            
        elif path in ["/download_excel", "/download_discrepancies_excel", "/download_counted_excel"]:
            if not is_logged_in:
                self.send_response(302)
                self.send_header("Location", "/login")
                self.end_headers()
                return
            with db_lock:
                if filepath and os.path.exists(filepath):
                    with open(filepath, "rb") as f:
                        file_bytes = f.read()
                    self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    self.send_header("Content-Disposition", f"attachment; filename={os.path.basename(filepath)}")
                    self.send_header("Content-Length", str(len(file_bytes)))
                    self.end_headers()
                    self.wfile.write(file_bytes)
                else:
                    self.send_response_data("text/plain; charset=utf-8", b"Excel fayli tapilmadi", 404)
            
        elif path == "/manifest.json":
            manifest_data = {
                "name": "Mobil Anbar Sayımı",
                "short_name": "Sayım",
                "start_url": "/",
                "display": "standalone",
                "background_color": "#0F172A",
                "theme_color": "#6366F1",
                "orientation": "portrait",
                "icons": [
                    {
                        "src": "data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 100 100'><rect width='100' height='100' rx='20' fill='%236366F1'/><path d='M30 30h40v40H30z' fill='none' stroke='white' stroke-width='6'/><circle cx='50' cy='50' r='10' fill='white'/></svg>",
                        "sizes": "192x192 512x512",
                        "type": "image/svg+xml"
                    }
                ]
            }
            self.send_response_data("application/json; charset=utf-8", json.dumps(manifest_data).encode("utf-8"))
            
        elif path == "/sw.js":
            sw_code = """
            const CACHE_NAME = 'anbar-sayimi-v8';
            const ASSETS = [
                '/',
                '/html5-qrcode.min.js',
                'https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Outfit:wght@600;700;800&display=swap'
            ];
            
            self.addEventListener('install', (e) => {
                e.waitUntil(
                    caches.open(CACHE_NAME).then((cache) => {
                        return cache.addAll(ASSETS);
                    }).then(() => self.skipWaiting())
                );
            });
            
            self.addEventListener('activate', (e) => {
                e.waitUntil(
                    caches.keys().then((keys) => {
                        return Promise.all(
                            keys.map((key) => {
                                if (key !== CACHE_NAME) {
                                    return caches.delete(key);
                                }
                            })
                        );
                    }).then(() => self.clients.claim())
                );
            });
            
            self.addEventListener('fetch', (e) => {
                if (e.request.method !== 'GET' || e.request.url.includes('/add_count') || e.request.url.includes('/get_products') || e.request.url.includes('/get_product_details') || e.request.url.includes('/api/active_operators')) {
                    return;
                }
                e.respondWith(
                    caches.match(e.request).then((cachedResponse) => {
                        if (cachedResponse) {
                            return cachedResponse;
                        }
                        return fetch(e.request).then((response) => {
                            if (response && response.status === 200 && response.type === 'basic') {
                                const responseToCache = response.clone();
                                caches.open(CACHE_NAME).then((cache) => {
                                    cache.put(e.request, responseToCache);
                                });
                            }
                            return response;
                        }).catch(() => {
                            if (e.request.mode === 'navigate') {
                                return caches.match('/');
                            }
                        });
                    })
                );
            });
            """
            self.send_response_data("application/javascript", sw_code.strip().encode("utf-8"))
            
        elif path == "/html5-qrcode.min.js":
            script_dir = get_base_dir()
            js_filepath = os.path.join(script_dir, "html5-qrcode.min.js")
            if os.path.exists(js_filepath):
                with open(js_filepath, "rb") as f:
                    js_data = f.read()
                self.send_response_data("application/javascript", js_data)
            else:
                self.send_response(302)
                self.send_header("Location", "https://unpkg.com/html5-qrcode/html5-qrcode.min.js")
                self.end_headers()
                
        elif path == "/get_allowed_operators":
            load_operators()
            names = list(allowed_operators.keys())
            self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "success", "operators": names}).encode("utf-8"))
            
        elif path == "/log_js_error":
            msg = query.get("msg", [""])[0]
            src = query.get("src", [""])[0]
            line = query.get("line", ["0"])[0]
            col = query.get("col", ["0"])[0]
            stack = query.get("stack", [""])[0]
            
            error_msg = f"[JS ERROR] {msg} at {src}:{line}:{col}\nStack: {stack}\n"
            print("\n" + error_msg)
            
            try:
                log_path = os.path.join(get_base_dir(), "js_errors.log")
                with open(log_path, "a", encoding="utf-8") as f:
                    f.write(error_msg + "\n")
            except Exception:
                pass
                
            self.send_response_data("image/gif", b"GIF89a\x01\x00\x01\x00\x80\x00\x00\xff\xff\xff\x00\x00\x00!\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;")
            
        elif path == "/ping":
            self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "success"}).encode("utf-8"))
            
        elif path == "/api/active_operators":
            ops = get_active_operators()
            now = time.time()
            ops_list = []
            for name in sorted(ops):
                last_act = active_operators.get(name, now)
                elapsed = int(now - last_act)
                if elapsed < 5:
                    time_str = "İndi"
                elif elapsed < 60:
                    time_str = f"{elapsed} saniyə əvvəl"
                else:
                    time_str = f"{elapsed // 60} dəqiqə əvvəl"
                ops_list.append({
                    "name": name,
                    "last_activity": last_act,
                    "time_str": time_str
                })
            self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "success", "operators": ops_list}).encode("utf-8"))
            
        elif path == "/login_operator":
            name = query.get("name", [""])[0].strip()
            pin = query.get("pin", [""])[0].strip()
            
            load_operators()
            with db_lock:
                correct_pin = allowed_operators.get(name)
                
            def verify_pin(stored_pin, provided_pin):
                if not stored_pin or not provided_pin:
                    return False
                if stored_pin == provided_pin:
                    return True
                try:
                    from werkzeug.security import check_password_hash
                    return check_password_hash(stored_pin, provided_pin)
                except Exception:
                    return False

            if correct_pin and verify_pin(correct_pin, pin):
                response_data = {"status": "success"}
            else:
                response_data = {"status": "error", "message": "Yanlış PIN kod!"}
                
            self.send_response_data("application/json; charset=utf-8", json.dumps(response_data).encode("utf-8"))
 
        elif path == "/get_product_details":
            barcode = query.get("barcode", [""])[0].strip()
            operator = query.get("operator", [""])[0].strip()
            with db_lock:
                p = products.get(barcode)
            if p:
                own_count = operator_counts.get((barcode, operator))
                
                # Get breakdown of counts
                breakdown_list = []
                for (bc, op), qty in operator_counts.items():
                    if bc == barcode and qty > 0:
                        breakdown_list.append(f"{op}")
                breakdown_str = ", ".join(breakdown_list)
                
                if not breakdown_str and p['yeni'] is not None:
                    op_val = p['operator'] or 'Naməlum'
                    clean_parts = []
                    for part in op_val.split(","):
                        clean_parts.append(part.split(":")[0].strip())
                    breakdown_str = ", ".join(clean_parts)
                
                response_data = {
                    "status": "success",
                    "product": {
                        "barcode": barcode,
                        "brend": p["brend"],
                        "adi": p["adi"],
                        "qaliq": p["qaliq"],
                        "qiymet": p["qiymet"],
                        "yeni": p["yeni"],
                        "operator": breakdown_str if breakdown_str else (p["operator"] or ""),
                        "own_qty": own_count
                     }
                }
            else:
                response_data = {"status": "error", "message": "Məhsul tapılmadı"}
            self.send_response_data("application/json; charset=utf-8", json.dumps(response_data).encode("utf-8"))
            
        elif path == "/get_products":
            with db_lock:
                total_items = len(products)
                total_counted = sum(1 for p in products.values() if p['yeni'] is not None)
                surplus_qty = sum((p['yeni'] - p['qaliq']) for p in products.values() if p['yeni'] is not None and (p['yeni'] - p['qaliq']) > 0)
                shortage_qty = sum(abs(p['yeni'] - p['qaliq']) for p in products.values() if p['yeni'] is not None and (p['yeni'] - p['qaliq']) < 0)
                total_diff_val = sum((p['yeni'] - p['qaliq']) * p['qiymet'] for p in products.values() if p['yeni'] is not None)
                
                response_data = {
                    "products": products,
                    "stats": {
                        "total_items": total_items,
                        "total_counted": total_counted,
                        "surplus_qty": surplus_qty,
                        "shortage_qty": shortage_qty,
                        "total_diff_val": total_diff_val
                    }
                }
            self.send_response_data("application/json; charset=utf-8", json.dumps(response_data).encode("utf-8"))
            
        elif path == "/add_count":
            barcode = query.get("barcode", [""])[0]
            qty_str = query.get("qty", ["0"])[0]
            mode = query.get("mode", ["add"])[0]
            
            with db_lock:
                if barcode in products:
                    try:
                        p = products[barcode]
                        
                        if mode == "reset" or mode == "delete":
                            p['yeni'] = None
                            p['operator'] = ""
                            p['order'] = 0
                            # Remove entries for this barcode from operator_counts
                            to_del = [key for key in operator_counts if key[0] == barcode]
                            for key in to_del:
                                del operator_counts[key]
                        else:
                            qty = float(qty_str)
                            op_name = operator if operator else "Mərkəz"
                            current_own = operator_counts.get((barcode, op_name), 0.0)
                            
                            if mode == "add":
                                new_own = current_own + qty
                            else:
                                new_own = qty
                                
                            operator_counts[(barcode, op_name)] = new_own
                            
                            # sum all operator counts for this barcode
                            total_qty = sum(qty for (bc, op), qty in operator_counts.items() if bc == barcode)
                            p['yeni'] = total_qty
                            
                            # operator list/breakdown
                            active_ops_list = sorted([f"{op}" for (bc, op), q in operator_counts.items() if bc == barcode and q > 0])
                            operators_list = ", ".join(active_ops_list) if active_ops_list else f"{op_name}"
                            p['operator'] = operators_list
                            
                        global count_order
                        if mode != "reset" and mode != "delete":
                            count_order += 1
                            p['order'] = count_order
                        
                        row = p['row']
                        sheet.cell(row=row, column=yeni_sayim_col, value=p['yeni'])
                        sheet.cell(row=row, column=say_ferqi_col, value=f"={yeni_letter}{row}-{qaliq_letter}{row}")
                        sheet.cell(row=row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{row}*{qiymet_letter}{row}")
                        if operator_col:
                            sheet.cell(row=row, column=operator_col, value=p['operator'] if p['yeni'] is not None else None)
                        
                        root.after(0, update_stats)
                        trigger_autosave()
                        play_pc_beep(1000, 150)
                        response_data = {"status": "success", "product": p}
                    except ValueError:
                        play_pc_beep(400, 300)
                        response_data = {"status": "error", "message": "Xətalı say"}
                else:
                    if mode == "reset" or mode == "delete":
                        response_data = {"status": "success", "message": "Məhsul tapılmadı, sıfırlamağa ehtiyac yoxdur"}
                    else:
                        name = query.get("name", [""])[0]
                        price_str = query.get("price", ["0"])[0]
                        stock_str = query.get("stock", ["0"])[0]
                        if name:
                            try:
                                price = float(price_str)
                                stock = float(stock_str)
                                qty = float(qty_str)
                                
                                op_name = operator if operator else "Mərkəz"
                                operator_counts[(barcode, op_name)] = qty
                                operators_list = op_name if qty > 0 else ""
                                
                                new_row = sheet.max_row + 1
                                sheet.cell(row=new_row, column=barkod_col, value=barcode)
                                sheet.cell(row=new_row, column=brend_col, value=name)
                                if adi_col:
                                    sheet.cell(row=new_row, column=adi_col, value="")
                                sheet.cell(row=new_row, column=anbar_qaligi_col, value=stock)
                                sheet.cell(row=new_row, column=qiymet_col, value=price)
                                
                                sheet.cell(row=new_row, column=yeni_sayim_col, value=qty if qty > 0 else None)
                                sheet.cell(row=new_row, column=say_ferqi_col, value=f"={yeni_letter}{new_row}-{qaliq_letter}{new_row}")
                                sheet.cell(row=new_row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{new_row}*{qiymet_letter}{new_row}")
                                if operator_col:
                                    sheet.cell(row=new_row, column=operator_col, value=operators_list)
                                
                                p = {
                                    'row': new_row,
                                    'kod': query.get("kod", [""])[0],
                                    'brend': name,
                                    'adi': '',
                                    'qaliq': stock,
                                    'qiymet': price,
                                    'yeni': qty if qty > 0 else None,
                                    'order': count_order + 1,
                                    'operator': operators_list
                                }
                                count_order += 1
                                products[barcode] = p
                                
                                root.after(0, update_stats)
                                trigger_autosave()
                                play_pc_beep(1200, 250)
                                response_data = {"status": "success", "product": p}
                            except ValueError:
                                play_pc_beep(400, 300)
                                response_data = {"status": "error", "message": "Xətalı parametrlər"}
                        else:
                            play_pc_beep(400, 300)
                            response_data = {"status": "error", "message": "Məhsul tapılmadı"}
                        
            self.send_response_data("application/json; charset=utf-8", json.dumps(response_data).encode("utf-8"))
            
        else:
            self.send_response_data("text/plain; charset=utf-8", b"Not Found", 404)

    def do_POST(self):
        parsed_url = urllib.parse.urlparse(self.path)
        path = parsed_url.path
        
        from http.cookies import SimpleCookie
        cookie_header = self.headers.get('Cookie', '')
        cookie = SimpleCookie(cookie_header)
        is_logged_in = cookie.get('admin_logged_in') and cookie['admin_logged_in'].value == 'true'
        
        if path == "/login":
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                data = json.loads(post_data.decode('utf-8'))
                username = data.get("username", "").strip()
                password = data.get("password", "").strip()
            except Exception:
                username = ""
                password = ""
                
            admin_user = os.environ.get("FLASK_ADMIN_USER", "DoreGence")
            admin_pass = os.environ.get("FLASK_ADMIN_PASS", "Aslanov1991")

            if username == admin_user and password == admin_pass:
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Set-Cookie", "admin_logged_in=true; Path=/; HttpOnly")
                self.end_headers()
                self.wfile.write(json.dumps({"status": "success", "message": "Giris ugurludur."}).encode("utf-8"))
            else:
                self.send_response(401)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.end_headers()
                self.wfile.write(json.dumps({"status": "error", "message": "Istifadeci adi ve ya sifre yanlisdir!"}).encode("utf-8"))
                
        elif path == "/api/operators":
            if not is_logged_in:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Giris icazesi yoxdur!"}).encode("utf-8"), 401)
                return
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                data = json.loads(post_data.decode('utf-8'))
                username = data.get("username", "").strip()
                pin = data.get("pin", "").strip()
                old_username = data.get("old_username", "").strip()
            except Exception:
                username = ""
                pin = ""
                old_username = ""
                
            if not username or not pin:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Ad ve PIN daxil edilmelidir!"}).encode("utf-8"), 400)
                return
                
            try:
                load_operators()
                if old_username and old_username in allowed_operators:
                    if old_username != username:
                        del allowed_operators[old_username]
                allowed_operators[username] = pin
                
                # Write back to operatorlar.txt
                current_dir = get_base_dir()
                file_path = os.path.join(current_dir, "operatorlar.txt")
                with open(file_path, "w", encoding="utf-8") as f:
                    for op, p in allowed_operators.items():
                        f.write(f"{op}:{p}\n")
                        
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "success", "message": "Operator melumatlari yadda saxlanildi."}).encode("utf-8"))
            except Exception as e:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": str(e)}).encode("utf-8"), 500)

        elif path == "/merge_excel_files":
            if not is_logged_in:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Giris icazesi yoxdur!"}).encode("utf-8"), 401)
                return
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            content_type = self.headers.get('Content-Type', '')

            import email.parser
            msg_data = f"Content-Type: {content_type}\r\n\r\n".encode('latin1') + post_data
            p = email.parser.BytesFeedParser()
            p.feed(msg_data)
            msg = p.close()

            files_list = []
            if msg.is_multipart():
                for part in msg.walk():
                    fn = part.get_filename()
                    if fn and fn.endswith('.xlsx'):
                        data = part.get_payload(decode=True)
                        if data:
                            files_list.append((fn, data))

            if not files_list:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Secilmis fayllar arasinda kecerli Excel (.xlsx) tapilmadi!"}).encode("utf-8"), 400)
                return

            try:
                out_wb = merge_excel_files_data(files_list)
                out_stream = io.BytesIO()
                out_wb.save(out_stream)
                excel_bytes = out_stream.getvalue()

                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", "attachment; filename=Yekun_Birlesdirilmis_Sayim.xlsx")
                self.send_header("Content-Length", str(len(excel_bytes)))
                self.end_headers()
                self.wfile.write(excel_bytes)
            except Exception as e:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": str(e)}).encode("utf-8"), 500)

        elif path == "/reset_inventory":
            if not is_logged_in:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Giris icazesi yoxdur!"}).encode("utf-8"), 401)
                return
            try:
                with db_lock:
                    for bc, p in products.items():
                        p['yeni'] = None
                        p['operator'] = ""
                        p['order'] = 0
                    operator_counts.clear()
                    
                    # Write reset changes back to the Excel file
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row=row, column=yeni_sayim_col, value=None)
                        sheet.cell(row=row, column=say_ferqi_col, value=f"={yeni_letter}{row}-{qaliq_letter}{row}")
                        sheet.cell(row=row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{row}*{qiymet_letter}{row}")
                        if operator_col:
                            sheet.cell(row=row, column=operator_col, value=None)
                            
                root.after(0, update_stats)
                trigger_autosave()
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "success", "message": "Butun sayimlar sifirlandi!"}).encode("utf-8"))
            except Exception as e:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": str(e)}).encode("utf-8"), 500)
        else:
            self.send_response_data("text/plain; charset=utf-8", b"Not Found", 404)

    def do_DELETE(self):
        parsed_url = urllib.parse.urlparse(self.path)
        path = parsed_url.path
        
        from http.cookies import SimpleCookie
        cookie_header = self.headers.get('Cookie', '')
        cookie = SimpleCookie(cookie_header)
        is_logged_in = cookie.get('admin_logged_in') and cookie['admin_logged_in'].value == 'true'
        
        if path.startswith("/api/operators/"):
            if not is_logged_in:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Giris icazesi yoxdur!"}).encode("utf-8"), 401)
                return
            username = urllib.parse.unquote(path[len("/api/operators/"):].strip())
            if not username:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Istifadeci adi bosdur!"}).encode("utf-8"), 400)
                return
            try:
                load_operators()
                if username in allowed_operators:
                    del allowed_operators[username]
                    current_dir = get_base_dir()
                    file_path = os.path.join(current_dir, "operatorlar.txt")
                    with open(file_path, "w", encoding="utf-8") as f:
                        for op, p in allowed_operators.items():
                            f.write(f"{op}:{p}\n")
                    self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "success", "message": "Operator silindi."}).encode("utf-8"))
                else:
                    self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": "Operator tapilmadi!"}).encode("utf-8"), 404)
            except Exception as e:
                self.send_response_data("application/json; charset=utf-8", json.dumps({"status": "error", "message": str(e)}).encode("utf-8"), 500)
        else:
            self.send_response_data("text/plain; charset=utf-8", b"Not Found", 404)

def choose_mode_gui():
    mode_win = tk.Tk()
    mode_win.title("Bağlantı Rejimi")
    mode_win.geometry("450x300")
    mode_win.configure(bg=BG_COLOR)
    mode_win.resizable(False, False)
    
    mode_win.update_idletasks()
    x = (mode_win.winfo_screenwidth() - mode_win.winfo_width()) // 2
    y = (mode_win.winfo_screenheight() - mode_win.winfo_height()) // 2
    mode_win.geometry(f"+{x}+{y}")
    
    selected_mode = [None]
    
    tk.Label(mode_win, text="ANBAR SAYIMI BOTU", font=('Segoe UI', 16, 'bold'), bg=BG_COLOR, fg=ACCENT_COLOR).pack(pady=(25, 10))
    tk.Label(mode_win, text="Zəhmət olmasa iş rejimini seçin:", font=('Segoe UI', 11), bg=BG_COLOR, fg=TEXT_COLOR).pack(pady=5)
    
    def set_mode(mode):
        selected_mode[0] = mode
        mode_win.destroy()
        
    def prompt_password(parent):
        pw_dialog = tk.Toplevel(parent)
        pw_dialog.title("Giriş Parolu")
        pw_dialog.geometry("380x200")
        pw_dialog.configure(bg=CARD_BG)
        pw_dialog.resizable(False, False)
        pw_dialog.grab_set()
        
        pw_dialog.update_idletasks()
        dx = parent.winfo_x() + (parent.winfo_width() - pw_dialog.winfo_width()) // 2
        dy = parent.winfo_y() + (parent.winfo_height() - pw_dialog.winfo_height()) // 2
        pw_dialog.geometry(f"+{dx}+{dy}")
        
        tk.Label(pw_dialog, text="Mərkəz Rejimi üçün Parol:", font=('Segoe UI', 11, 'bold'), bg=CARD_BG, fg=ACCENT_COLOR).pack(pady=(20, 10))
        
        ent_pw = tk.Entry(
            pw_dialog, 
            show="*", 
            bg=ENTRY_BG, 
            fg=TEXT_COLOR, 
            insertbackground=TEXT_COLOR, 
            font=('Segoe UI', 12), 
            bd=0, 
            relief="flat", 
            highlightthickness=1, 
            highlightbackground=BORDER_COLOR, 
            highlightcolor=ACCENT_COLOR,
            justify="center"
        )
        ent_pw.pack(fill="x", padx=40, pady=10, ipady=4)
        ent_pw.focus_set()
        
        result = [False]
        
        def check_pw(event=None):
            if ent_pw.get() == "Aslanov1991":
                result[0] = True
                pw_dialog.destroy()
            else:
                messagebox.showerror("Xəta", "Yanlış parol! Giriş rədd edildi.", parent=pw_dialog)
                ent_pw.delete(0, tk.END)
                ent_pw.focus_set()
                
        ent_pw.bind("<Return>", check_pw)
        
        btn_submit = tk.Button(
            pw_dialog,
            text="Təsdiqlə",
            command=check_pw,
            bg=SUCCESS_COLOR,
            fg=BG_COLOR,
            activebackground="#059669",
            activeforeground=BG_COLOR,
            font=('Segoe UI', 10, 'bold'),
            relief="flat",
            bd=0,
            padx=20,
            pady=6,
            cursor="hand2"
        )
        btn_submit.pack(pady=10)
        
        pw_dialog.wait_window(pw_dialog)
        return result[0]

    def check_local_access():
        if prompt_password(mode_win):
            set_mode('local')

    btn_local = tk.Button(
        mode_win,
        text="📂 LOKAL / SERVER REJİMİ (Əsas Kompüter)",
        command=check_local_access,
        bg=ACCENT_COLOR,
        fg=TEXT_COLOR,
        activebackground=ACCENT_HOVER,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 10, 'bold'),
        relief="flat",
        bd=0,
        padx=20,
        pady=10,
        cursor="hand2"
    )
    btn_local.pack(fill="x", padx=40, pady=10)
    
    btn_client = tk.Button(
        mode_win,
        text="🌐 ŞƏBƏKƏ MÜŞTƏRİ REJİMİ (Digər Kompüterlər)",
        command=lambda: set_mode('client'),
        bg=SUCCESS_COLOR,
        fg=BG_COLOR,
        activebackground=SUCCESS_COLOR,
        font=('Segoe UI', 10, 'bold'),
        relief="flat",
        bd=0,
        padx=20,
        pady=10,
        cursor="hand2"
    )
    btn_client.pack(fill="x", padx=40, pady=10)
    
    mode_win.protocol("WM_DELETE_WINDOW", sys.exit)
    mode_win.mainloop()
    return selected_mode[0]

def connect_client_gui():
    conn_win = tk.Tk()
    conn_win.title("Serverə Qoşul")
    conn_win.geometry("400x300")
    conn_win.configure(bg=CARD_BG)
    conn_win.resizable(False, False)
    
    conn_win.update_idletasks()
    x = (conn_win.winfo_screenwidth() - conn_win.winfo_width()) // 2
    y = (conn_win.winfo_screenheight() - conn_win.winfo_height()) // 2
    conn_win.geometry(f"+{x}+{y}")
    
    tk.Label(conn_win, text="Serverə Qoşulma", font=('Segoe UI', 12, 'bold'), bg=CARD_BG, fg=ACCENT_COLOR).pack(pady=15)
    
    form_frame = tk.Frame(conn_win, bg=CARD_BG)
    form_frame.pack(padx=20, fill="both", expand=True)
    form_frame.columnconfigure(1, weight=1)
    
    tk.Label(form_frame, text="Adınız (Operator):", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=0, column=0, sticky="w", pady=8)
    ent_op = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_op.grid(row=0, column=1, sticky="we", pady=8, padx=5)
    ent_op.focus_set()
    
    tk.Label(form_frame, text="Mərkəz IP:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=1, column=0, sticky="w", pady=8)
    ent_ip = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_ip.grid(row=1, column=1, sticky="we", pady=8, padx=5)
    ent_ip.insert(0, "192.168.1.15")
    
    tk.Label(form_frame, text="Port:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=2, column=0, sticky="w", pady=8)
    ent_port = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_port.grid(row=2, column=1, sticky="we", pady=8, padx=5)
    ent_port.insert(0, "5000")
    
    connection_info = [None, None, None]
    
    def on_connect():
        op_name = ent_op.get().strip()
        ip = ent_ip.get().strip()
        port_str = ent_port.get().strip()
        if not op_name:
            messagebox.showerror("Xəta", "Adınızı (Operator) daxil edin!", parent=conn_win)
            return
        if not ip or not port_str:
            messagebox.showerror("Xəta", "İP və Port boş ola bilməz!", parent=conn_win)
            return
        try:
            port = int(port_str)
        except ValueError:
            messagebox.showerror("Xəta", "Port düzgün ədəd olmalıdır!", parent=conn_win)
            return
            
        connection_info[0] = op_name
        connection_info[1] = ip
        connection_info[2] = port
        conn_win.destroy()
        
    btn_conn = tk.Button(
        conn_win,
        text="Qoşul",
        command=on_connect,
        bg=SUCCESS_COLOR,
        fg=BG_COLOR,
        font=('Segoe UI', 10, 'bold'),
        relief="flat",
        bd=0,
        padx=20,
        pady=6,
        cursor="hand2"
    )
    btn_conn.pack(pady=15)
    
    conn_win.protocol("WM_DELETE_WINDOW", sys.exit)
    conn_win.mainloop()
    return connection_info[0], connection_info[1], connection_info[2]

def init_client_data(op_name, ip, port):
    global server_ip, server_port, is_network_mode, products, operator_name
    server_ip = ip
    server_port = port
    is_network_mode = True
    operator_name = op_name
    
    try:
        url = f"http://{server_ip}:{server_port}/get_products?operator={urllib.parse.quote(operator_name)}"
        req = urllib.request.urlopen(url, timeout=3)
        res = json.loads(req.read().decode("utf-8"))
        products = res["products"]
        return True
    except Exception as e:
        messagebox.showerror("Bağlantı Xətası", f"Serverə qoşulmaq mümkün olmadı:\n{str(e)}")
        return False

ssh_is_first_start = True
ssh_tunnel_process = None
ssh_tunnel_url = ""
qr_photo = None

def start_ssh_tunnel():
    global ssh_tunnel_process, ssh_tunnel_url, is_server_active
    def tunnel_worker():
        global ssh_tunnel_process, ssh_tunnel_url, ssh_is_first_start
        import subprocess
        import sys
        
        while is_server_active:
            try:
                # Kill any existing ssh processes to avoid port conflicts on restart
                try:
                    if sys.platform == "win32":
                        subprocess.run(["taskkill", "/F", "/IM", "ssh.exe"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                except Exception:
                    pass
                    
                ssh_tunnel_process = subprocess.Popen(
                    ["ssh", "-o", "StrictHostKeyChecking=no", "-o", "ServerAliveInterval=30", "-o", "ServerAliveCountMax=3", "-R", "80:127.0.0.1:5000", "serveo.net"],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    bufsize=1,
                    stdin=subprocess.DEVNULL,
                    creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
                )
                
                for line in iter(ssh_tunnel_process.stdout.readline, ''):
                    if not is_server_active:
                        break
                    if "Forwarding HTTP traffic from" in line:
                        parts = line.split("from")
                        if len(parts) > 1:
                            ssh_tunnel_url = parts[1].strip()
                            root.after(0, update_server_status_ui)
                            root.after(0, copy_tunnel_url_to_clipboard)
                            if ssh_is_first_start:
                                ssh_is_first_start = False
                                root.after(0, show_qr_code_window)
                
                ssh_tunnel_process.wait()
            except Exception as e:
                print("SSH tunnel background error:", e)
                
            if is_server_active:
                ssh_tunnel_url = ""
                root.after(0, update_server_status_ui)
                time.sleep(5) # Wait 5 seconds before retrying reconnect

    threading.Thread(target=tunnel_worker, daemon=True).start()

def copy_tunnel_url_to_clipboard():
    global ssh_tunnel_url
    if ssh_tunnel_url:
        try:
            root.clipboard_clear()
            root.clipboard_append(ssh_tunnel_url)
            root.update()
            status_var.set(f"İnternet Linki Kopyalandı: {ssh_tunnel_url}")
        except Exception:
            pass

def play_pc_beep(freq=2000, dur=120):
    if freq >= 500:
        freq = 2200
        dur = 120
    def worker():
        try:
            import winsound
            winsound.Beep(freq, dur)
        except Exception:
            pass
    threading.Thread(target=worker, daemon=True).start()

def show_qr_code_window():
    global ssh_tunnel_url, server_ip, server_port, qr_photo
    target_url = ssh_tunnel_url if ssh_tunnel_url else f"http://{server_ip}:{server_port}"
    if not target_url:
        return
        
    import urllib.request
    import urllib.parse
    import base64
    
    try:
        encoded_url = urllib.parse.quote(target_url)
        qr_api_url = f"https://api.qrserver.com/v1/create-qr-code/?size=180x180&data={encoded_url}&format=png"
        req = urllib.request.Request(qr_api_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response:
            image_data = response.read()
            
        qr_photo = tk.PhotoImage(data=base64.b64encode(image_data))
        
        # Destroy duplicate window if already open
        for child in root.winfo_children():
            if isinstance(child, tk.Toplevel) and getattr(child, '_is_qr_window', False):
                child.destroy()
                
        qr_win = tk.Toplevel(root)
        qr_win._is_qr_window = True
        qr_win.title("Qoşulma QR Kodu")
        qr_win.geometry("260x300")
        qr_win.configure(bg=CARD_BG)
        qr_win.resizable(False, False)
        qr_win.grab_set()
        
        # Center dialog
        qr_win.update_idletasks()
        dx = root.winfo_x() + (root.winfo_width() - qr_win.winfo_width()) // 2
        dy = root.winfo_y() + (root.winfo_height() - qr_win.winfo_height()) // 2
        qr_win.geometry(f"+{dx}+{dy}")
        
        tk.Label(
            qr_win, 
            text="Skan Edib Qoşulun", 
            font=('Segoe UI', 12, 'bold'), 
            bg=CARD_BG, 
            fg=ACCENT_COLOR
        ).pack(pady=(15, 10))
        
        lbl_qr = tk.Label(qr_win, image=qr_photo, bg=CARD_BG)
        lbl_qr.pack()
        
        lbl_link = tk.Label(
            qr_win, 
            text=target_url, 
            font=('Segoe UI', 8), 
            bg=CARD_BG, 
            fg=TEXT_MUTED,
            wraplength=230
        )
        lbl_link.pack(pady=(10, 15))
        
    except Exception as e:
        messagebox.showerror("QR Kod Xətası", f"QR kodu yükləmək mümkün olmadı. İnternet bağlantınızı yoxlayın.\n\nXəta: {str(e)}", parent=root)

def stop_ssh_tunnel():
    global ssh_tunnel_process, ssh_tunnel_url
    if ssh_tunnel_process:
        try:
            ssh_tunnel_process.terminate()
            ssh_tunnel_process.wait(timeout=1)
        except Exception:
            try:
                ssh_tunnel_process.kill()
            except Exception:
                pass
        ssh_tunnel_process = None
    ssh_tunnel_url = ""

server_status_timer_id = None

def update_server_status_ui():
    global server_ip, server_port, ssh_tunnel_url, is_server_active, server_status_timer_id
    if not is_server_active:
        return
        
    if server_status_timer_id is not None:
        try:
            root.after_cancel(server_status_timer_id)
        except Exception:
            pass
        server_status_timer_id = None
        
    ops = get_active_operators()
    ops_str = f"Qoşulanlar: {len(ops)} ({', '.join(ops) if ops else 'yoxdur'})"
    status_text = f"Lokal: http://{server_ip}:{server_port}"
    if ssh_tunnel_url:
        status_text += f" | 🌐 İnternet: {ssh_tunnel_url}"
    status_text += f" | {ops_str}"
    
    if 'lbl_server_status' in globals() and lbl_server_status:
        lbl_server_status.config(text=f"🟢 {status_text}", fg=SUCCESS_COLOR)
        
    server_status_timer_id = root.after(5000, update_server_status_ui)

def stop_server():
    global httpd, is_server_active, server_status_timer_id, ssh_is_first_start
    stop_ssh_tunnel()
    ssh_is_first_start = True
    
    if server_status_timer_id is not None:
        try:
            root.after_cancel(server_status_timer_id)
        except Exception:
            pass
        server_status_timer_id = None
        
    if httpd:
        try:
            threading.Thread(target=httpd.shutdown, daemon=True).start()
            httpd.server_close()
        except Exception as e:
            print("Server dayandırma xətası:", e)
    httpd = None
    is_server_active = False
    status_var.set("Şəbəkə serveri deaktiv edildi.")
    if 'lbl_server_status' in globals() and lbl_server_status:
        lbl_server_status.config(text="🔴 Server Qapalıdır", fg=DANGER_COLOR)
    if 'btn_start_server' in globals() and btn_start_server:
        btn_start_server.config(text="🌐 Şəbəkəni Başlat", bg=ACCENT_COLOR, activebackground=ACCENT_HOVER)
    if 'btn_show_qr' in globals() and btn_show_qr:
        btn_show_qr.pack_forget()

def download_scanner_js():
    try:
        script_dir = get_base_dir()
        js_filepath = os.path.join(script_dir, "html5-qrcode.min.js")
        if not os.path.exists(js_filepath):
            import urllib.request
            url = "https://unpkg.com/html5-qrcode/html5-qrcode.min.js"
            urllib.request.urlretrieve(url, js_filepath)
    except Exception as e:
        print("Offline scanner library download failed:", e)

def start_server_thread():
    global httpd, server_ip, server_port, is_server_active
    server_port = 5000
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        server_ip = s.getsockname()[0]
        s.close()
    except Exception:
        server_ip = "127.0.0.1"
        
    server_address = ('', server_port)
    try:
        httpd = ThreadingHTTPServer(server_address, WarehouseHTTPRequestHandler)
        # Download library for offline camera scan support in background
        threading.Thread(target=download_scanner_js, daemon=True).start()
        
        t = threading.Thread(target=httpd.serve_forever, daemon=True)
        t.start()
        is_server_active = True
        status_var.set(f"Şəbəkə Serveri Aktivdir: {server_ip}:{server_port}")
        
        # Start SSH tunnel in the background
        start_ssh_tunnel()
        update_server_status_ui()
        
        if 'btn_show_qr' in globals() and btn_show_qr:
            btn_show_qr.pack(side="left", padx=5)
            
        if 'btn_start_server' in globals() and btn_start_server:
            btn_start_server.config(
                text="🛑 Şəbəkəni Deaktiv Et", 
                bg=DANGER_COLOR, 
                activebackground="#B91C1C"
            )
    except Exception as e:
        messagebox.showerror("Server Xətası", f"Serveri başlatmaq mümkün olmadı:\n{str(e)}")

def toggle_server():
    global is_server_active
    if is_server_active:
        stop_server()
    else:
        start_server_thread()

# Setup workbook and map columns
def init_excel_data():
    global filepath, wb, sheet, products
    global brend_col, adi_col, anbar_qaligi_col, barkod_col, kod_col, yeni_sayim_col, qiymet_col, say_ferqi_col, qiymet_ferqi_col, operator_col
    global yeni_letter, qaliq_letter, say_ferqi_letter, qiymet_letter

    try:
        wb = openpyxl.load_workbook(filepath, data_only=False)
        sheet = wb.active
    except Exception as e:
        messagebox.showerror("Xəta", f"Excel faylını açmaq mümkün olmadı:\n{str(e)}")
        sys.exit(0)
        
    # Map headers
    headers = {}
    for col in range(1, sheet.max_column + 1):
        val = sheet.cell(row=1, column=col).value
        if val:
            headers[normalize_header(val)] = col
            
    # Required columns check
    brend_col = headers.get('brend') or headers.get('brand')
    adi_col = headers.get('mehsulun adi') or headers.get('mehsulun_adi') or headers.get('adi') or headers.get('name') or headers.get('mehsul') or headers.get('description')
    
    if not brend_col:
        brend_col = adi_col
        if brend_col == adi_col:
            adi_col = None
            
    anbar_qaligi_col = headers.get('anbar qaligi') or headers.get('qaliq') or headers.get('stock') or headers.get('sistem qaligi')
    
    # Map barcode and code columns
    barkod_col = headers.get('barkod') or headers.get('barkkod') or headers.get('barcode')
    kod_col = headers.get('kod') or headers.get('kodu') or headers.get('mehsulun kodu') or headers.get('mehsulun_kodu') or headers.get('code') or headers.get('product code') or headers.get('product_code')
    if not barkod_col and kod_col:
        barkod_col = kod_col
        kod_col = None
    elif not barkod_col:
        # Fallback to older code column matching if none of above matches
        barkod_col = headers.get('kod')
        
    qiymet_col = headers.get('mehsulun qiymeti') or headers.get('qiymet') or headers.get('qiymeti') or headers.get('price') or headers.get('satis qiymeti') or headers.get('satis_qiymeti') or headers.get('cena') or headers.get('цена') or headers.get('satis qiymet') or headers.get('maya qiymeti') or headers.get('fiyat')
    
    missing_cols = []
    if not brend_col: missing_cols.append("Brend (və ya Adı/Məhsulun Adı)")
    if not anbar_qaligi_col: missing_cols.append("Anbar Qaligi (və ya Qalıq)")
    if not barkod_col: missing_cols.append("Barkod (və ya Kod)")
    
    if missing_cols:
        messagebox.showerror(
            "Başlıq Xətası", 
            f"Excel faylında aşağıdakı vacib sütun başlıqları tapılmadı:\n\n" + 
            "\n".join([f"- {c}" for c in missing_cols]) + 
            "\n\nZəhmət olmasa faylın ilk sətirində bu başlıqların mövcud olduğunu yoxlayın."
        )
        sys.exit(0)
        
    # Check or create calculation columns
    yeni_sayim_col = headers.get('yeni sayim') or headers.get('real say')
    say_ferqi_col = headers.get('say ferqi')
    qiymet_ferqi_col = headers.get('qiymet ferqi')
    
    max_col = sheet.max_column
    if not qiymet_col:
        max_col += 1
        sheet.cell(row=1, column=max_col, value="Qiyməti")
        qiymet_col = max_col
        
    if not yeni_sayim_col:
        max_col += 1
        sheet.cell(row=1, column=max_col, value="Yeni Sayim")
        yeni_sayim_col = max_col
    if not say_ferqi_col:
        max_col += 1
        sheet.cell(row=1, column=max_col, value="Say ferqi")
        say_ferqi_col = max_col
    if not qiymet_ferqi_col:
        max_col += 1
        sheet.cell(row=1, column=max_col, value="Qiymet ferqi")
        qiymet_ferqi_col = max_col
        
    # Check or create Operator column
    operator_col = headers.get('operator') or headers.get('sayimci') or headers.get('user')
    if not operator_col:
        max_col += 1
        sheet.cell(row=1, column=max_col, value="Operator")
        operator_col = max_col
        
    # Get column letters for formulas
    yeni_letter = get_column_letter(yeni_sayim_col)
    qaliq_letter = get_column_letter(anbar_qaligi_col)
    say_ferqi_letter = get_column_letter(say_ferqi_col)
    qiymet_letter = get_column_letter(qiymet_col)
    
    # Load products into memory
    products.clear()
    operator_counts.clear()
    for row_idx in range(2, sheet.max_row + 1):
        barcode_val = sheet.cell(row=row_idx, column=barkod_col).value
        if barcode_val is not None:
            # Handle float format issues (e.g. 1.23456e+12)
            if isinstance(barcode_val, float):
                barcode_str = str(int(barcode_val)).strip()
            else:
                barcode_str = str(barcode_val).strip()
                
            if not barcode_str:
                continue
                
            kod_val = ""
            if kod_col:
                k_val = sheet.cell(row=row_idx, column=kod_col).value
                if k_val is not None:
                    if isinstance(k_val, float):
                        kod_val = str(int(k_val)).strip()
                    else:
                        kod_val = str(k_val).strip()
                        
            brend_val = sheet.cell(row=row_idx, column=brend_col).value or "Naməlum Brend"
            adi_val = ""
            if adi_col:
                adi_val = sheet.cell(row=row_idx, column=adi_col).value
            adi_val = str(adi_val).strip() if adi_val else ""
            
            # Stock quantity
            qaliq_val = sheet.cell(row=row_idx, column=anbar_qaligi_col).value
            try:
                qaliq = float(qaliq_val) if qaliq_val is not None else 0.0
            except ValueError:
                qaliq = 0.0
                
            # Price
            qiymet_val = sheet.cell(row=row_idx, column=qiymet_col).value
            try:
                qiymet = float(qiymet_val) if qiymet_val is not None else 0.0
            except ValueError:
                qiymet = 0.0
                
            # Existing counted count
            yeni_val = sheet.cell(row=row_idx, column=yeni_sayim_col).value
            # Check if it starts with '=' (formula)
            if yeni_val is not None and str(yeni_val).startswith('='):
                yeni = None
            else:
                try:
                    yeni = float(yeni_val) if yeni_val is not None else None
                except ValueError:
                    yeni = None
                    
            op_val = sheet.cell(row=row_idx, column=operator_col).value if operator_col else ""
            
            # Load existing operator counts if any
            if yeni is not None:
                op_str = str(op_val).strip() if op_val else "Mərkəz"
                if ":" in op_str:
                    parts = op_str.split(",")
                    parsed_any = False
                    for part in parts:
                        if ":" in part:
                            subparts = part.split(":")
                            sub_op = subparts[0].strip()
                            try:
                                sub_qty = float(subparts[1].strip())
                                operator_counts[(barcode_str, sub_op)] = sub_qty
                                parsed_any = True
                            except ValueError:
                                pass
                    if not parsed_any:
                        operator_counts[(barcode_str, op_str)] = yeni
                else:
                    first_op = op_str.split(",")[0].strip()
                    operator_counts[(barcode_str, first_op if first_op else "Mərkəz")] = yeni

            clean_op = ""
            if op_val:
                op_str = str(op_val).strip()
                clean_parts = []
                for part in op_str.split(","):
                    clean_parts.append(part.split(":")[0].strip())
                clean_op = ", ".join(clean_parts)

            products[barcode_str] = {
                'row': row_idx,
                'kod': kod_val,
                'brend': str(brend_val).strip(),
                'adi': adi_val,
                'qaliq': qaliq,
                'qiymet': qiymet,
                'yeni': yeni,
                'order': 0,
                'operator': clean_op
            }

# Choose startup mode
startup_mode = choose_mode_gui()
if not startup_mode:
    sys.exit(0)

if startup_mode == 'client':
    op_name, ip, port = connect_client_gui()
    if not op_name or not ip or not port:
        sys.exit(0)
    if not init_client_data(op_name, ip, port):
        sys.exit(0)
    filepath = f"{ip}:{port}"
else:
    # Find appropriate Excel file (Local/Master mode)
    while True:
        files = scan_excel_files()
        if len(files) == 1:
            filepath = files[0]
            break
        elif len(files) > 1:
            filepath = select_file_gui(files)
            if not filepath:
                sys.exit(0)
            break
        else:
            retry = show_no_files_gui()
            if not retry:
                sys.exit(0)

    # Load data
    init_excel_data()

# --- BUILD MAIN WINDOW ---
root = tk.Tk()
root.title("Anbar Sayımı Botu v1.0")
root.geometry("1200x700")
root.configure(bg=BG_COLOR)

# Center main window
root.update_idletasks()
x = (root.winfo_screenwidth() - root.winfo_width()) // 2
y = (root.winfo_screenheight() - root.winfo_height()) // 2
root.geometry(f"+{x}+{y}")

# Variables for Tkinter
barcode_var = tk.StringVar()
qty_var = tk.StringVar()
current_product_var = tk.StringVar()
status_var = tk.StringVar(value="Sistem hazır. Skan gözlənilir...")
current_brand_filter = ""
history_search_var = tk.StringVar()
brand_search_var = tk.StringVar()

# Tkinter Layout Styles
style = ttk.Style()
style.theme_use('clam')
style.configure('.', background=BG_COLOR, foreground=TEXT_COLOR)
style.configure('TFrame', background=BG_COLOR)
style.configure('Card.TFrame', background=CARD_BG, relief="flat", borderwidth=0)
style.configure('TLabel', background=BG_COLOR, foreground=TEXT_COLOR, font=('Segoe UI', 10))

# Notebook Styling (Flat, matte)
style.configure('TNotebook', background=BG_COLOR, borderwidth=0, shiftwidth=0)
style.configure('TNotebook.Tab', 
                background=CARD_BG, 
                foreground=TEXT_MUTED, 
                padding=[12, 6], 
                font=('Segoe UI', 9, 'bold'),
                borderwidth=0,
                relief="flat")
style.map('TNotebook.Tab', 
          background=[('selected', ACCENT_COLOR)], 
          foreground=[('selected', TEXT_COLOR)])

# Treeview Styling (Flat, matte)
style.configure('Treeview', 
                background=CARD_BG, 
                fieldbackground=CARD_BG, 
                foreground=TEXT_COLOR, 
                font=('Segoe UI', 8), 
                rowheight=18,
                borderwidth=0,
                relief="flat")
style.configure('Treeview.Heading', 
                background=ENTRY_BG, 
                foreground=TEXT_COLOR, 
                font=('Segoe UI', 8, 'bold'),
                borderwidth=0,
                relief="flat")
style.map('Treeview.Heading',
          background=[('active', BORDER_COLOR)],
          foreground=[('active', TEXT_COLOR)])

# Remove Treeview border element in clam layout
style.layout("Treeview", [('Treeview.treearea', {'sticky': 'nswe'})])

# Scrollbar Styling
style.configure('Vertical.TScrollbar', 
                background=ENTRY_BG, 
                troughcolor=BG_COLOR, 
                arrowcolor=TEXT_MUTED, 
                borderwidth=0, 
                gripcount=0)

# --- FUNCTIONS ---

def clear_inventory_counts_gui():
    if is_network_mode:
        messagebox.showerror("Xəta", "Şəbəkə müştəri rejimində sayımı sıfırlamaq olmaz!")
        return
        
    confirm = messagebox.askyesno(
        "Sayımı Sıfırla", 
        "DİQQƏT: Bütün operatorların daxil etdiyi sayım nəticələri sıfırlanacaq və Excel faylı təmizlənəcək!\n\n"
        "Bunu etmək istədiyinizdən əminsiniz? Bu əməliyyat geri qaytarıla bilməz!",
        icon="warning"
    )
    if not confirm:
        return
        
    confirm2 = messagebox.askyesno(
        "Təsdiq", 
        "Sonuncu xəbərdarlıq: Bütün sayımları silmək istəyirsiniz?"
    )
    if not confirm2:
        return
        
    with db_lock:
        operator_counts.clear()
        for bc, p in products.items():
            p['yeni'] = None
            p['operator'] = ""
            p['order'] = 0
            
            row = p['row']
            sheet.cell(row=row, column=yeni_sayim_col, value=None)
            sheet.cell(row=row, column=say_ferqi_col, value=f"={yeni_letter}{row}-{qaliq_letter}{row}")
            sheet.cell(row=row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{row}*{qiymet_letter}{row}")
            if operator_col:
                sheet.cell(row=row, column=operator_col, value=None)
                
    update_stats()
    trigger_autosave()
    status_var.set("Bütün sayım məlumatları sıfırlandı!")

def on_history_double_click(event):
    selected_item = history_tree.selection()
    if selected_item:
        item_vals = history_tree.item(selected_item[0], "values")
        if item_vals:
            barcode_var.set(item_vals[1])
            on_barcode_enter()

def on_brand_product_double_click(event):
    selected_item = brand_tree.selection()
    if selected_item:
        item_vals = brand_tree.item(selected_item[0], "values")
        if item_vals:
            barcode_var.set(item_vals[1])
            on_barcode_enter()

def refresh_counted_products():
    if 'history_tree' not in globals():
        return
    for item in history_tree.get_children():
        history_tree.delete(item)
        
    search_q = history_search_var.get().strip().lower() if 'history_search_var' in globals() else ""
        
    counted_list = []
    with db_lock:
        for bc, p in products.items():
            if p['yeni'] is not None and p.get('operator'):
                display_name = f"{p['brend']} - {p['adi']}" if p.get('adi') else p['brend']
                if not search_q or search_q in bc.lower() or search_q in display_name.lower() or search_q in p.get('kod', '').lower():
                    counted_list.append((bc, p.copy()))
            
    # Sort by -order (most recent first), then by brand and adi
    counted_list.sort(key=lambda x: (-x[1].get('order', 0), x[1]['brend'].lower(), x[1].get('adi', '').lower()))
    
    for bc, p in counted_list:
        say_ferqi = p['yeni'] - p['qaliq']
        qiymet_ferqi = say_ferqi * p['qiymet']
        display_name = f"{p['brend']} - {p['adi']}" if p.get('adi') else p['brend']
        
        # Determine color tags and text labels: + Green, - Red, 0 Yellow
        if say_ferqi > 0:
            row_tag = 'surplus'
            ferq_text = f"{say_ferqi:+.2f} [ARTIQ]"
        elif say_ferqi < 0:
            row_tag = 'shortage'
            ferq_text = f"{say_ferqi:+.2f} [ƏSKİK]"
        else:
            row_tag = 'equal'
            ferq_text = f"{say_ferqi:+.2f} [DUZ]"
            
        op = p.get('operator', '')
        history_tree.insert("", "end", values=(
            p.get('kod', ''),
            bc,
            display_name,
            f"{p['qaliq']:.2f}",
            f"{p['yeni']:.2f}",
            ferq_text,
            f"{qiymet_ferqi:+.2f} AZN",
            op
        ), tags=(row_tag,))

def refresh_brand_products():
    if 'brand_tree' not in globals():
        return
    # Clear brand_tree
    for item in brand_tree.get_children():
        brand_tree.delete(item)
        
    search_q = brand_search_var.get().strip().lower() if 'brand_search_var' in globals() else ""
        
    # Get a safe copy of products
    with db_lock:
        local_products = {bc: p.copy() for bc, p in products.items()}
        
    # Populate
    for bc, p in local_products.items():
        if not current_brand_filter or p['brend'] == current_brand_filter:
            display_name = f"{p['brend']} - {p['adi']}" if p['adi'] else p['brend']
            if not search_q or search_q in bc.lower() or search_q in display_name.lower() or search_q in p.get('kod', '').lower():
                # Determine counted string with soft color tags: + Green, - Red, 0 Yellow
                row_tag = ''
                if p['yeni'] is not None:
                    say_ferqi = p['yeni'] - p['qaliq']
                    if say_ferqi > 0:
                        row_tag = 'surplus'
                        counted_str = f"{p['yeni']:.2f} [ARTIQ]"
                    elif say_ferqi < 0:
                        row_tag = 'shortage'
                        counted_str = f"{p['yeni']:.2f} [ƏSKİK]"
                    else:
                        row_tag = 'equal'
                        counted_str = f"{p['yeni']:.2f} [DUZ]"
                else:
                    counted_str = "-"
                
                op = p.get('operator', '')
                brand_tree.insert("", "end", values=(
                    p.get('kod', ''),
                    bc,
                    display_name,
                    f"{p['qaliq']:.2f}",
                    counted_str,
                    f"{p['qiymet']:.2f} AZN",
                    op
                ), tags=(row_tag,) if row_tag else ())

def set_brand_filter(brand_name):
    global current_brand_filter
    current_brand_filter = brand_name
    if brand_name:
        btn_brand_select.config(text=f"🔍 Brend: {brand_name}", bg=SUCCESS_COLOR)
        status_var.set(f"Brend süzgəci aktivləşdirildi: {brand_name}")
    else:
        btn_brand_select.config(text="🔍 Brend Seç: Hamısı", bg=ACCENT_COLOR)
        status_var.set("Bütün brendlər göstərilir.")
    refresh_brand_products()

def select_brand_gui():
    # Get unique brands
    with db_lock:
        brands = sorted(list(set(p['brend'] for p in products.values() if p['brend'])))
    
    dialog = tk.Toplevel(root)
    dialog.title("Brend Seçimi")
    dialog.geometry("400x500")
    dialog.configure(bg=BG_COLOR)
    dialog.resizable(False, False)
    dialog.grab_set()
    
    # Center
    dialog.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() - dialog.winfo_width()) // 2
    y = root.winfo_y() + (root.winfo_height() - dialog.winfo_height()) // 2
    dialog.geometry(f"+{x}+{y}")
    
    title_lbl = tk.Label(
        dialog, 
        text="Sayılacaq Brendi Seçin", 
        font=('Segoe UI', 13, 'bold'), 
        bg=BG_COLOR, 
        fg=ACCENT_COLOR
    )
    title_lbl.pack(pady=(15, 10))
    
    search_frame = tk.Frame(dialog, bg=BG_COLOR)
    search_frame.pack(fill="x", padx=20, pady=(0, 10))
    
    tk.Label(search_frame, text="Axtarış:", bg=BG_COLOR, fg=TEXT_MUTED).pack(side="left", padx=(0, 5))
    search_var = tk.StringVar()
    search_ent = tk.Entry(
        search_frame, 
        textvariable=search_var, 
        bg=ENTRY_BG, 
        fg=TEXT_COLOR, 
        insertbackground=TEXT_COLOR, 
        font=('Segoe UI', 10), 
        bd=0, 
        relief="flat", 
        highlightthickness=1, 
        highlightbackground=BORDER_COLOR, 
        highlightcolor=ACCENT_COLOR
    )
    search_ent.pack(side="left", fill="x", expand=True, ipady=4)
    search_ent.focus_set()
    
    list_frame = tk.Frame(dialog, bg=BG_COLOR)
    list_frame.pack(padx=20, fill="both", expand=True)
    
    scrollbar = ttk.Scrollbar(list_frame)
    scrollbar.pack(side="right", fill="y")
    
    listbox = tk.Listbox(
        list_frame, 
        bg=CARD_BG, 
        fg=TEXT_COLOR, 
        selectbackground=ACCENT_COLOR, 
        selectforeground=TEXT_COLOR,
        font=('Segoe UI', 10), 
        bd=0, 
        relief="flat",
        highlightthickness=1,
        highlightbackground=BORDER_COLOR,
        highlightcolor=ACCENT_COLOR,
        yscrollcommand=scrollbar.set
    )
    listbox.pack(side="left", fill="both", expand=True)
    scrollbar.config(command=listbox.yview)
    
    def populate_list(filter_text=""):
        listbox.delete(0, tk.END)
        listbox.insert(tk.END, "[ Bütün Brendlər ]")
        for b in brands:
            if not filter_text or filter_text.lower() in b.lower():
                listbox.insert(tk.END, b)
        listbox.selection_set(0)
        
    populate_list()
    
    def on_search_key(event):
        populate_list(search_var.get())
        
    search_ent.bind("<KeyRelease>", on_search_key)
    
    selected_brand = [None]
    
    def on_select(event=None):
        sel = listbox.curselection()
        if sel:
            choice = listbox.get(sel[0])
            if choice == "[ Bütün Brendlər ]":
                selected_brand[0] = ""
            else:
                selected_brand[0] = choice
            dialog.destroy()
            
    def on_cancel():
        dialog.destroy()
        
    listbox.bind("<Return>", on_select)
    listbox.bind("<Double-Button-1>", on_select)
    dialog.bind("<Escape>", lambda e: on_cancel())
    
    btn_frame = tk.Frame(dialog, bg=BG_COLOR)
    btn_frame.pack(pady=15)
    
    btn_choose = tk.Button(
        btn_frame, 
        text="Seç", 
        command=on_select, 
        bg=ACCENT_COLOR, 
        fg=TEXT_COLOR, 
        activebackground=ACCENT_HOVER,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 10, 'bold'), 
        relief="flat", 
        bd=0,
        padx=20, 
        pady=6,
        cursor="hand2"
    )
    btn_choose.pack(side="left", padx=5)
    
    btn_cancel = tk.Button(
        btn_frame, 
        text="Ləğv et", 
        command=on_cancel, 
        bg=ENTRY_BG, 
        fg=TEXT_COLOR, 
        activebackground=BORDER_COLOR,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 10), 
        relief="flat", 
        bd=0,
        padx=15, 
        pady=6,
        cursor="hand2"
    )
    btn_cancel.pack(side="left", padx=5)
    
    dialog.wait_window(dialog)
    
    if selected_brand[0] is not None:
        set_brand_filter(selected_brand[0])

def show_search_results_gui(candidates, search_term):
    result = [None]
    
    dialog = tk.Toplevel(root)
    dialog.title(f"Axtarış: '{search_term}'")
    dialog.geometry("600x420")
    dialog.configure(bg=BG_COLOR)
    dialog.resizable(False, False)
    dialog.grab_set()
    
    # Center
    dialog.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() - dialog.winfo_width()) // 2
    y = root.winfo_y() + (root.winfo_height() - dialog.winfo_height()) // 2
    dialog.geometry(f"+{x}+{y}")
    
    title_lbl = tk.Label(
        dialog, 
        text=f"Axtarış nəticələri: '{search_term}'", 
        font=('Segoe UI', 12, 'bold'), 
        bg=BG_COLOR, 
        fg=ACCENT_COLOR
    )
    title_lbl.pack(pady=(15, 5))
    
    desc_lbl = tk.Label(
        dialog, 
        text=f"Siyahıdan uyğun məhsulu seçin (Arrow keys + Enter və ya cüt klik):", 
        font=('Segoe UI', 9), 
        bg=BG_COLOR, 
        fg=TEXT_MUTED
    )
    desc_lbl.pack(pady=(0, 10))
    
    list_frame = tk.Frame(dialog, bg=BG_COLOR)
    list_frame.pack(padx=20, fill="both", expand=True)
    
    scrollbar = ttk.Scrollbar(list_frame)
    scrollbar.pack(side="right", fill="y")
    
    listbox = tk.Listbox(
        list_frame, 
        bg=CARD_BG, 
        fg=TEXT_COLOR, 
        selectbackground=ACCENT_COLOR, 
        selectforeground=TEXT_COLOR,
        font=('Segoe UI', 10), 
        bd=0, 
        relief="flat",
        highlightthickness=1,
        highlightbackground=BORDER_COLOR,
        highlightcolor=ACCENT_COLOR,
        yscrollcommand=scrollbar.set
    )
    listbox.pack(side="left", fill="both", expand=True)
    scrollbar.config(command=listbox.yview)
    
    for bc, p in candidates:
        display_name = f"{p['brend']} - {p['adi']}" if p['adi'] else p['brend']
        listbox.insert(tk.END, f"{display_name} ({bc}) - Qalıq: {p['qaliq']:.2f} | Qiymət: {p['qiymet']:.2f} AZN")
        
    listbox.selection_set(0)
    listbox.focus_set()
    
    def on_select(event=None):
        sel = listbox.curselection()
        if sel:
            result[0] = candidates[sel[0]][0]
            dialog.destroy()
            
    def on_cancel():
        dialog.destroy()
        
    listbox.bind("<Return>", on_select)
    listbox.bind("<Double-Button-1>", on_select)
    dialog.bind("<Escape>", lambda e: on_cancel())
    
    btn_frame = tk.Frame(dialog, bg=BG_COLOR)
    btn_frame.pack(pady=15)
    
    btn_choose = tk.Button(
        btn_frame, 
        text="Seç", 
        command=on_select, 
        bg=ACCENT_COLOR, 
        fg=TEXT_COLOR, 
        activebackground=ACCENT_HOVER,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 10, 'bold'), 
        relief="flat", 
        bd=0,
        padx=20, 
        pady=6,
        cursor="hand2"
    )
    btn_choose.pack(side="left", padx=5)
    
    btn_cancel = tk.Button(
        btn_frame, 
        text="Ləğv et", 
        command=on_cancel, 
        bg=ENTRY_BG, 
        fg=TEXT_COLOR, 
        activebackground=BORDER_COLOR,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 10), 
        relief="flat", 
        bd=0,
        padx=15, 
        pady=6,
        cursor="hand2"
    )
    btn_cancel.pack(side="left", padx=5)
    
    dialog.wait_window(dialog)
    return result[0]

def on_barcode_enter(event=None):
    barcode = barcode_var.get().strip()
    if not barcode:
        return
        
    matched_barcode = None
    
    with db_lock:
        # 1. Exact match check
        if barcode in products:
            p = products[barcode]
            # Only select if it fits the brand filter (if active)
            if not current_brand_filter or p['brend'] == current_brand_filter:
                matched_barcode = barcode
                
        # 2. Candidate match check
        if not matched_barcode:
            candidates = []
            is_numeric_suffix = barcode.isdigit() and len(barcode) >= 4
            
            for bc_key, p in products.items():
                # If brand filter is active, only consider products of this brand
                if current_brand_filter and p['brend'] != current_brand_filter:
                    continue
                    
                # Check barcode suffix match
                if is_numeric_suffix and bc_key.endswith(barcode):
                    candidates.append((bc_key, p.copy()))
                    continue
                    
                # Check brand name or product name or kod substring match
                search_str = barcode.lower()
                if search_str in p['brend'].lower() or (p['adi'] and search_str in p['adi'].lower()) or (p.get('kod') and search_str in p['kod'].lower()):
                    candidates.append((bc_key, p.copy()))
                    
            if len(candidates) == 1:
                matched_barcode = candidates[0][0]
            elif len(candidates) > 1:
                matched_barcode = show_search_results_gui(candidates, barcode)
                if not matched_barcode:
                    ent_barcode.focus_set()
                    return
            else:
                matched_barcode = None

    if matched_barcode:
        play_pc_beep(800, 100)
        with db_lock:
            p = products[matched_barcode].copy()
        current_product_var.set(matched_barcode)
        
        # Display product details (Brand + Name)
        display_name = f"{p['brend']} - {p['adi']}" if p['adi'] else p['brend']
        lbl_product_name.config(text=display_name, fg=TEXT_COLOR)
        lbl_system_stock.config(text=f"{p['qaliq']:.2f}")
        lbl_price.config(text=f"{p['qiymet']:.2f} AZN")
        
        if p['yeni'] is not None:
            lbl_prev_count.config(text=f"{p['yeni']:.2f}", fg=SUCCESS_COLOR)
            try:
                yeni_val = float(p['yeni'])
                qty_var.set(str(int(yeni_val) if yeni_val.is_integer() else yeni_val))
            except (ValueError, TypeError):
                qty_var.set(str(p['yeni']))
        else:
            lbl_prev_count.config(text="Sayılmayıb", fg=TEXT_MUTED)
            qty_var.set("")
            
        # Enable quantity input, focus, and select all
        ent_qty.config(state="normal")
        ent_qty.focus_set()
        ent_qty.select_range(0, tk.END)
        
        # Reset difference labels to blank until new qty entered
        lbl_diff_qty.config(text="-", fg=TEXT_MUTED)
        lbl_diff_val.config(text="-", fg=TEXT_MUTED)
        
        status_var.set("Məhsul tapıldı. Sayı daxil edib Enter sıxın.")
    else:
        root.bell()
        lbl_product_name.config(text="Məhsul tapılmadı!", fg=DANGER_COLOR)
        lbl_system_stock.config(text="-")
        lbl_price.config(text="-")
        lbl_prev_count.config(text="-", fg=TEXT_MUTED)
        ent_qty.config(state="disabled")
        qty_var.set("")
        
        lbl_diff_qty.config(text="-", fg=TEXT_MUTED)
        lbl_diff_val.config(text="-", fg=TEXT_MUTED)
        
        status_var.set("XƏTA: Bu məhsul/barkod bazada yoxdur və ya hazırkı brendə aid deyil!")

def on_qty_enter(event=None):
    barcode = current_product_var.get()
    qty_str = qty_var.get().strip()
    
    if not barcode or barcode not in products:
        return
        
    try:
        qty = float(qty_str)
        if qty < 0:
            raise ValueError()
    except ValueError:
        messagebox.showerror("Xəta", "Zəhmət olmasa düzgün say daxil edin (müsbət ədəd).")
        ent_qty.focus_set()
        return
        
    if is_network_mode:
        try:
            url = f"http://{server_ip}:{server_port}/add_count?barcode={barcode}&qty={qty}&mode=set&operator={urllib.parse.quote(operator_name)}"
            req = urllib.request.urlopen(url, timeout=3)
            res = json.loads(req.read().decode("utf-8"))
            if res.get("status") == "success":
                update_stats()
                barcode_var.set("")
                qty_var.set("")
                ent_qty.config(state="disabled")
                ent_barcode.focus_set()
                p = products[barcode]
                display_name = f"{p['brend']} - {p['adi']}" if p.get('adi') else p['brend']
                status_var.set(f"Yadda saxlanıldı: '{display_name}' - Say: {qty:.2f}")
            else:
                messagebox.showerror("Xəta", f"Server xətası: {res.get('message')}")
        except Exception as e:
            messagebox.showerror("Xəta", f"Serverlə bağlantı kəsildi:\n{str(e)}")
        return

    with db_lock:
        p = products[barcode]
        
        op_name = operator_name if operator_name else "Mərkəz"
        operator_counts[(barcode, op_name)] = qty
        
        # Calculate total count from all operators
        total_qty = sum(q for (bc, op), q in operator_counts.items() if bc == barcode)
        p['yeni'] = total_qty
        
        # Operator breakdown string
        active_ops_list = sorted([f"{op}: {q:.2f}" for (bc, op), q in operator_counts.items() if bc == barcode and q > 0])
        operators_list = ", ".join(active_ops_list) if active_ops_list else f"{op_name}: {qty:.2f}"
        p['operator'] = operators_list
        
        global count_order
        count_order += 1
        p['order'] = count_order
        
        # Calculate differences for display
        say_ferqi = total_qty - p['qaliq']
        qiymet_ferqi = say_ferqi * p['qiymet']
        
        # Display difference immediately
        if say_ferqi > 0:
            lbl_diff_qty.config(text=f"+{say_ferqi:.2f} (Artıq)", fg=SUCCESS_COLOR)
        elif say_ferqi < 0:
            lbl_diff_qty.config(text=f"{say_ferqi:.2f} (Əskik)", fg=DANGER_COLOR)
        else:
            lbl_diff_qty.config(text="0.00 (Fərq yoxdur)", fg=TEXT_MUTED)
            
        if qiymet_ferqi > 0:
            lbl_diff_val.config(text=f"+{qiymet_ferqi:.2f} AZN", fg=SUCCESS_COLOR)
        elif qiymet_ferqi < 0:
            lbl_diff_val.config(text=f"{qiymet_ferqi:.2f} AZN", fg=DANGER_COLOR)
        else:
            lbl_diff_val.config(text="0.00 AZN", fg=TEXT_MUTED)
            
        # Save directly to Excel in-memory sheet with formulas
        row = p['row']
        sheet.cell(row=row, column=yeni_sayim_col, value=total_qty)
        sheet.cell(row=row, column=say_ferqi_col, value=f"={yeni_letter}{row}-{qaliq_letter}{row}")
        sheet.cell(row=row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{row}*{qiymet_letter}{row}")
        if operator_col:
            sheet.cell(row=row, column=operator_col, value=operators_list)
    
    # Update Stats
    update_stats()
    trigger_autosave()
    play_pc_beep(1000, 150)
    
    # Reset input fields and focus barcode
    barcode_var.set("")
    qty_var.set("")
    ent_qty.config(state="disabled")
    
    # Focus back to barcode
    ent_barcode.focus_set()
    display_name = f"{p['brend']} - {p['adi']}" if p.get('adi') else p['brend']
    status_var.set(f"Yadda saxlanıldı: '{display_name}' - Say: {qty:.2f}")

def update_stats():
    if is_network_mode:
        try:
            url = f"http://{server_ip}:{server_port}/get_products?operator={urllib.parse.quote(operator_name)}"
            req = urllib.request.urlopen(url, timeout=3)
            res = json.loads(req.read().decode("utf-8"))
            global products
            products = res["products"]
            s = res["stats"]
            
            lbl_stat_counted.config(text=f"{s['total_counted']} / {s['total_items']}")
            lbl_stat_surplus.config(text=f"+{s['surplus_qty']:.2f}")
            lbl_stat_shortage.config(text=f"-{s['shortage_qty']:.2f}")
            
            total_diff_val = s['total_diff_val']
            if total_diff_val > 0:
                lbl_stat_diff_val.config(text=f"+{total_diff_val:.2f} AZN", fg=SUCCESS_COLOR)
            elif total_diff_val < 0:
                lbl_stat_diff_val.config(text=f"{total_diff_val:.2f} AZN", fg=DANGER_COLOR)
            else:
                lbl_stat_diff_val.config(text="0.00 AZN", fg=TEXT_MUTED)
                
            refresh_counted_products()
            refresh_brand_products()
        except Exception as e:
            status_var.set(f"Bağlantı xətası: {str(e)}")
        return

    with db_lock:
        total_items = len(products)
        total_counted = 0
        surplus_qty = 0.0
        shortage_qty = 0.0
        total_diff_val = 0.0
        
        for bc, p in products.items():
            if p['yeni'] is not None:
                total_counted += 1
                diff_q = p['yeni'] - p['qaliq']
                diff_v = diff_q * p['qiymet']
                
                if diff_q > 0:
                    surplus_qty += diff_q
                elif diff_q < 0:
                    shortage_qty += abs(diff_q)
                    
                total_diff_val += diff_v
            
    lbl_stat_counted.config(text=f"{total_counted} / {total_items}")
    lbl_stat_surplus.config(text=f"+{surplus_qty:.2f}")
    lbl_stat_shortage.config(text=f"-{shortage_qty:.2f}")
    
    if total_diff_val > 0:
        lbl_stat_diff_val.config(text=f"+{total_diff_val:.2f} AZN", fg=SUCCESS_COLOR)
    elif total_diff_val < 0:
        lbl_stat_diff_val.config(text=f"{total_diff_val:.2f} AZN", fg=DANGER_COLOR)
    else:
        lbl_stat_diff_val.config(text="0.00 AZN", fg=TEXT_MUTED)
        
    refresh_counted_products()
    refresh_brand_products()

    # Update active operators list and server status label if in server mode
    if 'lbl_server_status' in globals() and lbl_server_status and is_server_active:
        update_server_status_ui()

def add_new_product():
    dialog = tk.Toplevel(root)
    dialog.title("Yeni Məhsul Əlavə Et")
    dialog.geometry("420x400")
    dialog.configure(bg=CARD_BG)
    dialog.resizable(False, False)
    dialog.grab_set()
    
    # Center
    dialog.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() - dialog.winfo_width()) // 2
    y = root.winfo_y() + (root.winfo_height() - dialog.winfo_height()) // 2
    dialog.geometry(f"+{x}+{y}")
    
    title_lbl = tk.Label(
        dialog, 
        text="Yeni Məhsul Qeydiyyatı", 
        font=('Segoe UI', 13, 'bold'), 
        bg=CARD_BG, 
        fg=ACCENT_COLOR
    )
    title_lbl.pack(pady=15)
    
    form_frame = tk.Frame(dialog, bg=CARD_BG)
    form_frame.pack(padx=20, fill="both", expand=True)
    form_frame.columnconfigure(1, weight=1)
    
    # Fields
    tk.Label(form_frame, text="Barkod:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=0, column=0, sticky="w", pady=8)
    ent_new_barcode = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_new_barcode.grid(row=0, column=1, sticky="we", pady=8, padx=5)
    
    # Autofill current typed barcode if any
    curr_bc = barcode_var.get().strip()
    if curr_bc:
        ent_new_barcode.insert(0, curr_bc)
        
    tk.Label(form_frame, text="Kod:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=1, column=0, sticky="w", pady=8)
    ent_new_kod = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_new_kod.grid(row=1, column=1, sticky="we", pady=8, padx=5)
        
    tk.Label(form_frame, text="Brend / Adı:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=2, column=0, sticky="w", pady=8)
    ent_new_name = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_new_name.grid(row=2, column=1, sticky="we", pady=8, padx=5)
    ent_new_name.focus_set()
    
    tk.Label(form_frame, text="Qiymət (AZN):", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=3, column=0, sticky="w", pady=8)
    ent_new_price = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_new_price.grid(row=3, column=1, sticky="we", pady=8, padx=5)
    ent_new_price.insert(0, "0.00")
    
    tk.Label(form_frame, text="Sistem Qalığı:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10)).grid(row=4, column=0, sticky="w", pady=8)
    ent_new_stock = tk.Entry(form_frame, bg=ENTRY_BG, fg=TEXT_COLOR, insertbackground=TEXT_COLOR, font=('Segoe UI', 11), bd=0, relief="flat", highlightthickness=1, highlightbackground=BORDER_COLOR, highlightcolor=ACCENT_COLOR)
    ent_new_stock.grid(row=4, column=1, sticky="we", pady=8, padx=5)
    ent_new_stock.insert(0, "0.00")
    
    def on_submit():
        bc = ent_new_barcode.get().strip()
        kod_val = ent_new_kod.get().strip()
        name = ent_new_name.get().strip()
        price_str = ent_new_price.get().strip()
        stock_str = ent_new_stock.get().strip()
        
        if not bc or not name:
            messagebox.showerror("Xəta", "Barkod və Brend/Ad sahələri boş qala bilməz!", parent=dialog)
            return
            
        if bc in products:
            messagebox.showerror("Xəta", "Bu barkod artıq sistemdə mövcuddur!", parent=dialog)
            return
            
        try:
            price = float(price_str)
            stock = float(stock_str)
            if price < 0 or stock < 0:
                raise ValueError()
        except ValueError:
            messagebox.showerror("Xəta", "Qiymət və Qalıq üçün düzgün müsbət ədəd daxil edin!", parent=dialog)
            return
            
        if is_network_mode:
            try:
                encoded_name = urllib.parse.quote(name)
                encoded_kod = urllib.parse.quote(kod_val)
                url = f"http://{server_ip}:{server_port}/add_count?barcode={bc}&qty=0&mode=set&name={encoded_name}&price={price}&stock={stock}&kod={encoded_kod}&operator={urllib.parse.quote(operator_name)}"
                req = urllib.request.urlopen(url, timeout=3)
                res = json.loads(req.read().decode("utf-8"))
                if res.get("status") == "success":
                    dialog.destroy()
                    update_stats()
                    barcode_var.set(bc)
                    on_barcode_enter()
                else:
                    messagebox.showerror("Xəta", f"Server xətası: {res.get('message')}", parent=dialog)
            except Exception as e:
                messagebox.showerror("Xəta", f"Serverlə bağlantı kəsildi:\n{str(e)}", parent=dialog)
            return

        with db_lock:
            # Append to Excel file
            new_row = sheet.max_row + 1
            sheet.cell(row=new_row, column=barkod_col, value=bc)
            if kod_col:
                sheet.cell(row=new_row, column=kod_col, value=kod_val)
            sheet.cell(row=new_row, column=brend_col, value=name)
            if adi_col:
                sheet.cell(row=new_row, column=adi_col, value="")
            sheet.cell(row=new_row, column=anbar_qaligi_col, value=stock)
            sheet.cell(row=new_row, column=qiymet_col, value=price)
            
            # Add formulas for calculation columns
            sheet.cell(row=new_row, column=yeni_sayim_col, value=None)
            sheet.cell(row=new_row, column=say_ferqi_col, value=f"={yeni_letter}{new_row}-{qaliq_letter}{new_row}")
            sheet.cell(row=new_row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{new_row}*{qiymet_letter}{new_row}")
            if operator_col:
                sheet.cell(row=new_row, column=operator_col, value=operator_name)
            
            # Add to memory dict
            products[bc] = {
                'row': new_row,
                'kod': kod_val,
                'brend': name,
                'adi': '',
                'qaliq': stock,
                'qiymet': price,
                'yeni': None,
                'order': 0,
                'operator': operator_name
            }
        
        trigger_autosave()
        dialog.destroy()
        
        # Load the added barcode
        barcode_var.set(bc)
        on_barcode_enter()
        
    btn_submit = tk.Button(
        dialog, 
        text="Əlavə Et", 
        command=on_submit, 
        bg=SUCCESS_COLOR, 
        fg=BG_COLOR, 
        font=('Segoe UI', 10, 'bold'), 
        relief="flat", 
        bd=0,
        padx=20, 
        pady=6,
        cursor="hand2"
    )
    btn_submit.pack(pady=15)
    make_hoverable(btn_submit, "#059669")

def save_and_close():
    if is_network_mode:
        if messagebox.askyesno("Çıxış", "Şəbəkə bağlantısını kəsib çıxmaq istəyirsiniz?"):
            root.destroy()
        return

    try:
        # Stop SSH Tunnel if active
        stop_ssh_tunnel()
        
        # Save main file under lock
        with db_lock:
            wb.save(filepath)
            
        # Copy and style outside lock
        import shutil
        dir_name = os.path.dirname(filepath)
        new_file_path = os.path.join(dir_name, "yeni sayim neticesi.xlsx")
        shutil.copy2(filepath, new_file_path)
        save_output_excel(new_file_path)
        
        messagebox.showinfo(
            "Uğurlu", 
            f"Bütün sayım nəticələri yadda saxlanıldı!\n\n"
            f"1. Əsas fayl yeniləndi: {os.path.basename(filepath)}\n"
            f"2. Yeni fayl yaradıldı: yeni sayim neticesi.xlsx"
        )
        root.destroy()
    except Exception as e:
        messagebox.showerror(
            "Yadda Saxlama Xətası", 
            f"Faylları yadda saxlamaq mümkün olmadı!\n\n"
            f"Ehtimal olunan səbəb: Excel fayllarından biri hazırda başqa bir proqramda (məs. MS Excel) açıqdır.\n"
            f"Zəhmət olmasa həmin proqramı bağlayıb yenidən yoxlayın.\n\nXəta təfərrüatı: {str(e)}"
        )

# --- HEADER SECTION ---
header_frame = tk.Frame(root, bg=BG_COLOR)
header_frame.pack(fill="x", padx=30, pady=(20, 10))

# Top row inside header_frame
header_top = tk.Frame(header_frame, bg=BG_COLOR)
header_top.pack(fill="x")

lbl_title = tk.Label(
    header_top, 
    text="ANBAR SAYIMI YARDIMÇISI", 
    font=('Segoe UI', 16, 'bold'), 
    bg=BG_COLOR, 
    fg=TEXT_COLOR
)
lbl_title.pack(side="left")

lbl_filename = tk.Label(
    header_top, 
    text=f"Aktiv Excel: {os.path.basename(filepath)}" if not is_network_mode else f"Şəbəkə Rejimi (Müştəri): {operator_name}", 
    font=('Segoe UI', 10, 'italic'), 
    bg=BG_COLOR, 
    fg=TEXT_MUTED
)
lbl_filename.pack(side="right", ipady=3)

# Bottom row inside header_frame for controls/status
header_bottom = tk.Frame(header_frame, bg=BG_COLOR)
header_bottom.pack(fill="x", pady=(8, 0))

btn_brand_select = tk.Button(
    header_bottom, 
    text="🔍 Brend Seç: Hamısı", 
    command=select_brand_gui, 
    bg=ACCENT_COLOR, 
    fg=TEXT_COLOR, 
    activebackground=ACCENT_HOVER,
    activeforeground=TEXT_COLOR,
    font=('Segoe UI', 9, 'bold'), 
    relief="flat", 
    bd=0,
    padx=12, 
    pady=5,
    cursor="hand2"
)
btn_brand_select.pack(side="left")
make_hoverable(btn_brand_select, ACCENT_HOVER)

# Reset Count Button
btn_clear_list = tk.Button(
    header_bottom,
    text="🗑️ Sayımı Sıfırla",
    command=clear_inventory_counts_gui,
    bg=DANGER_COLOR,
    fg=TEXT_COLOR,
    activebackground="#B91C1C",
    activeforeground=TEXT_COLOR,
    font=('Segoe UI', 9, 'bold'),
    relief="flat",
    bd=0,
    padx=12,
    pady=5,
    cursor="hand2"
)
btn_clear_list.pack(side="left", padx=15)
make_hoverable(btn_clear_list, "#B91C1C")

# Network status and start button frame
network_controls_frame = tk.Frame(header_bottom, bg=BG_COLOR)
network_controls_frame.pack(side="right")

if is_network_mode:
    lbl_client_status = tk.Label(
        network_controls_frame, 
        text=f"🟢 Mərkəzə qoşulub: {server_ip}:{server_port} | Operator: {operator_name}", 
        font=('Segoe UI', 9, 'bold'), 
        bg=BG_COLOR, 
        fg=SUCCESS_COLOR
    )
    lbl_client_status.pack(side="right", padx=5)
else:
    lbl_server_status = tk.Label(
        network_controls_frame, 
        text="🔴 Server qapalıdır", 
        font=('Segoe UI', 9, 'bold'), 
        bg=BG_COLOR, 
        fg=DANGER_COLOR
    )
    lbl_server_status.pack(side="left", padx=10)

    btn_start_server = tk.Button(
        network_controls_frame,
        text="🌐 Şəbəkəni Başlat",
        command=toggle_server,
        bg=ACCENT_COLOR,
        fg=TEXT_COLOR,
        activebackground=ACCENT_HOVER,
        activeforeground=TEXT_COLOR,
        font=('Segoe UI', 9, 'bold'),
        relief="flat",
        bd=0,
        padx=12,
        pady=4,
        cursor="hand2"
    )
    btn_start_server.pack(side="left", padx=5)
    make_hoverable(btn_start_server, ACCENT_HOVER)

    btn_show_qr = tk.Button(
        network_controls_frame,
        text="📱 QR Kod",
        command=show_qr_code_window,
        bg=SUCCESS_COLOR,
        fg=BG_COLOR,
        activebackground="#059669",
        activeforeground=BG_COLOR,
        font=('Segoe UI', 9, 'bold'),
        relief="flat",
        bd=0,
        padx=12,
        pady=4,
        cursor="hand2"
    )
    make_hoverable(btn_show_qr, "#059669")

# Separator line
sep = tk.Frame(root, height=1, bg=BORDER_COLOR)
sep.pack(fill="x", padx=30, pady=5)

# --- MAIN CONTENT PANEL ---
main_panel = tk.Frame(root, bg=BG_COLOR)
main_panel.pack(fill="both", expand=True, padx=30, pady=10)

# Grid Layout config
main_panel.columnconfigure(0, weight=3, uniform="group1") # Left input
main_panel.columnconfigure(1, weight=5, uniform="group1") # Right history/stats
main_panel.rowconfigure(0, weight=1)

# --- LEFT PANEL (INPUT & INTERACTION) ---
left_panel = tk.Frame(main_panel, bg=BG_COLOR)
left_panel.grid(row=0, column=0, sticky="nswe", padx=(0, 15))

# Card 1: Scanning Inputs
scan_card = tk.Frame(
    left_panel, 
    bg=CARD_BG, 
    highlightthickness=1, 
    highlightbackground=BORDER_COLOR
)
scan_card.pack(fill="both", expand=True, ipady=10)

# Card Title
lbl_scan_title = tk.Label(
    scan_card,
    text="📥 SKAN VƏ DAXİLETMƏ",
    font=('Segoe UI', 11, 'bold'),
    bg=CARD_BG,
    fg=ACCENT_COLOR
)
lbl_scan_title.pack(anchor="w", padx=20, pady=(15, 5))

# Barcode Label & Entry
lbl_bc_prompt = tk.Label(scan_card, text="Barkodu skan edin və ya yazın (Enter sıxın):", font=('Segoe UI', 10), bg=CARD_BG, fg=TEXT_COLOR)
lbl_bc_prompt.pack(anchor="w", padx=20, pady=(15, 5))

ent_barcode = tk.Entry(
    scan_card, 
    textvariable=barcode_var, 
    bg=ENTRY_BG, 
    disabledbackground=CARD_BG,
    disabledforeground=TEXT_MUTED,
    fg=TEXT_COLOR, 
    insertbackground=TEXT_COLOR, 
    font=('Segoe UI', 13, 'bold'), 
    bd=0, 
    relief="flat", 
    highlightthickness=1, 
    highlightbackground=BORDER_COLOR, 
    highlightcolor=ACCENT_COLOR
)
ent_barcode.pack(fill="x", padx=20, pady=5, ipady=8)
ent_barcode.bind("<Return>", on_barcode_enter)
ent_barcode.focus_set()

# Product Details Display
details_frame = tk.Frame(scan_card, bg=CARD_BG)
details_frame.pack(fill="x", padx=20, pady=15)
details_frame.columnconfigure(0, weight=1)
details_frame.columnconfigure(1, weight=1)

# Product Name (Spans across both columns)
lbl_product_name = tk.Label(
    details_frame, 
    text="Skan gözlənilir...", 
    font=('Segoe UI', 13, 'bold'), 
    bg=CARD_BG, 
    fg=TEXT_MUTED, 
    wraplength=380, 
    justify="left"
)
lbl_product_name.grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 15))

# System Stock Info
tk.Label(details_frame, text="Sistem Qalığı:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 9)).grid(row=1, column=0, sticky="w", pady=3)
lbl_system_stock = tk.Label(details_frame, text="-", bg=CARD_BG, fg=TEXT_COLOR, font=('Segoe UI', 10, 'bold'))
lbl_system_stock.grid(row=1, column=1, sticky="w", pady=3)

# Price Info
tk.Label(details_frame, text="Məhsulun Qiyməti:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 9)).grid(row=2, column=0, sticky="w", pady=3)
lbl_price = tk.Label(details_frame, text="-", bg=CARD_BG, fg=TEXT_COLOR, font=('Segoe UI', 10, 'bold'))
lbl_price.grid(row=2, column=1, sticky="w", pady=3)

# Previously Counted Info
tk.Label(details_frame, text="Əvvəlki Sayım:", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 9)).grid(row=3, column=0, sticky="w", pady=3)
lbl_prev_count = tk.Label(details_frame, text="-", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 10, 'bold'))
lbl_prev_count.grid(row=3, column=1, sticky="w", pady=3)

# Real Count input field
lbl_qty_prompt = tk.Label(scan_card, text="Real sayı daxil edin (Enter sıxın):", font=('Segoe UI', 10), bg=CARD_BG, fg=TEXT_COLOR)
lbl_qty_prompt.pack(anchor="w", padx=20, pady=(10, 5))

ent_qty = tk.Entry(
    scan_card, 
    textvariable=qty_var, 
    bg=ENTRY_BG, 
    disabledbackground=CARD_BG,
    disabledforeground=TEXT_MUTED,
    fg=TEXT_COLOR, 
    insertbackground=TEXT_COLOR, 
    font=('Segoe UI', 13, 'bold'), 
    bd=0, 
    relief="flat", 
    highlightthickness=1, 
    highlightbackground=BORDER_COLOR, 
    highlightcolor=ACCENT_COLOR,
    state="disabled"
)
ent_qty.pack(fill="x", padx=20, pady=5, ipady=8)
ent_qty.bind("<Return>", on_qty_enter)

# Difference Display (Immediate visual feedback)
diff_frame = tk.Frame(scan_card, bg=CARD_BG)
diff_frame.pack(fill="x", padx=20, pady=(20, 10))
diff_frame.columnconfigure(0, weight=1)
diff_frame.columnconfigure(1, weight=1)

tk.Label(diff_frame, text="SAY FƏRQİ", font=('Segoe UI', 9, 'bold'), bg=CARD_BG, fg=TEXT_MUTED).grid(row=0, column=0, sticky="w")
lbl_diff_qty = tk.Label(diff_frame, text="-", font=('Segoe UI', 14, 'bold'), bg=CARD_BG, fg=TEXT_MUTED)
lbl_diff_qty.grid(row=1, column=0, sticky="w", pady=5)

tk.Label(diff_frame, text="QİYMƏT FƏRQİ", font=('Segoe UI', 9, 'bold'), bg=CARD_BG, fg=TEXT_MUTED).grid(row=0, column=1, sticky="w")
lbl_diff_val = tk.Label(diff_frame, text="-", font=('Segoe UI', 14, 'bold'), bg=CARD_BG, fg=TEXT_MUTED)
lbl_diff_val.grid(row=1, column=1, sticky="w", pady=5)


# --- RIGHT PANEL (STATS & SCAN HISTORY) ---
right_panel = tk.Frame(main_panel, bg=BG_COLOR)
right_panel.grid(row=0, column=1, sticky="nswe", padx=(15, 0))

# Stats Grid (Top of right panel)
stats_frame = tk.Frame(right_panel, bg=BG_COLOR)
stats_frame.pack(fill="x", pady=(0, 15))
stats_frame.columnconfigure(0, weight=1)
stats_frame.columnconfigure(1, weight=1)
stats_frame.columnconfigure(2, weight=1)
stats_frame.columnconfigure(3, weight=1)

def create_stat_card(parent, col, title, initial_val, color=TEXT_COLOR):
    frame = tk.Frame(parent, bg=CARD_BG, highlightthickness=1, highlightbackground=BORDER_COLOR)
    frame.grid(row=0, column=col, sticky="nswe", padx=6, ipady=10)
    
    lbl_title = tk.Label(frame, text=title.upper(), font=('Segoe UI', 8, 'bold'), bg=CARD_BG, fg=TEXT_MUTED)
    lbl_title.pack(pady=(8, 2))
    
    lbl_val = tk.Label(frame, text=initial_val, font=('Segoe UI', 14, 'bold'), bg=CARD_BG, fg=color)
    lbl_val.pack(pady=(2, 8))
    return lbl_val

lbl_stat_counted = create_stat_card(stats_frame, 0, "SAYILAN", "0 / 0", ACCENT_COLOR)
lbl_stat_surplus = create_stat_card(stats_frame, 1, "ARTIQ (QTY)", "0.00", SUCCESS_COLOR)
lbl_stat_shortage = create_stat_card(stats_frame, 2, "ƏSKİK (QTY)", "0.00", DANGER_COLOR)
lbl_stat_diff_val = create_stat_card(stats_frame, 3, "QIYMƏT FƏRQİ", "0.00 AZN", TEXT_MUTED)

# Tab Control (Notebook) for History and Brand catalog
notebook = ttk.Notebook(right_panel)
notebook.pack(fill="both", expand=True)

# Tab 1: Sayılan Məhsullar
tab_history = ttk.Frame(notebook)
notebook.add(tab_history, text=" Sayılan Məhsullar ")

# Tab 2: Brendin Məhsulları
tab_brand_products = ttk.Frame(notebook)
notebook.add(tab_brand_products, text=" Brendin Məhsulları ")

# History Search Area
history_search_frame = tk.Frame(tab_history, bg=BG_COLOR)
history_search_frame.pack(fill="x", padx=5, pady=(5, 0))

lbl_hist_search = tk.Label(history_search_frame, text="🔍 Axtarış:", font=('Segoe UI', 9), bg=BG_COLOR, fg=TEXT_MUTED)
lbl_hist_search.pack(side="left", padx=5)

ent_hist_search = tk.Entry(
    history_search_frame,
    textvariable=history_search_var,
    bg=ENTRY_BG,
    fg=TEXT_COLOR,
    insertbackground=TEXT_COLOR,
    font=('Segoe UI', 9),
    bd=0,
    relief="flat",
    highlightthickness=1,
    highlightbackground=BORDER_COLOR,
    highlightcolor=ACCENT_COLOR
)
ent_hist_search.pack(side="left", fill="x", expand=True, ipady=3, padx=5)
ent_hist_search.bind("<KeyRelease>", lambda event: refresh_counted_products())

# History Treeview Setup (inside tab_history)
columns = ("kod", "barcode", "brand", "stock", "count", "diff_qty", "diff_val", "operator")
history_tree = ttk.Treeview(tab_history, columns=columns, show="headings", selectmode="browse")

history_tree.heading("kod", text="Kod")
history_tree.heading("barcode", text="Barkod")
history_tree.heading("brand", text="Brend / Adı")
history_tree.heading("stock", text="Sist.")
history_tree.heading("count", text="Real")
history_tree.heading("diff_qty", text="Fərq")
history_tree.heading("diff_val", text="Dəyər Fərqi")
history_tree.heading("operator", text="Operator")

history_tree.column("kod", width=60, anchor="center")
history_tree.column("barcode", width=85, anchor="center")
history_tree.column("brand", width=140, anchor="w")
history_tree.column("stock", width=40, anchor="center")
history_tree.column("count", width=40, anchor="center")
history_tree.column("diff_qty", width=85, anchor="center")
history_tree.column("diff_val", width=85, anchor="center")
history_tree.column("operator", width=95, anchor="center")

history_tree.pack(fill="both", expand=True, padx=5, pady=5)
history_tree.bind("<Double-Button-1>", on_history_double_click)
history_tree.tag_configure('surplus', foreground=TEXT_COLOR)
history_tree.tag_configure('shortage', foreground=TEXT_COLOR)
history_tree.tag_configure('equal', foreground=TEXT_COLOR)

# Brand Search Area
brand_search_frame = tk.Frame(tab_brand_products, bg=BG_COLOR)
brand_search_frame.pack(fill="x", padx=5, pady=(5, 0))

lbl_brand_search = tk.Label(brand_search_frame, text="🔍 Axtarış:", font=('Segoe UI', 9), bg=BG_COLOR, fg=TEXT_MUTED)
lbl_brand_search.pack(side="left", padx=5)

ent_brand_search = tk.Entry(
    brand_search_frame,
    textvariable=brand_search_var,
    bg=ENTRY_BG,
    fg=TEXT_COLOR,
    insertbackground=TEXT_COLOR,
    font=('Segoe UI', 9),
    bd=0,
    relief="flat",
    highlightthickness=1,
    highlightbackground=BORDER_COLOR,
    highlightcolor=ACCENT_COLOR
)
ent_brand_search.pack(side="left", fill="x", expand=True, ipady=3, padx=5)
ent_brand_search.bind("<KeyRelease>", lambda event: refresh_brand_products())

# Brand Products Treeview Setup (inside tab_brand_products)
brand_columns = ("kod", "barcode", "brand", "stock", "count", "price", "operator")
brand_tree = ttk.Treeview(tab_brand_products, columns=brand_columns, show="headings", selectmode="browse")

brand_tree.heading("kod", text="Kod")
brand_tree.heading("barcode", text="Barkod")
brand_tree.heading("brand", text="Brend / Adı")
brand_tree.heading("stock", text="Sist.")
brand_tree.heading("count", text="Real")
brand_tree.heading("price", text="Qiymət")
brand_tree.heading("operator", text="Operator")

brand_tree.column("kod", width=60, anchor="center")
brand_tree.column("barcode", width=85, anchor="center")
brand_tree.column("brand", width=170, anchor="w")
brand_tree.column("stock", width=40, anchor="center")
brand_tree.column("count", width=40, anchor="center")
brand_tree.column("price", width=65, anchor="center")
brand_tree.column("operator", width=95, anchor="center")

brand_tree.pack(fill="both", expand=True, padx=5, pady=5)
brand_tree.bind("<Double-Button-1>", on_brand_product_double_click)
brand_tree.tag_configure('surplus', foreground=TEXT_COLOR)
brand_tree.tag_configure('shortage', foreground=TEXT_COLOR)
brand_tree.tag_configure('equal', foreground=TEXT_COLOR)

# Run initial stats calculation
update_stats()


# --- FOOTER / ACTIONS SECTION ---
footer_frame = tk.Frame(root, bg=BG_COLOR)
footer_frame.pack(fill="x", padx=30, pady=(10, 20))

# Status Bar
lbl_status = tk.Label(
    footer_frame, 
    textvariable=status_var, 
    font=('Segoe UI', 9, 'italic'), 
    bg=BG_COLOR, 
    fg=TEXT_MUTED
)
lbl_status.pack(side="left", ipady=8)

# Buttons on the right
btn_save = tk.Button(
    footer_frame, 
    text="💾 Yadda Saxla və Çıx" if not is_network_mode else "🚪 Bağlantını Kəs və Çıx", 
    command=save_and_close, 
    bg=SUCCESS_COLOR if not is_network_mode else DANGER_COLOR, 
    fg=BG_COLOR if not is_network_mode else TEXT_COLOR, 
    activebackground="#059669" if not is_network_mode else "#DC2626", 
    activeforeground=BG_COLOR if not is_network_mode else TEXT_COLOR,
    font=('Segoe UI', 10, 'bold'), 
    relief="flat", 
    bd=0, 
    padx=18, 
    pady=8,
    cursor="hand2"
)
btn_save.pack(side="right", padx=(10, 0))
make_hoverable(btn_save, "#059669" if not is_network_mode else "#DC2626")

btn_new_prod = tk.Button(
    footer_frame, 
    text="➕ Yeni Məhsul Əlavə Et", 
    command=add_new_product, 
    bg=ACCENT_COLOR, 
    fg=TEXT_COLOR, 
    activebackground=ACCENT_HOVER,
    activeforeground=TEXT_COLOR,
    font=('Segoe UI', 10, 'bold'), 
    relief="flat", 
    bd=0, 
    padx=15, 
    pady=8,
    cursor="hand2"
)
btn_new_prod.pack(side="right")
make_hoverable(btn_new_prod, ACCENT_HOVER)


# Bind focus to barcode entry on click inside window
def reset_focus(event):
    if event.widget == root or event.widget == main_panel:
        ent_barcode.focus_set()

root.bind("<Button-1>", reset_focus)

# Start Main Loop (moved to end of file)

MOBILE_HTML_CONTENT = """<!DOCTYPE html>
<html lang="az">
<head>
    <meta charset="UTF-8">
    <meta name="google" content="notranslate">
    <meta name="viewport" content="width=device-width, initial-scale=1.0, maximum-scale=1.0, user-scalable=no">
    <title>Mobil Anbar Sayımı</title>
    <link rel="manifest" href="/manifest.json">
    <meta name="apple-mobile-web-app-capable" content="yes">
    <meta name="apple-mobile-web-app-status-bar-style" content="black-translucent">
    <meta name="apple-mobile-web-app-title" content="Sayım">
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Outfit:wght@600;700;800&display=swap&subset=latin,latin-ext" rel="stylesheet">
    <style>
        :root {
            --bg-color: #0F172A;
            --card-bg: #1E293B;
            --input-bg: #334155;
            --text-color: #F8FAFC;
            --text-muted: #94A3B8;
            --accent: #6366F1;
            --accent-hover: #4F46E5;
            --success: #10B981;
            --danger: #EF4444;
            --border: #475569;
        }
        * {
            box-sizing: border-box;
            margin: 0;
            padding: 0;
            font-family: 'Inter', sans-serif;
            -webkit-tap-highlight-color: transparent;
        }
        body {
            background-color: var(--bg-color);
            color: var(--text-color);
            padding: 15px;
            display: flex;
            flex-direction: column;
            align-items: center;
            min-height: 100vh;
        }
        .container {
            width: 100%;
            max-width: 480px;
            display: flex;
            flex-direction: column;
            gap: 15px;
        }
        header {
            text-align: center;
            padding: 10px 0;
        }
        header h1 {
            font-family: 'Outfit', sans-serif;
            font-size: 1.6rem;
            font-weight: 800;
            color: var(--accent);
            letter-spacing: 0.5px;
        }
        header p {
            font-size: 0.85rem;
            color: var(--text-muted);
            margin-top: 3px;
        }
        .card {
            background-color: var(--card-bg);
            border-radius: 12px;
            padding: 20px;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1), 0 2px 4px -1px rgba(0, 0, 0, 0.06);
            border: 1px solid var(--border);
        }
        .card-title {
            font-size: 0.85rem;
            font-weight: 600;
            margin-bottom: 15px;
            color: var(--accent);
            text-transform: uppercase;
            letter-spacing: 1px;
        }
        .form-group {
            display: flex;
            flex-direction: column;
            gap: 6px;
            margin-bottom: 12px;
        }
        label {
            font-size: 0.85rem;
            color: var(--text-muted);
            font-weight: 500;
        }
        input, select {
            background-color: var(--input-bg);
            border: 1px solid var(--border);
            color: var(--text-color);
            padding: 12px;
            border-radius: 8px;
            font-size: 1rem;
            width: 100%;
            outline: none;
            transition: border-color 0.2s;
        }
        input:focus, select:focus {
            border-color: var(--accent);
        }
        .btn {
            background-color: var(--accent);
            color: white;
            border: none;
            padding: 12px;
            font-size: 1rem;
            font-weight: 600;
            border-radius: 8px;
            cursor: pointer;
            width: 100%;
            transition: background-color 0.2s, transform 0.1s;
            display: flex;
            align-items: center;
            justify-content: center;
            gap: 8px;
        }
        .btn:active {
            transform: scale(0.98);
        }
        .btn-success {
            background-color: var(--success);
        }
        .btn-success:hover {
            background-color: #059669;
        }
        .btn-danger {
            background-color: var(--danger);
        }
        .btn-secondary {
            background-color: #475569;
        }
        .row {
            display: flex;
            gap: 10px;
        }
        .info-row {
            display: flex;
            justify-content: space-between;
            padding: 10px 0;
            border-bottom: 1px dashed var(--border);
            font-size: 0.95rem;
        }
        .info-row:last-child {
            border-bottom: none;
        }
        .info-label {
            color: var(--text-muted);
        }
        .info-val {
            font-weight: 600;
        }
        #scanner-container {
            width: 100%;
            border-radius: 12px;
            overflow: hidden;
            background: #000;
            display: none;
            position: relative;
            aspect-ratio: 4/3;
            border: 2px solid var(--border);
            box-shadow: 0 8px 24px rgba(0, 0, 0, 0.4);
            margin-bottom: 15px;
        }
        #reader {
            border: none !important;
        }
        #reader a, #reader button, #reader select, #reader span, #reader img {
            display: none !important;
        }
        @keyframes scan-laser {
            0% { top: 0%; opacity: 0; }
            10% { opacity: 0.8; }
            90% { opacity: 0.8; }
            100% { top: 100%; opacity: 0; }
        }
        .scanner-laser {
            position: absolute;
            left: 0;
            right: 0;
            height: 3px;
            background: linear-gradient(90deg, transparent, var(--success), transparent);
            animation: scan-laser 2.2s infinite linear;
            z-index: 4;
            pointer-events: none;
            box-shadow: 0 0 10px var(--success), 0 0 20px var(--success);
        }
        .scanner-vignette {
            position: absolute;
            top: 0;
            left: 0;
            width: 100%;
            height: 100%;
            box-shadow: inset 0 0 40px rgba(0,0,0,0.6);
            z-index: 3;
            pointer-events: none;
        }
        .scanner-actions {
            position: absolute;
            bottom: 15px;
            left: 10px;
            right: 10px;
            display: flex;
            gap: 10px;
            justify-content: center;
            z-index: 10;
        }
        .scanner-action-btn {
            padding: 8px 16px;
            font-size: 0.9rem;
            border-radius: 20px;
            border: 1px solid rgba(255, 255, 255, 0.2);
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.3);
            font-weight: 600;
            backdrop-filter: blur(4px);
            transition: all 0.2s;
            outline: none;
            cursor: pointer;
            width: auto;
            min-width: 120px;
            display: flex;
            align-items: center;
            justify-content: center;
            gap: 6px;
        }
        #stop-scan-btn {
            background: rgba(239, 68, 68, 0.85) !important;
            color: #fff !important;
        }
        #stop-scan-btn:active {
            background: rgba(239, 68, 68, 1) !important;
            transform: scale(0.95);
        }
        #toggle-torch-btn {
            background: rgba(255, 255, 255, 0.15) !important;
            color: #fff !important;
            display: none;
        }
        #toggle-torch-btn:active {
            background: rgba(255, 255, 255, 0.3) !important;
            transform: scale(0.95);
        }
        .status-msg {
            padding: 10px;
            border-radius: 6px;
            font-size: 0.9rem;
            font-weight: 500;
            text-align: center;
            display: none;
            animation: fadeIn 0.3s;
        }
        .status-success {
            background-color: rgba(16, 185, 129, 0.15);
            color: var(--success);
            border: 1px solid rgba(16, 185, 129, 0.3);
        }
        .status-error {
            background-color: rgba(239, 68, 68, 0.15);
            color: var(--danger);
            border: 1px solid rgba(239, 68, 68, 0.3);
        }
        .history-item {
            display: flex;
            justify-content: space-between;
            padding: 12px;
            background: rgba(255, 255, 255, 0.03);
            border-radius: 8px;
            margin-bottom: 8px;
            border-left: 4px solid var(--accent);
            font-size: 0.9rem;
        }
        .history-name {
            font-weight: 600;
        }
        .history-details {
            font-size: 0.75rem;
            color: var(--text-muted);
            margin-top: 2px;
        }
        .history-qty {
            font-size: 1.1rem;
            font-weight: 700;
            color: var(--success);
            align-self: center;
        }
        @keyframes fadeIn {
            from { opacity: 0; transform: translateY(-5px); }
            to { opacity: 1; transform: translateY(0); }
        }
        .tab-btn {
            border: 1px solid var(--border);
            border-radius: 8px;
            padding: 10px;
            font-size: 0.9rem;
            transition: all 0.2s;
        }
        .tab-btn.active {
            background-color: var(--accent);
            color: white;
            border-color: var(--accent);
        }
        .product-list-item {
            transition: background-color 0.2s;
            display: flex;
            justify-content: space-between;
            align-items: center;
            padding: 12px;
            background: rgba(255,255,255,0.03);
            border-radius: 8px;
            border: 1px solid var(--border);
            margin-bottom: 8px;
            font-size: 0.9rem;
        }
        .product-list-item:active {
            background-color: rgba(255,255,255,0.08) !important;
        }
        .product-list-item-left {
            flex: 1;
            padding-right: 12px;
            text-align: left;
        }
        .product-list-item-right {
            text-align: right;
            min-width: 90px;
            display: flex;
            flex-direction: column;
            gap: 2px;
            align-items: flex-end;
            white-space: nowrap;
        }
        #autocomplete-results {
            display: none;
            position: absolute;
            top: 100%;
            left: 0;
            right: 0;
            background: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 8px;
            max-height: 200px;
            overflow-y: auto;
            z-index: 100;
            margin-top: 4px;
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.5);
        }
        .autocomplete-item {
            padding: 10px 12px;
            border-bottom: 1px solid var(--border);
            cursor: pointer;
            text-align: left;
            font-size: 0.9rem;
            transition: background-color 0.1s;
        }
        .autocomplete-item:last-child {
            border-bottom: none;
        }
        .autocomplete-item:hover {
            background-color: rgba(99, 102, 241, 0.15);
        }
        .autocomplete-item-title {
            font-weight: 600;
            color: var(--text-color);
        }
        .autocomplete-item-subtitle {
            font-size: 0.75rem;
            color: var(--text-muted);
            margin-top: 2px;
        }
    </style>
    <script src="https://unpkg.com/html5-qrcode/html5-qrcode.min.js"></script>
</head>
<body>
    <div class="container">
        <header>
            <h1>Mobil Anbar Sayımı</h1>
            <p>Bulud Serverinə qoşulub</p>
        </header>

        <!-- Sync Status Bar -->
        <div id="sync-status-bar" style="display: none; background-color: var(--danger); color: white; text-align: center; padding: 10px; font-size: 0.9rem; font-weight: bold; border-radius: 8px; border: 1px solid rgba(255,255,255,0.2); animation: fadeIn 0.3s; margin-bottom: 12px;">
            ⚠️ İnternet yoxdur! <span id="offline-count-lbl">0</span> sayım yaddaşda gözləyir.
        </div>

        <!-- Step 1: Operator Name Setup -->
        <div id="operator-setup-card" class="card">
            <div class="card-title">Operator Girişi</div>
            <div class="form-group">
                <label for="op-name-select">Operator seçin:</label>
                <select id="op-name-select">
                    <option value="">Yüklənir...</option>
                </select>
            </div>
            <div class="form-group" style="margin-top: 15px;">
                <label for="op-pin-input">PIN Kod:</label>
                <input type="password" id="op-pin-input" pattern="[0-9]*" inputmode="numeric" placeholder="PIN kodunuzu daxil edin">
            </div>
            <button class="btn btn-success" id="save-op-btn" style="margin-top: 15px;">Daxil ol</button>
        </div>

        <!-- Step 2: Main Count Interface -->
        <div id="main-count-interface" class="card" style="display: none;">
            <div class="card-title">Skan və Daxiletmə (<span id="op-display-name"></span>)</div>
            
            <!-- Brand Selection dropdown -->
            <div class="form-group" style="margin-bottom: 15px;">
                <label for="brand-select">Brend Seçin:</label>
                <select id="brand-select">
                    <option value="">[ Bütün Brendlər ]</option>
                </select>
            </div>
            
            <div class="status-msg" id="status-box"></div>

            <!-- Tabs -->
            <div class="row" style="margin-bottom: 15px;">
                <button class="btn tab-btn active" id="tab-scan-btn" style="flex: 1; border-radius: 8px 0 0 8px; padding: 10px 5px; font-size: 0.85rem;">Skan və Sayım</button>
                <button class="btn tab-btn" id="tab-list-btn" style="flex: 1; border-radius: 0 8px 8px 0; background-color: var(--card-bg); border: 1px solid var(--border); color: var(--text-muted); padding: 10px 5px; font-size: 0.85rem;">Məhsullar</button>
            </div>

            <!-- View 1: Scan & Count -->
            <div id="scan-view">
                <div class="form-group" style="margin-top: 10px; position: relative;">
                    <label for="barcode-input">Barkod:</label>
                    <div class="row">
                        <input type="text" id="barcode-input" placeholder="Barkodu yazın və ya skan edin" style="flex: 1;" autocomplete="off">
                        <button class="btn" id="search-btn" style="width: auto; padding: 0 15px;">Axtar</button>
                        <button class="btn" id="scan-camera-btn" style="width: auto; padding: 0 15px; background-color: var(--border);">📷</button>
                    </div>
                    <div id="autocomplete-results"></div>
                </div>

                <!-- Scanner Area moved to root level of body -->

                <!-- Product Details Display (Hidden by default) -->
                <div id="product-details-box" style="display: none; background: rgba(255,255,255,0.03); border-radius: 8px; padding: 12px; margin-bottom: 12px; border: 1px solid var(--border);">
                    <div class="info-row">
                        <span class="info-label">Məhsul:</span>
                        <span class="info-val" id="prod-name-lbl">-</span>
                    </div>
                    <div class="info-row">
                        <span class="info-label">Qiymət:</span>
                        <span class="info-val" id="prod-price-lbl">-</span>
                    </div>
                    <div class="info-row">
                        <span class="info-label">Sistem Qalığı:</span>
                        <span class="info-val" id="prod-stock-lbl">-</span>
                    </div>
                    <div class="info-row">
                        <span class="info-label">Əvvəlki Say:</span>
                        <span class="info-val" id="prod-prev-lbl" style="color: var(--success);">-</span>
                    </div>
                </div>

                <!-- Quantity Entry -->
                <div class="form-group" id="qty-group" style="display: none;">
                    <label for="qty-input">Yeni Say:</label>
                    <input type="number" step="any" id="qty-input" placeholder="Say daxil edin" inputmode="decimal">
                </div>

                <div class="row">
                    <button class="btn btn-secondary" id="cancel-btn" style="display: none; flex: 1;">Ləğv et</button>
                    <button class="btn btn-success" id="submit-btn" style="display: none; flex: 2;">Yadda Saxla</button>
                </div>
            </div>

            <!-- View 2: Product List -->
            <div id="list-view" style="display: none;">
                <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 10px;">
                    <span style="font-weight: 600; font-size: 0.9rem; color: var(--text-muted);">Məhsul siyahısı:</span>
                    <span style="background-color: var(--accent); color: white; padding: 2px 8px; border-radius: 20px; font-size: 0.75rem; font-weight: bold;"><span id="total-count-badge">0</span> məhsul</span>
                </div>
                <div class="form-group">
                    <input type="text" id="list-search-input" placeholder="Brend, ad və ya barkodla axtar...">
                </div>
                <div id="mobile-product-list" style="max-height: 420px; overflow-y: auto; display: flex; flex-direction: column; gap: 8px; margin-top: 10px;">
                    <div style="text-align: center; color: var(--text-muted); font-size: 0.85rem; padding: 10px;">Yüklənir...</div>
                </div>
            </div>
            
            <button class="btn btn-secondary" id="logout-btn" style="margin-top: 15px; font-size: 0.85rem; padding: 8px;">Adı Dəyişdir (Çıxış)</button>
        </div>

        <!-- History Card -->
        <div id="history-card" class="card" style="display: none;">
            <div class="card-title" style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 15px;">
                <span>Son Sayımlarım</span>
                <button id="clear-history-btn" style="background: rgba(239, 68, 68, 0.1); border: 1px solid rgba(239, 68, 68, 0.2); color: var(--danger); border-radius: 6px; padding: 4px 10px; font-size: 0.75rem; font-weight: 600; cursor: pointer; outline: none; text-transform: uppercase;">Təmizlə</button>
            </div>
            <div id="history-list">
                <div style="text-align: center; color: var(--text-muted); font-size: 0.85rem; padding: 10px;">Hələ ki sayım yoxdur</div>
            </div>
        </div>
    </div>

    <script>
        function safeGetStorage(key, fallback = "") {
            try { return localStorage.getItem(key) || fallback; } 
            catch(e) { return fallback; }
        }

        function safeSetStorage(key, value) {
            try { localStorage.setItem(key, value); } 
            catch(e) { console.error(e); }
        }

        function safeRemoveStorage(key) {
            try { localStorage.removeItem(key); } 
            catch(e) { console.error(e); }
        }

        function safeGetJSON(key, defaultValue = []) {
            try {
                const val = safeGetStorage(key);
                if (!val) return defaultValue;
                const parsed = JSON.parse(val);
                if (parsed === null) return defaultValue;
                if (Array.isArray(parsed)) {
                    return parsed.filter(x => x !== null && x !== undefined);
                }
                return parsed || defaultValue;
            } catch(e) {
                return defaultValue;
            }
        }

        let operatorName = safeGetStorage("operatorName");
        let currentProduct = null;
        let html5QrCode = null;
        let offlineQueue = safeGetJSON("offlineQueue").filter(item => item && item.barcode);
        let isSyncing = false;
        let selectedBrand = safeGetStorage("selectedBrand");
        let searchTimeout = null;
        let lastBarcodeValue = "";
        var allProducts = {};
        var isProductsLoading = false;

        // UI elements
        const syncStatusBar = document.getElementById("sync-status-bar");
        const offlineCountLbl = document.getElementById("offline-count-lbl");
        const opSetupCard = document.getElementById("operator-setup-card");
        const mainCountInterface = document.getElementById("main-count-interface");
        const historyCard = document.getElementById("history-card");
        
        const opNameSelect = document.getElementById("op-name-select");
        const opPinInput = document.getElementById("op-pin-input");
        const saveOpBtn = document.getElementById("save-op-btn");
        const opDisplayName = document.getElementById("op-display-name");
        const brandSelect = document.getElementById("brand-select");
        
        const barcodeInput = document.getElementById("barcode-input");
        const searchBtn = document.getElementById("search-btn");
        const scanCameraBtn = document.getElementById("scan-camera-btn");
        
        const tabScanBtn = document.getElementById("tab-scan-btn");
        const tabListBtn = document.getElementById("tab-list-btn");
        const scanView = document.getElementById("scan-view");
        const listView = document.getElementById("list-view");
        const listSearchInput = document.getElementById("list-search-input");
        const totalCountBadge = document.getElementById("total-count-badge");
        const scannerContainer = document.getElementById("scanner-container");
        const stopScanBtn = document.getElementById("stop-scan-btn");
        const switchCameraBtn = document.getElementById("switch-camera-btn");
        const autocompleteBox = document.getElementById("autocomplete-results");
        const toggleTorchBtn = document.getElementById("toggle-torch-btn");
        let isTorchOn = false;
        let scannerCameras = [];
        let currentCameraIndex = -1;
        
        const productDetailsBox = document.getElementById("product-details-box");
        const prodNameLbl = document.getElementById("prod-name-lbl");
        const prodPriceLbl = document.getElementById("prod-price-lbl");
        const prodStockLbl = document.getElementById("prod-stock-lbl");
        const prodPrevLbl = document.getElementById("prod-prev-lbl");
        
        const qtyGroup = document.getElementById("qty-group");
        const qtyInput = document.getElementById("qty-input");
        const cancelBtn = document.getElementById("cancel-btn");
        const submitBtn = document.getElementById("submit-btn");
        const statusBox = document.getElementById("status-box");
        const historyList = document.getElementById("history-list");
        const logoutBtn = document.getElementById("logout-btn");

        let globalAudioCtx = null;
        function initAudio() {
            if (!globalAudioCtx) {
                try {
                    globalAudioCtx = new (window.AudioContext || window.webkitAudioContext)();
                } catch(e) {}
            }
            if (globalAudioCtx && globalAudioCtx.state === 'suspended') {
                globalAudioCtx.resume();
            }
        }
        document.addEventListener('click', initAudio, { once: true });
        document.addEventListener('touchstart', initAudio, { once: true });
        document.addEventListener('keydown', initAudio, { once: true });

        function playBeep(freq = 1000, duration = 80, volume = 0.4) {
            try {
                initAudio();
                const ctx = globalAudioCtx || new (window.AudioContext || window.webkitAudioContext)();
                if (!ctx) return;
                
                const osc = ctx.createOscillator();
                const gain = ctx.createGain();
                
                osc.connect(gain);
                gain.connect(ctx.destination);
                
                osc.type = 'sine';
                osc.frequency.setValueAtTime(freq, ctx.currentTime);
                
                // Set volume and ramp down to prevent audio clicks/pops
                gain.gain.setValueAtTime(volume, ctx.currentTime);
                gain.gain.exponentialRampToValueAtTime(0.00001, ctx.currentTime + duration / 1000);
                
                osc.start(ctx.currentTime);
                osc.stop(ctx.currentTime + duration / 1000);
            } catch (e) {
                console.error("Audio feedback failed:", e);
            }
        }

        function updateOfflineBadge() {
            if (!syncStatusBar) return;
            const count = offlineQueue.length;
            if (count > 0) {
                syncStatusBar.style.display = "block";
                offlineCountLbl.textContent = count;
            } else {
                syncStatusBar.style.display = "none";
            }
        }

        function saveToOfflineQueue(barcode, qty, mode, display_name) {
            offlineQueue = offlineQueue.filter(item => item && item.barcode !== barcode);
            offlineQueue.push({
                barcode,
                qty,
                mode,
                display_name,
                operator: operatorName,
                time: new Date().toLocaleTimeString("az-AZ")
            });
            safeSetStorage("offlineQueue", JSON.stringify(offlineQueue));
            updateOfflineBadge();
        }

        function syncOfflineQueue() {
            if (isSyncing || offlineQueue.length === 0) return;
            if (!navigator.onLine) return;
            
            isSyncing = true;
            const item = offlineQueue[0];
            if (!item || !item.barcode) {
                offlineQueue.shift();
                safeSetStorage("offlineQueue", JSON.stringify(offlineQueue));
                isSyncing = false;
                syncOfflineQueue();
                return;
            }
            const url = `/add_count?barcode=${encodeURIComponent(item.barcode)}&qty=${item.qty}&mode=${item.mode}&operator=${encodeURIComponent(item.operator)}`;
            
            fetch(url)
                .then(res => res.json())
                .then(data => {
                    isSyncing = false;
                    if (data.status === "success") {
                        offlineQueue.shift();
                        safeSetStorage("offlineQueue", JSON.stringify(offlineQueue));
                        updateOfflineBadge();
                        syncOfflineQueue(); // sync next
                    } else {
                        console.log("Sync rejected:", data.message);
                    }
                })
                .catch(err => {
                    isSyncing = false;
                    console.log("Sync failed, retrying in 5s:", err);
                });
        }
        
        window.addEventListener("online", syncOfflineQueue);
        setInterval(syncOfflineQueue, 5000);
        
        window.addEventListener("load", () => {
            updateOfflineBadge();
            syncOfflineQueue();
        });

        // Fetch allowed operators on load
        function fetchAllowedOperators() {
            fetch("/get_allowed_operators?t=" + Date.now())
                .then(res => res.json())
                .then(data => {
                    if (data.status === "success") {
                        opNameSelect.innerHTML = "";
                        data.operators.forEach(name => {
                            const opt = document.createElement("option");
                            opt.value = name;
                            opt.textContent = name;
                            opNameSelect.appendChild(opt);
                        });
                    }
                })
                .catch(err => {
                    opNameSelect.innerHTML = "<option value=''>Y\u00fckl\u0259m\u0259 x\u0259tas\u0131</option>";
                });
        }

        // Initialize state
        if (operatorName) {
            showMainInterface();
        } else {
            fetchAllowedOperators();
        }

        saveOpBtn.addEventListener("click", () => {
            const name = opNameSelect.value;
            const pin = opPinInput.value.trim();
            if (!name) {
                alert("Z\u0259hm\u0259t olmasa operatoru se\u00e7in!");
                return;
            }
            if (!pin) {
                alert("Z\u0259hm\u0259t olmasa PIN kodu daxil edin!");
                return;
            }
            
            fetch(`/login_operator?name=${encodeURIComponent(name)}&pin=${encodeURIComponent(pin)}&t=` + Date.now())
                .then(res => res.json())
                .then(data => {
                    if (data.status === "success") {
                        operatorName = name;
                        safeSetStorage("operatorName", name);
                        opPinInput.value = "";
                        showMainInterface();
                    } else {
                        alert(data.message || "Giri\u015f r\u0259dd edildi!");
                    }
                })
                .catch(err => alert("Serverl\u0259 ba\u011flant\u0131 x\u0259tas\u0131!"));
        });

        logoutBtn.addEventListener("click", () => {
            operatorName = "";
            safeRemoveStorage("operatorName");
            opSetupCard.style.display = "block";
            mainCountInterface.style.display = "none";
            historyCard.style.display = "none";
            stopScanner();
            fetchAllowedOperators();
        });

        function showMainInterface() {
            opSetupCard.style.display = "none";
            mainCountInterface.style.display = "block";
            historyCard.style.display = "block";
            opDisplayName.textContent = operatorName;
            barcodeInput.focus();
            loadLocalHistory();
            loadProductsFromServer();
        }

        function showStatus(text, isError = false) {
            statusBox.textContent = text;
            statusBox.className = "status-msg " + (isError ? "status-error" : "status-success");
            statusBox.style.display = "block";
            setTimeout(() => {
                statusBox.style.display = "none";
            }, 4000);
        }

        // Search barcode (detect scanner speed vs manual typing)
        let keyTimes = [];

        barcodeInput.addEventListener("keydown", (e) => {
            if (e.key === "Enter" || e.keyCode === 13) {
                if (searchTimeout) clearTimeout(searchTimeout);
                const val = barcodeInput.value.trim();
                lastBarcodeValue = val;
                searchProduct(val);
                return;
            }
            keyTimes.push(Date.now());
            if (keyTimes.length > 5) keyTimes.shift();
        });

        barcodeInput.addEventListener("input", (e) => {
            const val = barcodeInput.value.trim();
            lastBarcodeValue = val;
            
            // Check key timing to identify physical scanner vs manual keyboard input
            let isFast = false;
            if (keyTimes.length >= 2) {
                const totalDiff = keyTimes[keyTimes.length - 1] - keyTimes[0];
                const avgDiff = totalDiff / (keyTimes.length - 1);
                if (avgDiff < 55) { // Average time between keypresses < 55ms
                    isFast = true;
                }
            }

            // Treat paste operations or extremely fast typing (like hardware scanners) as barcode scan
            if (e.inputType === "insertFromPaste" || (val.length >= 6 && isFast)) {
                if (searchTimeout) clearTimeout(searchTimeout);
                searchTimeout = setTimeout(() => {
                    searchProduct(val);
                }, 50); // Small delay to let input populate fully
            } else {
                // Manual input: do not auto-search. The user will select from autocomplete,
                // click "Axtar", or press Enter.
            }
        });

        searchBtn.addEventListener("click", () => {
            const val = barcodeInput.value.trim();
            lastBarcodeValue = val;
            searchProduct(val);
        });

        function searchProduct(barcode, playSound = true) {
            if (!barcode) return;
            if (searchTimeout) clearTimeout(searchTimeout);
            stopScanner();
            
            let localProd = allProducts[barcode];
            if (selectedBrand && localProd && localProd.brend !== selectedBrand) {
                showStatus(`X\u0259ta: Bu m\u0259hsul se\u00e7ilmi\u015f brend\u0259 (${selectedBrand}) aid deyil!`, true);
                playBeep(400, 300);
                resetInputs(false);
                return;
            }
            
            fetch(`/get_product_details?barcode=${encodeURIComponent(barcode)}&operator=${encodeURIComponent(operatorName)}`)
                .then(res => res.json())
                .then(data => {
                    if (data.status === "success") {
                        currentProduct = data.product;
                        if (selectedBrand && currentProduct.brend !== selectedBrand) {
                            showStatus(`X\u0259ta: Bu m\u0259hsul se\u00e7ilmi\u015f brend\u0259 (${selectedBrand}) aid deyil!`, true);
                            playBeep(400, 300);
                            resetInputs(false);
                            return;
                        }
                        
                        if (playSound) {
                            playBeep(1000, 80, 0.4);
                        }
                        
                        prodNameLbl.textContent = `${currentProduct.kod ? '[' + currentProduct.kod + '] ' : ''}${currentProduct.brend} - ${currentProduct.adi}`;
                        prodPriceLbl.textContent = `${currentProduct.qiymet.toFixed(2)} AZN`;
                        prodStockLbl.textContent = currentProduct.qaliq.toFixed(2);
                        prodPrevLbl.textContent = currentProduct.yeni !== null 
                            ? `${currentProduct.yeni.toFixed(2)} ${currentProduct.operator ? '(' + currentProduct.operator + ' t\u0259r\u0259find\u0259n)' : ''}` 
                            : "Say\u0131lmay\u0131b";
                        
                        productDetailsBox.style.display = "block";
                        qtyGroup.style.display = "block";
                        submitBtn.style.display = "block";
                        cancelBtn.style.display = "block";
                        
                        qtyInput.value = currentProduct.yeni !== null ? currentProduct.yeni : "";
                        qtyInput.focus();
                        qtyInput.select();
                    } else {
                        // Soft search client-side
                        let queryLower = barcode.toLowerCase().trim();
                        let matches = [];
                        for (let b in allProducts) {
                            let p = allProducts[b];
                            if (selectedBrand && p.brend !== selectedBrand) continue;
                            let name = `${p.brend} ${p.adi || ""}`.toLowerCase();
                            if (b.includes(queryLower) || name.includes(queryLower) || (p.kod && p.kod.toLowerCase().includes(queryLower))) {
                                matches.push(b);
                            }
                        }
                        
                        if (matches.length === 1) {
                            searchProduct(matches[0]);
                        } else if (matches.length > 1) {
                            listSearchInput.value = barcode;
                            switchToListTab();
                            renderProductList();
                            showStatus(`${matches.length} m\u0259hsul tap\u0131ld\u0131, siyah\u0131dan se\u00e7in.`);
                        } else {
                            showStatus("M\u0259hsul tap\u0131lmad\u0131! Yeni m\u0259hsul kimi \u0259lav\u0259 oluna bil\u0259r.", true);
                            resetInputs(false);
                            if (confirm(`"${barcode}" barkodu tap\u0131lmad\u0131. Yeni m\u0259hsul kimi \u0259lav\u0259 etm\u0259k ist\u0259yirsiniz?`)) {
                                addNewProductFlow(barcode);
                            }
                        }
                    }
                })
                .catch(err => {
                    // Offline fallback
                    if (localProd) {
                        currentProduct = {
                            barcode: barcode,
                            kod: localProd.kod || "",
                            brend: localProd.brend,
                            adi: localProd.adi,
                            qiymet: localProd.qiymet,
                            qaliq: localProd.qaliq,
                            yeni: localProd.yeni,
                            operator: localProd.operator
                        };
                        prodNameLbl.textContent = `${currentProduct.kod ? '[' + currentProduct.kod + '] ' : ''}${currentProduct.brend} - ${currentProduct.adi}`;
                        prodPriceLbl.textContent = `${currentProduct.qiymet.toFixed(2)} AZN`;
                        prodStockLbl.textContent = currentProduct.qaliq.toFixed(2);
                        prodPrevLbl.textContent = currentProduct.yeni !== null 
                            ? `${currentProduct.yeni.toFixed(2)} ${currentProduct.operator ? '(' + currentProduct.operator + ' t\u0259r\u0259find\u0259n)' : ''}` 
                            : "Say\u0131lmay\u0131b";
                        
                        productDetailsBox.style.display = "block";
                        qtyGroup.style.display = "block";
                        submitBtn.style.display = "block";
                        cancelBtn.style.display = "block";
                        
                        qtyInput.value = currentProduct.yeni !== null ? currentProduct.yeni : "";
                        qtyInput.focus();
                        qtyInput.select();
                        
                        if (playSound) {
                            playBeep(1000, 80, 0.4);
                        }
                        showStatus("Oflayn rejimd\u0259 tap\u0131ld\u0131", false);
                    } else {
                        showStatus("\u015e\u0259b\u0259k\u0259 x\u0259tas\u0131! Serverl\u0259 ba\u011flant\u0131n\u0131 yoxlay\u0131n.", true);
                    }
                });
        }
        
        function addNewProductFlow(barcode) {
            const defaultName = selectedBrand ? (selectedBrand + " - ") : "";
            const name = prompt("M\u0259hsulun Brend v\u0259 Ad\u0131:", defaultName);
            if (!name) return;
            const price = parseFloat(prompt("Qiym\u0259ti (AZN):", "0.00"));
            if (isNaN(price)) return;
            const stock = parseFloat(prompt("Sistem Qal\u0131\u011f\u0131:", "0.00"));
            if (isNaN(stock)) return;
            
            fetch(`/add_count?barcode=${encodeURIComponent(barcode)}&qty=0&mode=set&name=${encodeURIComponent(name)}&price=${price}&stock=${stock}&operator=${encodeURIComponent(operatorName)}`)
                .then(res => res.json())
                .then(data => {
                    if (data.status === "success") {
                        showStatus("M\u0259hsul \u0259lav\u0259 olundu!");
                        playBeep(2200, 150);
                        loadProductsFromServer().then(() => searchProduct(barcode));
                    } else {
                        alert("X\u0259ta: " + data.message);
                    }
                })
                .catch(err => alert("Ba\u011flant\u0131 x\u0259tas\u0131!"));
        }
        submitBtn.addEventListener("click", submitCount);
        qtyInput.addEventListener("keydown", (e) => {
            if (e.key === "Enter") submitCount();
        });

        function submitCount() {
            if (!currentProduct) return;
            const product = currentProduct;
            const qtyVal = parseFloat(qtyInput.value.trim());
            if (isNaN(qtyVal) || qtyVal < 0) {
                alert("D\u00fczg\u00fcn say daxil edin!");
                return;
            }
            
            fetch(`/add_count?barcode=${encodeURIComponent(product.barcode)}&qty=${qtyVal}&mode=set&operator=${encodeURIComponent(operatorName)}`)
                .then(res => res.json())
                .then(data => {
                    if (data.status === "success") {
                        showStatus(`Yadda saxlan\u0131ld\u0131: ${qtyVal.toFixed(2)}`);
                        playBeep(1200, 150);
                        saveToLocalHistory(product.barcode, `${product.brend} - ${product.adi}`, qtyVal);
                        if (allProducts[product.barcode] && data.product) {
                            allProducts[product.barcode].yeni = data.product.yeni;
                            allProducts[product.barcode].operator = data.product.operator;
                        }
                        renderProductList();
                        resetInputs();
                    } else {
                        showStatus("X\u0259ta: " + data.message, true);
                        playBeep(400, 300);
                    }
                })
                .catch(err => {
                    saveToOfflineQueue(product.barcode, qtyVal, "set", `${product.brend} - ${product.adi}`);
                    showStatus("\u26a0\ufe0f \u0130nternet yoxdur! Say\u0131m yadda\u015fda saxlan\u0131ld\u0131.", false);
                    playBeep(500, 300);
                    saveToLocalHistory(product.barcode, `${product.brend} - ${product.adi}`, qtyVal);
                    if (allProducts[product.barcode]) {
                        allProducts[product.barcode].yeni = qtyVal;
                    }
                    resetInputs();
                });
        }

        cancelBtn.addEventListener("click", () => resetInputs());

        function resetInputs(clearBarcode = true) {
            if (searchTimeout) clearTimeout(searchTimeout);
            currentProduct = null;
            if (clearBarcode) {
                barcodeInput.value = "";
                lastBarcodeValue = "";
            }
            qtyInput.value = "";
            productDetailsBox.style.display = "none";
            qtyGroup.style.display = "none";
            submitBtn.style.display = "none";
            cancelBtn.style.display = "none";
            if (clearBarcode) barcodeInput.focus();
        }

        // Camera Scanner integration
        scanCameraBtn.addEventListener("click", () => {
            if (scannerContainer.style.display === "block") {
                stopScanner();
            } else {
                startScanner();
            }
        });

        stopScanBtn.addEventListener("click", stopScanner);
        
        switchCameraBtn.addEventListener("click", () => {
            if (!html5QrCode || !html5QrCode.isScanning || scannerCameras.length <= 1) return;
            
            currentCameraIndex = (currentCameraIndex + 1) % scannerCameras.length;
            const nextCameraId = scannerCameras[currentCameraIndex].id;
            
            isTorchOn = false;
            if (toggleTorchBtn) {
                toggleTorchBtn.textContent = "💡 Fənər";
                toggleTorchBtn.style.backgroundColor = "rgba(255, 255, 255, 0.15)";
            }
            
            html5QrCode.stop().then(() => {
                html5QrCode.clear();
                html5QrCode = new Html5Qrcode("reader");
                runScannerWithId(nextCameraId);
            }).catch(err => {
                console.error("Camera switch error:", err);
            });
        });

        function runScannerWithId(cameraIdOrFacingMode) {
            const config = { fps: 15 };
            
            html5QrCode.start(
                cameraIdOrFacingMode, 
                config, 
                (decodedText) => {
                    if (navigator.vibrate) navigator.vibrate(100);
                    playBeep();
                    barcodeInput.value = decodedText;
                    stopScanner();
                    searchProduct(decodedText, false);
                },
                (errorMessage) => {}
            ).then(() => {
                setTimeout(() => {
                    try {
                        const capabilities = html5QrCode.getRunningTrackCapabilities();
                        if (capabilities.torch) {
                            toggleTorchBtn.style.display = "block";
                            isTorchOn = false;
                            toggleTorchBtn.textContent = "💡 Fənər";
                            toggleTorchBtn.style.backgroundColor = "rgba(255, 255, 255, 0.15)";
                        } else {
                            toggleTorchBtn.style.display = "none";
                        }
                    } catch(e) {
                        toggleTorchBtn.style.display = "none";
                    }
                }, 500);
            }).catch(err => {
                alert("Kamera açılmadı: " + err);
                stopScanner();
            });
        }

        function startScanner() {
            resetInputs(true);
            scannerContainer.style.display = "block";
            
            if (toggleTorchBtn) toggleTorchBtn.style.display = "none";
            if (switchCameraBtn) switchCameraBtn.style.display = "none";

            if (typeof Html5Qrcode === 'undefined') {
                alert("Skaner kitabxanası yüklənməyib. Zəhmət olmasa internet bağlantısını yoxlayın.");
                scannerContainer.style.display = "none";
                return;
            }
            
            html5QrCode = new Html5Qrcode("reader");
            
            Html5Qrcode.getCameras().then(devices => {
                scannerCameras = devices || [];
                if (scannerCameras.length > 1) {
                    if (switchCameraBtn) {
                        switchCameraBtn.style.display = "block";
                    }
                } else {
                    if (switchCameraBtn) {
                        switchCameraBtn.style.display = "none";
                    }
                }
                
                let defaultCameraId = { facingMode: "environment" };
                currentCameraIndex = -1;
                for (let i = 0; i < scannerCameras.length; i++) {
                    const label = scannerCameras[i].label.toLowerCase();
                    if (label.includes("back") || label.includes("rear") || label.includes("environment") || label.includes("arka")) {
                        defaultCameraId = scannerCameras[i].id;
                        currentCameraIndex = i;
                        break;
                    }
                }
                if (currentCameraIndex === -1 && scannerCameras.length > 0) {
                    defaultCameraId = scannerCameras[0].id;
                    currentCameraIndex = 0;
                }
                
                runScannerWithId(defaultCameraId);
            }).catch(err => {
                console.warn("Could not get cameras:", err);
                runScannerWithId({ facingMode: "environment" });
            });
        }

        function stopScanner() {
            if (toggleTorchBtn) toggleTorchBtn.style.display = "none";
            if (switchCameraBtn) switchCameraBtn.style.display = "none";
            isTorchOn = false;
            if (html5QrCode && html5QrCode.isScanning) {
                html5QrCode.stop().then(() => {
                    html5QrCode.clear();
                    html5QrCode = null;
                    scannerContainer.style.display = "none";
                }).catch(err => {
                    console.log("Stop error: ", err);
                    scannerContainer.style.display = "none";
                });
            } else {
                scannerContainer.style.display = "none";
            }
        }

        // Local History
        function saveToLocalHistory(barcode, name, qty) {
            let history = safeGetJSON("scanHistory");
            history = history.filter(item => item && item.barcode !== barcode);
            history.unshift({ barcode, name, qty, time: new Date().toLocaleTimeString("az-AZ") });
            if (history.length > 5) history.pop();
            safeSetStorage("scanHistory", JSON.stringify(history));
            loadLocalHistory();
        }

        function loadLocalHistory() {
            try {
                let history = safeGetJSON("scanHistory");
                history = history.filter(item => item && item.barcode);
                if (history.length === 0) {
                    historyList.innerHTML = `<div style="text-align: center; color: var(--text-muted); font-size: 0.85rem; padding: 10px;">H\u0259l\u0259 ki say\u0131m yoxdur</div>`;
                    return;
                }
                
                historyList.innerHTML = history.map(item => {
                    const qty = typeof item.qty === 'number' ? item.qty : parseFloat(item.qty) || 0;
                    return `
                        <div class="history-item" onclick="searchProduct('${item.barcode}')" style="cursor: pointer; display: flex; align-items: center; justify-content: space-between;">
                            <div style="flex: 1; text-align: left; padding-right: 10px;">
                                <div class="history-name">${item.name || 'M\u0259hsul'}</div>
                                <div class="history-details">Barkod: ${item.barcode || ''} | Saat: ${item.time || ''}</div>
                            </div>
                            <div style="display: flex; align-items: center; gap: 12px; margin-left: auto;">
                                <div class="history-qty">${qty.toFixed(2)}</div>
                                <button onclick="event.stopPropagation(); resetProductCount('${item.barcode}')" style="background: rgba(239, 68, 68, 0.1); border: 1px solid rgba(239, 68, 68, 0.15); color: var(--danger); border-radius: 6px; padding: 4px 8px; font-size: 0.85rem; cursor: pointer; outline: none; display: flex; align-items: center; justify-content: center;">\u1f5d1\ufe0f</button>
                            </div>
                        </div>
                    `;
                }).join("");
            } catch(e) {
                historyList.innerHTML = `<div style="text-align: center; color: var(--text-muted); font-size: 0.85rem; padding: 10px;">H\u0259l\u0259 ki say\u0131m yoxdur</div>`;
            }
        }

        function resetProductCount(barcode) {
            let history = safeGetJSON("scanHistory");
            let item = history.find(x => x && x.barcode === barcode);
            let name = item ? item.name : barcode;
            
            if (!confirm(`"${name}" m\u0259hsulunun say\u0131m\u0131n\u0131 s\u0131f\u0131rlamaq (silm\u0259k) ist\u0259yirsiniz?`)) return;
            
            fetch(`/add_count?barcode=${encodeURIComponent(barcode)}&mode=reset&operator=${encodeURIComponent(operatorName)}`)
                .then(res => res.json())
                .then(data => {
                    if (data.status === "success") {
                        showStatus("Say\u0131m s\u0131f\u0131rland\u0131!");
                        playBeep(800, 150);
                        history = history.filter(item => item && item.barcode !== barcode);
                        safeSetStorage("scanHistory", JSON.stringify(history));
                        loadLocalHistory();
                        if (allProducts[barcode]) {
                            allProducts[barcode].yeni = null;
                            allProducts[barcode].operator = "";
                        }
                        renderProductList();
                        if (currentProduct && currentProduct.barcode === barcode) resetInputs();
                    } else {
                        showStatus("X\u0259ta: " + data.message, true);
                        playBeep(400, 300);
                    }
                })
                .catch(err => {
                    saveToOfflineQueue(barcode, 0, "reset", name);
                    showStatus("\u26a0\ufe0f \u0130nternet yoxdur! S\u0131f\u0131rlama yadda\u015fda saxlan\u0131ld\u0131.", false);
                    playBeep(500, 300);
                    history = history.filter(item => item && item.barcode !== barcode);
                    safeSetStorage("scanHistory", JSON.stringify(history));
                    loadLocalHistory();
                    if (allProducts[barcode]) allProducts[barcode].yeni = null;
                    if (currentProduct && currentProduct.barcode === barcode) resetInputs();
                });
        }

        // Product list and tabs
        function loadProductsFromServer(retries = 5) {
            if (isProductsLoading && retries === 5) return Promise.resolve();
            isProductsLoading = true;
            return fetch(`/get_products?operator=${encodeURIComponent(operatorName || '')}&_t=${Date.now()}`)
                .then(res => {
                    if (!res.ok) throw new Error("HTTP error " + res.status);
                    return res.json();
                })
                .then(data => {
                    isProductsLoading = false;
                    const fetchedProducts = data.products || {};
                    if (Object.keys(fetchedProducts).length > 0) {
                        allProducts = fetchedProducts;
                        if (totalCountBadge) totalCountBadge.textContent = Object.keys(allProducts).length;
                        updateBrandSelectOptions();
                        renderProductList();
                    } else if (retries > 0) {
                        setTimeout(() => loadProductsFromServer(retries - 1), 1500);
                    }
                })
                .catch(err => {
                    isProductsLoading = false;
                    console.error("Products load error:", err);
                    if (retries > 0) {
                        setTimeout(() => loadProductsFromServer(retries - 1), 2000);
                    }
                });
        }

        function updateBrandSelectOptions() {
            if (!brandSelect) return;
            const currentSelected = selectedBrand;
            const brands = new Set();
            for (let b in allProducts) {
                const item = allProducts[b];
                if (item && item.brend) {
                    const bName = String(item.brend).trim();
                    if (bName) brands.add(bName);
                }
            }
            
            const sortedBrands = Array.from(brands).sort();
            brandSelect.innerHTML = '<option value="">[ B\u00fct\u00fcn Brendl\u0259r ]</option>';
            sortedBrands.forEach(brand => {
                const opt = document.createElement("option");
                opt.value = brand;
                opt.textContent = brand;
                brandSelect.appendChild(opt);
            });
            
            if (brands.has(currentSelected)) {
                brandSelect.value = currentSelected;
                selectedBrand = currentSelected;
            } else {
                brandSelect.value = "";
                selectedBrand = "";
                safeRemoveStorage("selectedBrand");
            }
        }

        if (brandSelect) {
            brandSelect.addEventListener("focus", () => {
                if (Object.keys(allProducts).length === 0) {
                    loadProductsFromServer();
                }
            });
            brandSelect.addEventListener("click", () => {
                if (Object.keys(allProducts).length === 0) {
                    loadProductsFromServer();
                }
            });
        }

        brandSelect.addEventListener("change", () => {
            selectedBrand = brandSelect.value;
            if (selectedBrand) {
                safeSetStorage("selectedBrand", selectedBrand);
            } else {
                safeRemoveStorage("selectedBrand");
            }
            renderProductList();
        });

        function renderProductList() {
            const query = listSearchInput.value.toLowerCase().trim();
            const listContainer = document.getElementById("mobile-product-list");
            listContainer.innerHTML = "";
            
            const items = [];
            for (let barcode in allProducts) {
                const p = allProducts[barcode];
                if (selectedBrand && p.brend !== selectedBrand) continue;
                const name = `${p.brend} ${p.adi || ""}`.toLowerCase();
                if (!query || name.includes(query) || barcode.includes(query) || (p.kod && p.kod.toLowerCase().includes(query))) {
                    items.push({ barcode, ...p });
                }
            }
            
            items.sort((a, b) => a.brend.localeCompare(b.brend));
            
            const renderedItems = items.slice(0, 50).map(p => {
                const diffQty = p.yeni !== null ? (p.yeni - p.qaliq) : 0;
                let qtyColor = "var(--text-muted)";
                let statusText = "";
                if (p.yeni !== null) {
                    if (diffQty > 0) {
                        qtyColor = "var(--success)";
                        statusText = `+${diffQty.toFixed(2)} Art\u0131q`;
                    } else if (diffQty < 0) {
                        qtyColor = "var(--danger)";
                        statusText = `${diffQty.toFixed(2)} \u018fskik`;
                    } else {
                        qtyColor = "var(--text-color)";
                        statusText = "D\u00fcz";
                    }
                }
                
                return `
                    <div class="product-list-item" onclick="selectProductFromList('${p.barcode}')" style="cursor: pointer; border-left: 4px solid ${p.yeni !== null ? 'var(--success)' : 'var(--border)'};">
                        <div class="product-list-item-left">
                            <div style="font-weight: 600;">${p.kod ? '[' + p.kod + '] ' : ''}${p.brend} ${p.adi}</div>
                            <div style="font-size: 0.75rem; color: var(--text-muted); margin-top: 3px;">Barkod: ${p.barcode} | Qal\u0131q: ${p.qaliq.toFixed(2)}</div>
                        </div>
                        <div class="product-list-item-right">
                            <div style="font-weight: 700; color: ${qtyColor};">${p.yeni !== null ? p.yeni.toFixed(2) : '-'}</div>
                            <div style="font-size: 0.7rem; color: ${qtyColor}; font-weight: 600;">${statusText}</div>
                        </div>
                    </div>
                `;
            }).join("");
            
            listContainer.innerHTML = renderedItems || `<div style="text-align: center; color: var(--text-muted); padding: 15px; font-size: 0.85rem;">M\u0259hsul tap\u0131lmad\u0131</div>`;
            if (items.length > 50) {
                listContainer.innerHTML += `<div style="text-align: center; color: var(--text-muted); font-size: 0.8rem; padding: 5px;">Daha ${items.length - 50} m\u0259hsul var, axtar\u0131\u015f\u0131 d\u0259qiql\u0259\u015fdirin...</div>`;
            }
        }

        function selectProductFromList(barcode) {
            switchToScanTab();
            barcodeInput.value = barcode;
            searchProduct(barcode);
        }

        tabScanBtn.addEventListener("click", switchToScanTab);
        tabListBtn.addEventListener("click", switchToListTab);
        listSearchInput.addEventListener("input", renderProductList);

        if (clearHistoryBtn = document.getElementById("clear-history-btn")) {
            clearHistoryBtn.addEventListener("click", () => {
                let history = safeGetJSON("scanHistory");
                if (history.length === 0) {
                    showStatus("Siyah\u0131 onsuz da bo\u015fdur.", true);
                    return;
                }
                
                const resetServer = confirm("B\u00fct\u00fcn son say\u0131mlar\u0131 SERVERD\u018f s\u0131f\u0131rlamaq (silm\u0259k) ist\u0259yirsiniz?\n\n[OK] - Say\u0131mlar\u0131 serverd\u0259n sil v\u0259 siyah\u0131n\u0131 t\u0259mizl\u0259\n[Cancel] - Yaln\u0131z telefondak\u0131 bu siyah\u0131n\u0131 t\u0259mizl\u0259");
                if (resetServer) {
                    let promises = history.filter(item => item && item.barcode).map(item => {
                        return fetch(`/add_count?barcode=${encodeURIComponent(item.barcode)}&mode=reset&operator=${encodeURIComponent(operatorName)}`)
                            .then(res => res.json())
                            .catch(err => {
                                saveToOfflineQueue(item.barcode, 0, "reset", item.name);
                                if (allProducts[item.barcode]) allProducts[item.barcode].yeni = null;
                                return { status: "success" };
                            });
                    });
                    
                    Promise.all(promises).then(() => {
                        safeRemoveStorage("scanHistory");
                        loadLocalHistory();
                        history.forEach(item => {
                            if (item && item.barcode && allProducts[item.barcode]) {
                                allProducts[item.barcode].yeni = null;
                                allProducts[item.barcode].operator = "";
                            }
                        });
                        renderProductList();
                        resetInputs();
                        showStatus("B\u00fct\u00fcn son say\u0131mlar\u0131 serverd\u0259 s\u0131f\u0131rland\u0131!");
                        playBeep(800, 200);
                    });
                } else {
                    safeRemoveStorage("scanHistory");
                    loadLocalHistory();
                    showStatus("Siyah\u0131 telefondan t\u0259mizl\u0259ndi.");
                    playBeep(600, 100);
                }
            });
        }

        function switchToScanTab() {
            tabScanBtn.className = "btn tab-btn active";
            tabListBtn.className = "btn tab-btn";
            tabListBtn.style.backgroundColor = "var(--card-bg)";
            tabListBtn.style.color = "var(--text-muted)";
            tabScanBtn.style.backgroundColor = "";
            tabScanBtn.style.color = "";
            scanView.style.display = "block";
            listView.style.display = "none";
            historyCard.style.display = "block";
        }

        function switchToListTab() {
            tabListBtn.className = "btn tab-btn active";
            tabScanBtn.className = "btn tab-btn";
            tabScanBtn.style.backgroundColor = "var(--card-bg)";
            tabScanBtn.style.color = "var(--text-muted)";
            tabListBtn.style.backgroundColor = "";
            tabListBtn.style.color = "";
            scanView.style.display = "none";
            listView.style.display = "block";
            historyCard.style.display = "none";
            loadProductsFromServer();
        }

        if (toggleTorchBtn) {
            toggleTorchBtn.addEventListener("click", toggleTorch);
        }

        function toggleTorch() {
            if (!html5QrCode || !html5QrCode.isScanning) return;
            try {
                const capabilities = html5QrCode.getRunningTrackCapabilities();
                if (capabilities.torch) {
                    isTorchOn = !isTorchOn;
                    html5QrCode.applyVideoConstraints({
                        advanced: [{ torch: isTorchOn }]
                    }).then(() => {
                        toggleTorchBtn.textContent = isTorchOn ? "\u1f526 F\u0259n\u0259ri S\u00f6nd\u00fcr" : "\u1f526 F\u0259n\u0259r";
                        toggleTorchBtn.style.backgroundColor = isTorchOn ? "var(--accent)" : "rgba(255, 255, 255, 0.15)";
                    }).catch(err => {
                        console.error("Torch application error:", err);
                    });
                }
            } catch (e) {
                console.error("Torch capability error:", e);
            }
        }

        barcodeInput.addEventListener("input", () => {
            const val = barcodeInput.value.trim().toLowerCase();
            if (!val || val.length < 2) {
                autocompleteBox.innerHTML = "";
                autocompleteBox.style.display = "none";
                return;
            }
            
            const matches = [];
            for (let bc in allProducts) {
                const p = allProducts[bc];
                if (selectedBrand && p.brend !== selectedBrand) continue;
                const name = `${p.brend} ${p.adi || ""}`.toLowerCase();
                if (bc.toLowerCase().includes(val) || name.includes(val) || (p.kod && p.kod.toLowerCase().includes(val))) {
                    matches.push({ barcode: bc, ...p });
                }
            }
            
            if (matches.length > 0) {
                autocompleteBox.innerHTML = matches.slice(0, 5).map(p => `
                    <div class="autocomplete-item" onclick="window.selectAutocomplete('${p.barcode}')">
                        <div class="autocomplete-item-title">${p.kod ? '[' + p.kod + '] ' : ''}${p.brend} ${p.adi}</div>
                        <div class="autocomplete-item-subtitle">Barkod: ${p.barcode} | Qal\u0131q: ${p.qaliq.toFixed(2)}</div>
                    </div>
                `).join("");
                autocompleteBox.style.display = "block";
            } else {
                autocompleteBox.innerHTML = "";
                autocompleteBox.style.display = "none";
            }
        });

        window.selectAutocomplete = function(barcode) {
            barcodeInput.value = barcode;
            autocompleteBox.innerHTML = "";
            autocompleteBox.style.display = "none";
            searchProduct(barcode);
        };

        // Close autocomplete when clicking outside
        document.addEventListener("click", (e) => {
            if (e.target !== barcodeInput && e.target !== autocompleteBox) {
                autocompleteBox.innerHTML = "";
                autocompleteBox.style.display = "none";
            }
        });

        // Register Service Worker for PWA
        if ('serviceWorker' in navigator) {
            window.addEventListener('load', () => {
                navigator.serviceWorker.register('/sw.js')
                    .then(reg => console.log('Service Worker registered successfully!', reg.scope))
                    .catch(err => console.error('Service Worker registration failed:', err));
            });
        }
    </script>

    <!-- Scanner Area -->
    <div id="scanner-container">
        <div id="reader"></div>
        <div class="scanner-laser"></div>
        <div class="scanner-vignette"></div>
        <div class="scanner-actions">
            <button class="scanner-action-btn" id="toggle-torch-btn" style="display: none;">💡 Fənər</button>
            <button class="scanner-action-btn" id="switch-camera-btn" style="display: none;">🔄 Dəyiş</button>
            <button class="scanner-action-btn" id="stop-scan-btn">Dayandır</button>
        </div>
    </div>
</body>
</html>
"""

# Start Main Loop
root.protocol("WM_DELETE_WINDOW", save_and_close)
root.mainloop()

