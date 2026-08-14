import os
import io
import json
import sqlite3
import unicodedata
try:
    import psycopg2
except ImportError:
    psycopg2 = None
from flask import Flask, request, jsonify, render_template, send_file, redirect, session, url_for
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
import time
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max upload size
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "DoreGenceAdminSecretKey1991")

ADMIN_USERNAME = os.environ.get("FLASK_ADMIN_USER", "DoreGence")
ADMIN_PASSWORD = os.environ.get("FLASK_ADMIN_PASS", "Aslanov1991")

@app.after_request
def add_cors_headers(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Expose-Headers"] = "Content-Disposition"
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('admin_logged_in'):
            if request.path in ["/admin", "/download_excel", "/download_discrepancies_excel", "/download_counted_excel"]:
                return redirect(url_for("login"))
            else:
                return jsonify({"status": "error", "message": "Giriş icazəsi yoxdur!"}), 401
        return f(*args, **kwargs)
    return decorated_function

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        if request.is_json:
            data = request.json or {}
            username = data.get("username", "").strip()
            password = data.get("password", "").strip()
        else:
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "").strip()
            
        if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
            session['admin_logged_in'] = True
            return jsonify({"status": "success", "message": "Giriş uğurludur."})
        else:
            return jsonify({"status": "error", "message": "İstifadəçi adı və ya şifrə yanlışdır!"}), 401
            
    if session.get('admin_logged_in'):
        return redirect(url_for("admin"))
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.pop('admin_logged_in', None)
    return redirect(url_for("login"))

DATABASE_URL = os.environ.get("DATABASE_URL")

# --- DATABASE SETUP ---

import threading
import time

_db_pool = None
_db_lock = threading.Lock()
_sqlite_conn = None
_use_sqlite_fallback = False

def get_db_type():
    global _use_sqlite_fallback
    if _use_sqlite_fallback:
        return "sqlite"
    return "postgres" if DATABASE_URL else "sqlite"

def _setup_pg_pool():
    global _db_pool
    if psycopg2 is None:
        raise ImportError("DATABASE_URL konfiqurasiya edilib, lakin 'psycopg2' kitabxanası quraşdırılmayıb!")
    from psycopg2.pool import ThreadedConnectionPool
    _db_pool = ThreadedConnectionPool(minconn=2, maxconn=10, dsn=DATABASE_URL)

def _setup_sqlite():
    global _sqlite_conn
    _sqlite_conn = sqlite3.connect("inventory.db", check_same_thread=False)
    _sqlite_conn.row_factory = sqlite3.Row
    _sqlite_conn.execute("PRAGMA journal_mode=WAL")
    _sqlite_conn.execute("PRAGMA synchronous=NORMAL")
    _sqlite_conn.execute("PRAGMA cache_size=-8000")

def get_db_connection():
    global _db_pool, _sqlite_conn, _use_sqlite_fallback
    if DATABASE_URL and not _use_sqlite_fallback:
        try:
            if _db_pool is None:
                with _db_lock:
                    if _db_pool is None:
                        _setup_pg_pool()
            return _db_pool.getconn()
        except Exception as e:
            print("PostgreSQL connection error, falling back to SQLite:", e)
            _use_sqlite_fallback = True
            init_db()
            
    if _sqlite_conn is None:
        with _db_lock:
            if _sqlite_conn is None:
                _setup_sqlite()
    return _sqlite_conn

def _return_connection(conn):
    if DATABASE_URL and _db_pool and not _use_sqlite_fallback:
        try:
            _db_pool.putconn(conn)
        except Exception:
            pass

def execute_query(query, params=None, commit=True, fetch=False):
    db_type = get_db_type()
    if db_type == "sqlite":
        query = query.replace("%s", "?")
    
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        if params:
            cursor.execute(query, params)
        else:
            cursor.execute(query)
            
        if fetch:
            rows = cursor.fetchall()
            if db_type == "sqlite":
                result = [dict(row) for row in rows]
            else:
                colnames = [desc[0] for desc in cursor.description]
                result = [dict(zip(colnames, row)) for row in rows]
            return result
        
        if commit:
            conn.commit()
            
        return cursor.lastrowid if db_type == "sqlite" else None
    finally:
        cursor.close()
        if db_type == "sqlite":
            pass  # SQLite connection persists
        else:
            _return_connection(conn)

def init_db():
    global _use_sqlite_fallback
    db_type = get_db_type()
    print(f"Database type: {db_type}")
    
    if db_type == "postgres":
        try:
            # Test connection to PostgreSQL
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT 1")
            cursor.close()
            _return_connection(conn)
        except Exception as e:
            print("PostgreSQL initialization failed. Falling back to SQLite:", e)
            _use_sqlite_fallback = True
            db_type = "sqlite"
            print("Verilənlər bazası növü SQLite-a dəyişdirildi.")
            
    if db_type == "postgres":
        create_excel_table = """
        CREATE TABLE IF NOT EXISTS excel_template (
            id INT PRIMARY KEY,
            filename TEXT,
            file_bytes BYTEA
        )
        """
        create_products_table = """
        CREATE TABLE IF NOT EXISTS products (
            barcode VARCHAR(100) PRIMARY KEY,
            kod VARCHAR(100),
            brend VARCHAR(255),
            adi VARCHAR(255),
            qaliq NUMERIC,
            qiymet NUMERIC,
            yeni NUMERIC,
            operator VARCHAR(100),
            order_num INTEGER,
            row_idx INTEGER
        )
        """
        create_operators_table = """
        CREATE TABLE IF NOT EXISTS operators (
            username VARCHAR(100) PRIMARY KEY,
            pin VARCHAR(50)
        )
        """
        create_scan_logs_table = """
        CREATE TABLE IF NOT EXISTS scan_logs (
            id SERIAL PRIMARY KEY,
            barcode VARCHAR(100),
            operator VARCHAR(100),
            qty NUMERIC,
            mode VARCHAR(50),
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
        create_operator_counts_table = """
        CREATE TABLE IF NOT EXISTS operator_counts (
            barcode VARCHAR(100),
            operator VARCHAR(100),
            qty NUMERIC,
            PRIMARY KEY (barcode, operator)
        )
        """
    else:
        create_excel_table = """
        CREATE TABLE IF NOT EXISTS excel_template (
            id INTEGER PRIMARY KEY,
            filename TEXT,
            file_bytes BLOB
        )
        """
        create_products_table = """
        CREATE TABLE IF NOT EXISTS products (
            barcode TEXT PRIMARY KEY,
            kod TEXT,
            brend TEXT,
            adi TEXT,
            qaliq REAL,
            qiymet REAL,
            yeni REAL,
            operator TEXT,
            order_num INTEGER,
            row_idx INTEGER
        )
        """
        create_operators_table = """
        CREATE TABLE IF NOT EXISTS operators (
            username TEXT PRIMARY KEY,
            pin TEXT
        )
        """
        create_scan_logs_table = """
        CREATE TABLE IF NOT EXISTS scan_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            barcode TEXT,
            operator TEXT,
            qty REAL,
            mode TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
        create_operator_counts_table = """
        CREATE TABLE IF NOT EXISTS operator_counts (
            barcode TEXT,
            operator TEXT,
            qty REAL,
            PRIMARY KEY (barcode, operator)
        )
        """
        
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(create_excel_table)
        cursor.execute(create_products_table)
        cursor.execute(create_operators_table)
        cursor.execute(create_scan_logs_table)
        cursor.execute(create_operator_counts_table)
        conn.commit()
    finally:
        cursor.close()
        if get_db_type() == "postgres":
            _return_connection(conn)
        
    # Migration: Add 'kod' column to products table if missing
    try:
        execute_query("ALTER TABLE products ADD COLUMN kod TEXT")
    except Exception:
        pass
    
    # Create performance indexes
    try:
        execute_query("CREATE INDEX IF NOT EXISTS idx_operator_counts_barcode ON operator_counts(barcode)")
    except Exception:
        pass
    try:
        execute_query("CREATE INDEX IF NOT EXISTS idx_operator_counts_barcode_operator ON operator_counts(barcode, operator)")
    except Exception:
        pass
    try:
        execute_query("CREATE INDEX IF NOT EXISTS idx_products_operator ON products(operator)")
    except Exception:
        pass
    try:
        execute_query("CREATE INDEX IF NOT EXISTS idx_scan_logs_barcode ON scan_logs(barcode)")
    except Exception:
        pass
        
    # Seed default operators if empty
    ops = execute_query("SELECT * FROM operators", fetch=True)
    if not ops:
        default_ops = [
            ("Orxan", "1234"),
            ("Elmir", "2222"),
            ("Perviz", "3333"),
            ("Namiq", "4444"),
            ("Vuqar", "5555")
        ]
        
        # Try to load from local file if exists
        if os.path.exists("operatorlar.txt"):
            try:
                loaded_ops = []
                with open("operatorlar.txt", "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line and ":" in line:
                            parts = line.split(":", 1)
                            loaded_ops.append((parts[0].strip(), parts[1].strip()))
                if loaded_ops:
                    default_ops = loaded_ops
            except Exception as e:
                print("operatorlar.txt read error:", e)
                
        for username, pin in default_ops:
            try:
                if db_type == "postgres":
                    execute_query("INSERT INTO operators (username, pin) VALUES (%s, %s) ON CONFLICT (username) DO NOTHING", (username, pin))
                else:
                    execute_query("INSERT OR IGNORE INTO operators (username, pin) VALUES (%s, %s)", (username, pin))
            except Exception as e:
                print(f"Failed to insert operator {username}:", e)

# Run DB initialization on startup
init_db()

# --- UTILITIES ---

def normalize_header(text):
    if not text:
        return ""
    text = str(text).strip()
    replacements = {
        'İ': 'i', 'I': 'i', 'İ': 'i', 'ı': 'i',
        'Ə': 'e', 'ə': 'e',
        'Ö': 'o', 'ö': 'o',
        'Ü': 'u', 'ü': 'u',
        'Ğ': 'g', 'ğ': 'g',
        'Ç': 'c', 'ç': 'c',
        'Ş': 's', 'ş': 's'
    }
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)
    text = text.lower()
    text = unicodedata.normalize('NFD', text)
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    return " ".join(text.split())

def clean_operator_string(operator_str):
    if not operator_str:
        return ""
    parts = str(operator_str).split(",")
    clean_parts = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if ":" in part:
            name = part.split(":")[0].strip()
            if name:
                clean_parts.append(name)
        else:
            clean_parts.append(part)
    return ", ".join(clean_parts)

import time

active_operators = {}

# --- SIMPLE CACHE for /get_products ---
_products_cache = None
_products_cache_time = 0
CACHE_TTL = 5  # seconds

def invalidate_products_cache():
    global _products_cache, _products_cache_time
    _products_cache = None
    _products_cache_time = 0

@app.before_request
def update_operator_activity():
    if request.method == "GET":
        operator = request.args.get("operator")
    elif request.method == "POST" and request.path in ["/add_count", "/api/sync"]:
        if request.is_json:
            operator = (request.json or {}).get("operator")
        else:
            operator = request.args.get("operator")
    else:
        return
    if operator:
        operator = operator.strip()
        if operator:
            active_operators[operator] = time.time()

# --- PWA ROUTES ---
@app.route("/manifest.json")
def serve_manifest():
    manifest_data = {
        "name": "Mobil Anbar Sayımı",
        "short_name": "Sayım",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#0F172A",
        "theme_color": "#6366F1",
        "orientation": "portrait",
        "icons": [
            {
                "src": "data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 100 100'><rect width='100' height='100' rx='20' fill='%236366F1'/><path d='M30 30h40v40H30z' fill='none' stroke='white' stroke-width='6'/><circle cx='50' cy='50' r='10' fill='white'/></svg>",
                "sizes": "192x192 512x512",
                "type": "image/svg+xml"
            }
        ]
    }
    return jsonify(manifest_data)

@app.route("/sw.js")
def serve_sw():
    sw_code = """
    const CACHE_NAME = 'anbar-sayimi-v23';
    const ASSETS = [
        '/',
        '/html5-qrcode.min.js',
        'https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Outfit:wght@600;700;800&display=swap'
    ];
    
    self.addEventListener('install', (e) => {
        e.waitUntil(
            caches.open(CACHE_NAME).then((cache) => {
                return cache.addAll(ASSETS);
            }).then(() => self.skipWaiting())
        );
    });
    
    self.addEventListener('activate', (e) => {
        e.waitUntil(
            caches.keys().then((keys) => {
                return Promise.all(
                    keys.map((key) => {
                        if (key !== CACHE_NAME) {
                            return caches.delete(key);
                        }
                    })
                );
            }).then(() => self.clients.claim())
        );
    });
    
    self.addEventListener('fetch', (e) => {
        if (e.request.method !== 'GET' || e.request.url.includes('/add_count') || e.request.url.includes('/get_products') || e.request.url.includes('/get_product_details') || e.request.url.includes('/api/active_operators') || e.request.url.includes('/download_')) {
            return;
        }
        e.respondWith(
            caches.match(e.request).then((cachedResponse) => {
                if (cachedResponse) {
                    return cachedResponse;
                }
                return fetch(e.request).then((response) => {
                    if (response && response.status === 200 && response.type === 'basic') {
                        const responseToCache = response.clone();
                        caches.open(CACHE_NAME).then((cache) => {
                            cache.put(e.request, responseToCache);
                        });
                    }
                    return response;
                }).catch(() => {
                    if (e.request.mode === 'navigate') {
                        return caches.match('/');
                    }
                });
            })
        );
    });
    """
    return app.response_class(sw_code.strip(), mimetype="application/javascript")

@app.route("/api/active_operators")
def get_active_operators_api():
    now = time.time()
    expired = [op for op, last_act in active_operators.items() if now - last_act > 60]
    for op in expired:
        active_operators.pop(op, None)
        
    ops_list = []
    for name, last_act in sorted(active_operators.items(), key=lambda x: x[0]):
        elapsed = int(now - last_act)
        if elapsed < 5:
            time_str = "İndi"
        elif elapsed < 60:
            time_str = f"{elapsed} saniyə əvvəl"
        else:
            time_str = f"{elapsed // 60} dəqiqə əvvəl"
        ops_list.append({
            "name": name,
            "last_activity": last_act,
            "time_str": time_str
        })
        
    return jsonify({"status": "success", "operators": ops_list})

# --- CONTROLLER / ROUTER ---

@app.route("/")
def index():
    return render_template("operator.html")

@app.route("/admin")
@admin_required
def admin():
    return render_template("admin.html")

@app.route("/html5-qrcode.min.js")
def serve_qrcode_js():
    if os.path.exists("html5-qrcode.min.js"):
        return send_file("html5-qrcode.min.js", mimetype="application/javascript")
    else:
        return redirect("https://unpkg.com/html5-qrcode/html5-qrcode.min.js")

def sync_operators_to_file():
    try:
        ops = execute_query("SELECT username, pin FROM operators ORDER BY username ASC", fetch=True)
        with open("operatorlar.txt", "w", encoding="utf-8") as f:
            for op in ops:
                f.write(f"{op['username']}:{op['pin']}\n")
    except Exception as e:
        print("Failed to sync operators to file:", e)

@app.route("/api/operators", methods=["GET", "POST"])
@admin_required
def manage_operators():
    if request.method == "POST":
        data = request.json or {}
        username = data.get("username", "").strip()
        pin = data.get("pin", "").strip()
        old_username = data.get("old_username", "").strip()
        if not username or not pin:
            return jsonify({"status": "error", "message": "Ad və PIN daxil edilməlidir!"}), 400
        try:
            if old_username and old_username != username:
                execute_query("UPDATE operators SET username = %s, pin = %s WHERE username = %s", (username, pin, old_username))
                execute_query("UPDATE products SET operator = %s WHERE operator = %s", (username, old_username))
            else:
                db_type = get_db_type()
                if db_type == "postgres":
                    execute_query("INSERT INTO operators (username, pin) VALUES (%s, %s) ON CONFLICT (username) DO UPDATE SET pin = EXCLUDED.pin", (username, pin))
                else:
                    execute_query("INSERT OR REPLACE INTO operators (username, pin) VALUES (%s, %s)", (username, pin))
            
            sync_operators_to_file()
            return jsonify({"status": "success", "message": "Operator məlumatları yadda saxlanıldı."})
        except Exception as e:
            return jsonify({"status": "error", "message": str(e)}), 500
    else:
        try:
            ops = execute_query("SELECT username, pin FROM operators ORDER BY username ASC", fetch=True)
            return jsonify({"status": "success", "operators": ops})
        except Exception as e:
            return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/operators/<username>", methods=["DELETE"])
@admin_required
def delete_operator(username):
    username = username.strip()
    if not username:
        return jsonify({"status": "error", "message": "İstifadəçi adı boşdur!"}), 400
    try:
        execute_query("DELETE FROM operators WHERE username = %s", (username,))
        sync_operators_to_file()
        return jsonify({"status": "success", "message": "Operator silindi."})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API ENDPOINTS ---

@app.route("/get_allowed_operators", methods=["GET"])
def get_allowed_operators():
    try:
        ops = execute_query("SELECT username FROM operators ORDER BY username ASC", fetch=True)
        names = [o['username'] for o in ops] if ops else ["Orxan", "Elmir", "Perviz", "Namiq", "Vuqar"]
        return jsonify({"status": "success", "operators": names})
    except Exception as e:
        return jsonify({"status": "success", "operators": ["Orxan", "Elmir", "Perviz", "Namiq", "Vuqar"]})

def verify_pin(stored_pin, provided_pin):
    if not stored_pin or not provided_pin:
        return False
    if stored_pin == provided_pin:
        return True
    try:
        return check_password_hash(stored_pin, provided_pin)
    except Exception:
        return False

@app.route("/login_operator", methods=["GET"])
def login_operator():
    name = request.args.get("name", "").strip()
    pin = request.args.get("pin", "").strip()
    
    if not name or not pin:
        return jsonify({"status": "error", "message": "Ad və PIN daxil edilməlidir!"}), 400
        
    try:
        res = execute_query("SELECT pin FROM operators WHERE username = %s", (name,), fetch=True)
        if res and verify_pin(res[0]['pin'], pin):
            return jsonify({"status": "success"})
        return jsonify({"status": "error", "message": "Yanlış PIN kod!"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/get_product_details", methods=["GET"])
def get_product_details():
    barcode = request.args.get("barcode", "").strip()
    operator = request.args.get("operator", "").strip()
    if not barcode:
        return jsonify({"status": "error", "message": "Barkod boş ola bilməz!"}), 400
        
    try:
        res = execute_query("SELECT * FROM products WHERE barcode = %s", (barcode,), fetch=True)
        if res:
            p = res[0]
            
            own_count = None
            if operator:
                own_res = execute_query("SELECT qty FROM operator_counts WHERE barcode = %s AND operator = %s", (barcode, operator), fetch=True)
                if own_res:
                    own_count = float(own_res[0]['qty'])
            
            breakdown_res = execute_query("SELECT operator FROM operator_counts WHERE barcode = %s AND qty > 0 ORDER BY operator ASC", (barcode,), fetch=True)
            breakdown_str = ", ".join([r['operator'] for r in breakdown_res]) if breakdown_res else ""
            
            if not breakdown_str and p['yeni'] is not None:
                op_val = p['operator'] or 'Naməlum'
                clean_parts = []
                for part in op_val.split(","):
                    clean_parts.append(part.split(":")[0].strip())
                breakdown_str = ", ".join(clean_parts)
                
            return jsonify({
                "status": "success",
                "product": {
                    "barcode": p["barcode"],
                    "kod": p.get("kod", ""),
                    "brend": p["brend"],
                    "adi": p["adi"],
                    "qaliq": float(p["qaliq"]) if p["qaliq"] is not None else 0.0,
                    "qiymet": float(p["qiymet"]) if p["qiymet"] is not None else 0.0,
                    "yeni": float(p["yeni"]) if p["yeni"] is not None else None,
                    "operator": breakdown_str if breakdown_str else (p["operator"] or ""),
                    "own_qty": own_count
                }
            })
        else:
            return jsonify({"status": "error", "message": "Məhsul tapılmadı"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/scan_logs/<barcode>", methods=["GET"])
def get_scan_logs(barcode):
    barcode = barcode.strip()
    try:
        logs = execute_query("SELECT operator, qty, mode, created_at FROM scan_logs WHERE barcode = %s ORDER BY id DESC", (barcode,), fetch=True)
        formatted_logs = []
        for log in logs:
            dt_val = log.get("created_at")
            if isinstance(dt_val, str):
                dt_str = dt_val.split(".")[0]
            elif dt_val:
                dt_str = dt_val.strftime("%Y-%m-%d %H:%M:%S")
            else:
                dt_str = "-"
            formatted_logs.append({
                "operator": log["operator"],
                "qty": float(log["qty"]) if log["qty"] is not None else 0.0,
                "mode": log["mode"],
                "created_at": dt_str
            })
        return jsonify({"status": "success", "logs": formatted_logs})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/operator_counts/<barcode>", methods=["GET"])
def get_operator_counts(barcode):
    barcode = barcode.strip()
    try:
        rows = execute_query(
            "SELECT operator, qty FROM operator_counts WHERE barcode = %s AND qty > 0 ORDER BY operator ASC",
            (barcode,), fetch=True
        )
        counts = [{"operator": r["operator"], "qty": float(r["qty"])} for r in rows]
        return jsonify({"status": "success", "counts": counts})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/get_products", methods=["GET"])
def get_products():
    global _products_cache, _products_cache_time
    now = time.time()
    if _products_cache is not None and (now - _products_cache_time) < CACHE_TTL:
        return jsonify(_products_cache)
    
    try:
        prods = execute_query("SELECT * FROM products ORDER BY row_idx ASC, barcode ASC", fetch=True)
        
        total_items = len(prods)
        total_counted = sum(1 for p in prods if p['yeni'] is not None)
        
        surplus_qty = 0.0
        shortage_qty = 0.0
        total_diff_val = 0.0
        
        products_dict = {}
        for p in prods:
            yeni = float(p['yeni']) if p['yeni'] is not None else None
            qaliq = float(p['qaliq']) if p['qaliq'] is not None else 0.0
            qiymet = float(p['qiymet']) if p['qiymet'] is not None else 0.0
            
            p_formatted = {
                'row': p['row_idx'],
                'kod': p.get('kod', ''),
                'brend': p['brend'],
                'adi': p['adi'],
                'qaliq': qaliq,
                'qiymet': qiymet,
                'yeni': yeni,
                'order': p['order_num'] or 0,
                'operator': p['operator'] or ""
            }
            products_dict[p['barcode']] = p_formatted
            
            if yeni is not None:
                diff = yeni - qaliq
                if diff > 0:
                    surplus_qty += diff
                else:
                    shortage_qty += abs(diff)
                total_diff_val += diff * qiymet
                
        _products_cache = {
            "products": products_dict,
            "stats": {
                "total_items": total_items,
                "total_counted": total_counted,
                "surplus_qty": surplus_qty,
                "shortage_qty": shortage_qty,
                "total_diff_val": total_diff_val
            }
        }
        _products_cache_time = now
        return jsonify(_products_cache)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/add_count", methods=["GET", "POST"])
def add_count():
    # Supports both GET (for compatibility with mobile client code) and POST
    if request.method == "POST":
        data = request.json or {}
        barcode = data.get("barcode", "").strip()
        qty_str = str(data.get("qty", "0")).strip()
        mode = data.get("mode", "add")
        operator = data.get("operator", "").strip()
        name = data.get("name", "").strip()
        price_str = str(data.get("price", "0")).strip()
        stock_str = str(data.get("stock", "0")).strip()
    else:
        barcode = request.args.get("barcode", "").strip()
        qty_str = request.args.get("qty", "0").strip()
        mode = request.args.get("mode", "add")
        operator = request.args.get("operator", "").strip()
        name = request.args.get("name", "").strip()
        price_str = request.args.get("price", "0").strip()
        stock_str = request.args.get("stock", "0").strip()
        
    if not barcode:
        return jsonify({"status": "error", "message": "Barkod daxil edin!"}), 400
        
    try:
        # Check if product exists
        res = execute_query("SELECT * FROM products WHERE barcode = %s", (barcode,), fetch=True)
        
        if res:
            p = res[0]
            if mode in ["reset", "delete"]:
                execute_query(
                    "UPDATE products SET yeni = NULL, operator = '', order_num = 0 WHERE barcode = %s",
                    (barcode,)
                )
                execute_query("DELETE FROM operator_counts WHERE barcode = %s", (barcode,))
                op_log = operator if operator else "Admin"
                execute_query(
                    "INSERT INTO scan_logs (barcode, operator, qty, mode) VALUES (%s, %s, 0, %s)",
                    (barcode, op_log, mode)
                )
                p['yeni'] = None
                p['operator'] = ""
                invalidate_products_cache()
            else:
                try:
                    qty = float(qty_str)
                except ValueError:
                    return jsonify({"status": "error", "message": "Yanlış say miqdarı!"}), 400
                    
                op_name = operator if operator else "Admin"
                current_own = 0.0
                own_res = execute_query("SELECT qty FROM operator_counts WHERE barcode = %s AND operator = %s", (barcode, op_name), fetch=True)
                if own_res:
                    current_own = float(own_res[0]['qty'])
                    
                if mode == "add":
                    new_own = current_own + qty
                else:  # 'set' mode
                    new_own = qty
                    
                db_type = get_db_type()
                if db_type == "postgres":
                    execute_query(
                        "INSERT INTO operator_counts (barcode, operator, qty) VALUES (%s, %s, %s) ON CONFLICT (barcode, operator) DO UPDATE SET qty = EXCLUDED.qty",
                        (barcode, op_name, new_own)
                    )
                else:
                    execute_query(
                        "INSERT OR REPLACE INTO operator_counts (barcode, operator, qty) VALUES (%s, %s, %s)",
                        (barcode, op_name, new_own)
                    )
                
                # Calculate total summed count
                sum_res = execute_query("SELECT SUM(qty) as total FROM operator_counts WHERE barcode = %s", (barcode,), fetch=True)
                new_yeni = float(sum_res[0]['total']) if sum_res and sum_res[0]['total'] is not None else new_own
                
                # Generate breakdown list of operators who counted this product
                ops_res = execute_query("SELECT operator, qty FROM operator_counts WHERE barcode = %s AND qty > 0 ORDER BY operator ASC", (barcode,), fetch=True)
                operators_list = ", ".join([f"{row['operator']}:{float(row['qty']):.2f}" for row in ops_res]) if ops_res else op_name
                
                execute_query(
                    "UPDATE products SET yeni = %s, operator = %s, order_num = (SELECT COALESCE(MAX(order_num), 0) + 1 FROM products) WHERE barcode = %s",
                    (new_yeni, operators_list, barcode)
                )
                op_log = operator if operator else "Admin"
                execute_query(
                    "INSERT INTO scan_logs (barcode, operator, qty, mode) VALUES (%s, %s, %s, %s)",
                    (barcode, op_log, qty, mode)
                )
                p['yeni'] = new_yeni
                p['operator'] = operators_list
                
            invalidate_products_cache()
            return jsonify({"status": "success", "product": {
                "barcode": p["barcode"],
                "brend": p["brend"],
                "adi": p["adi"],
                "qaliq": float(p["qaliq"]) if p["qaliq"] is not None else 0.0,
                "qiymet": float(p["qiymet"]) if p["qiymet"] is not None else 0.0,
                "yeni": float(p["yeni"]) if p["yeni"] is not None else None,
                "operator": p["operator"]
            }})
        else:
            # Manual addition of a new product
            if mode in ["reset", "delete"]:
                return jsonify({"status": "success", "message": "Məhsul tapılmadı, sıfırlamağa ehtiyac yoxdur"})
                
            if not name:
                return jsonify({"status": "error", "message": "Məhsul tapılmadı və brend adı yazılmadı!"}), 404
                
            try:
                price = float(price_str)
                stock = float(stock_str)
                qty = float(qty_str)
            except ValueError:
                return jsonify({"status": "error", "message": "Parametrlər düzgün ədəd formatında deyil!"}), 400
                
            op_name = operator if operator else "Admin"
            db_type = get_db_type()
            if db_type == "postgres":
                execute_query(
                    "INSERT INTO operator_counts (barcode, operator, qty) VALUES (%s, %s, %s) ON CONFLICT (barcode, operator) DO UPDATE SET qty = EXCLUDED.qty",
                    (barcode, op_name, qty)
                )
            else:
                execute_query(
                    "INSERT OR REPLACE INTO operator_counts (barcode, operator, qty) VALUES (%s, %s, %s)",
                    (barcode, op_name, qty)
                )
                
            # Insert into database with row_idx = NULL (manually added, will be appended to Excel)
            kod_param = request.args.get("kod", "").strip() or None
            operators_list = op_name if qty > 0 else ""
            execute_query(
                "INSERT INTO products (barcode, kod, brend, adi, qaliq, qiymet, yeni, operator, order_num, row_idx) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, (SELECT COALESCE(MAX(order_num), 0) + 1 FROM products), NULL)",
                (barcode, kod_param, name, "", stock, price, qty if qty > 0 else None, operators_list)
            )
            op_log = operator if operator else "Admin"
            execute_query(
                "INSERT INTO scan_logs (barcode, operator, qty, mode) VALUES (%s, %s, %s, %s)",
                (barcode, op_log, qty, "add")
            )
            
            invalidate_products_cache()
            return jsonify({"status": "success", "product": {
                "barcode": barcode,
                "brend": name,
                "adi": "",
                "qaliq": stock,
                "qiymet": price,
                "yeni": qty if qty > 0 else None,
                "operator": operators_list
            }})
            
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/sync", methods=["POST"])
def sync_offline_queue():
    data = request.json or {}
    queue = data.get("queue", [])
    operator = data.get("operator", "").strip()
    
    if not queue:
        return jsonify({"status": "success", "message": "Sinxronizasiya ediləcək məlumat yoxdur"})
        
    success_count = 0
    errors = []
    
    try:
        success_count = 0
        errors = []
        
        op_name = operator if operator else "Admin"
        
        for item in queue:
            barcode = item.get("barcode", "").strip()
            qty = float(item.get("qty", 0.0))
            mode = item.get("mode", "set")
            
            if not barcode:
                continue
                
            res = execute_query("SELECT * FROM products WHERE barcode = %s", (barcode,), fetch=True)
            if res:
                p = res[0]
                
                # Fetch own count
                current_own = 0.0
                own_res = execute_query("SELECT qty FROM operator_counts WHERE barcode = %s AND operator = %s", (barcode, op_name), fetch=True)
                if own_res:
                    current_own = float(own_res[0]['qty'])
                    
                if mode == "add":
                    new_own = current_own + qty
                else:
                    new_own = qty
                    
                db_type = get_db_type()
                if db_type == "postgres":
                    execute_query(
                        "INSERT INTO operator_counts (barcode, operator, qty) VALUES (%s, %s, %s) ON CONFLICT (barcode, operator) DO UPDATE SET qty = EXCLUDED.qty",
                        (barcode, op_name, new_own)
                    )
                else:
                    execute_query(
                        "INSERT OR REPLACE INTO operator_counts (barcode, operator, qty) VALUES (%s, %s, %s)",
                        (barcode, op_name, new_own)
                    )
                
                # Total sum
                sum_res = execute_query("SELECT SUM(qty) as total FROM operator_counts WHERE barcode = %s", (barcode,), fetch=True)
                new_yeni = float(sum_res[0]['total']) if sum_res and sum_res[0]['total'] is not None else new_own
                
                # Operators list (breakdown format)
                ops_res = execute_query("SELECT operator, qty FROM operator_counts WHERE barcode = %s AND qty > 0 ORDER BY operator ASC", (barcode,), fetch=True)
                operators_list = ", ".join([f"{row['operator']}:{float(row['qty']):.2f}" for row in ops_res]) if ops_res else op_name
                
                execute_query(
                    "UPDATE products SET yeni = %s, operator = %s, order_num = (SELECT COALESCE(MAX(order_num), 0) + 1 FROM products) WHERE barcode = %s",
                    (new_yeni, operators_list, barcode)
                )
                
                # Log scan
                execute_query(
                    "INSERT INTO scan_logs (barcode, operator, qty, mode) VALUES (%s, %s, %s, %s)",
                    (barcode, op_name, qty, mode)
                )
                success_count += 1
            else:
                errors.append(f"Məhsul tapılmadı: {barcode}")
                
        invalidate_products_cache()
        return jsonify({"status": "success", "synced": success_count, "errors": errors})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/upload_excel", methods=["POST"])
@admin_required
def upload_excel():
    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "Fayl seçilməyib!"}), 400
        
    file = request.files['file']
    if file.filename == '':
        return jsonify({"status": "error", "message": "Boş fayl adı!"}), 400
        
    if not file.filename.endswith('.xlsx'):
        return jsonify({"status": "error", "message": "Yalnız .xlsx (Excel) faylları qəbul olunur!"}), 400
        
    try:
        file_bytes = file.read()
        
        # Verify it loads correctly before saving
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
        sheet = wb.active
        
        # Map headers with Smart Alias Matcher
        headers = {}
        for col in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col).value
            if val:
                headers[normalize_header(val)] = col

        def find_col_index(header_dict, synonyms, exclude_indices=None):
            if exclude_indices is None:
                exclude_indices = set()
            # 1. Exact Match first
            for syn in synonyms:
                norm_syn = normalize_header(syn)
                for h_key, col_idx in header_dict.items():
                    if col_idx in exclude_indices:
                        continue
                    if norm_syn == h_key:
                        return col_idx
            # 2. Substring Match
            for syn in synonyms:
                norm_syn = normalize_header(syn)
                for h_key, col_idx in header_dict.items():
                    if col_idx in exclude_indices:
                        continue
                    if norm_syn in h_key or h_key in norm_syn:
                        return col_idx
            return None

        barkod_synonyms = [
            'barkod', 'barkkod', 'barcode', 'strixkod', 'strix-kod', 'strix_kod',
            'ean', 'shtrihkod', 'штрихкод', 'штрих-код', 'bar_code',
            'sh_kod', 'bar_kod', 'barkod_no', 'barkkod_no'
        ]
        
        kod_synonyms = [
            'kod', 'kodu', 'product_code', 'mal_kodu', 'mehsul_kodu',
            'kod_mehsul', 'kod_mal', 'код', 'артикул', 'artikul', 'item_code', 'item_id'
        ]

        brend_synonyms = [
            'brend', 'brand', 'marka', 'istehsalci', 'производитель',
            'фирма', 'firma', 'vendor', 'manufacturer', 'malin_markasi',
            'taminatci', 'supplier'
        ]

        adi_synonyms = [
            'mehsulun adi', 'mehsulun_adi', 'malin adi', 'malin_adi',
            'adi', 'name', 'mehsul', 'description', 'nomenklatura',
            'наименование', 'название', 'товар', 'urun_adi', 'tam_adi',
            'mehsul_adi', 'item_name', 'aciglama'
        ]

        qaliq_synonyms = [
            'anbar qaligi', 'anbar_qaligi', 'qaliq', 'stock', 'sistem qaligi',
            'sistem_qaligi', 'miqdar', 'sayi', 'qaliq_miqdari', 'остаток',
            'количество', 'stok_miktari', 'stok', 'balance', 'qty',
            'quantity', 'son_qaliq', 'tek_qaliq', 'mevcut'
        ]

        qiymet_synonyms = [
            'mehsulun qiymeti', 'mehsulun_qiymeti', 'qiymet', 'qiymeti',
            'price', 'satis_qiymeti', 'satis qiymeti', 'satis_qiymet', 'satis qiymet',
            'satiş qiyməti', 'satiş qiymət', 'maya_qiymeti', 'maya qiymeti', 'цена',
            'стоимость', 'fiyat', 'cost', 'unit_price', 'qiymat', 'mebleg',
            'perakende', 'pərakəndə', 'retail', 'qiymet azn', 'qiymeti azn',
            'qiymet (azn)', 'qiymeti (azn)'
        ]

        barkod_col = find_col_index(headers, barkod_synonyms)
        kod_col = find_col_index(headers, kod_synonyms, exclude_indices={barkod_col} if barkod_col else set())

        # If barkod_col is missing but kod_col exists, use kod_col as fallback barcode
        if not barkod_col and kod_col:
            barkod_col = kod_col
            kod_col = None

        brend_col = find_col_index(headers, brend_synonyms)
        adi_col = find_col_index(headers, adi_synonyms, exclude_indices={brend_col} if brend_col else set())

        if not brend_col:
            brend_col = adi_col
            if brend_col == adi_col:
                adi_col = None

        anbar_qaligi_col = find_col_index(headers, qaliq_synonyms)
        qiymet_col = find_col_index(headers, qiymet_synonyms)

        missing_cols = []
        if not brend_col: missing_cols.append("Brend (və ya Adı/Məhsulun Adı)")
        if not anbar_qaligi_col: missing_cols.append("Anbar Qalıq (və ya Qalıq/Miqdar/Остаток)")
        if not barkod_col: missing_cols.append("Barkod (və ya Штрихкод/Code)")

        if missing_cols:
            return jsonify({
                "status": "error",
                "message": f"Excel faylında aşağıdakı vacib sütunlar tapılmadı: {', '.join(missing_cols)}"
            }), 400
            
        # Clear products table
        execute_query("DELETE FROM products")
        execute_query("DELETE FROM operator_counts")
        
        # Store template in database
        db_type = get_db_type()
        if db_type == "postgres":
            param_bytes = psycopg2.Binary(file_bytes)
            execute_query(
                "INSERT INTO excel_template (id, filename, file_bytes) VALUES (1, %s, %s) ON CONFLICT (id) DO UPDATE SET filename = EXCLUDED.filename, file_bytes = EXCLUDED.file_bytes",
                (file.filename, param_bytes)
            )
        else:
            param_bytes = sqlite3.Binary(file_bytes)
            execute_query(
                "INSERT OR REPLACE INTO excel_template (id, filename, file_bytes) VALUES (1, %s, %s)",
                (file.filename, param_bytes)
            )
            
        def parse_num(val):
            if val is None:
                return 0.0
            if isinstance(val, (int, float)):
                return float(val)
            s = str(val).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
            try:
                return float(s)
            except ValueError:
                return 0.0

        # Load products from Excel
        products_to_insert = []
        for row_idx in range(2, sheet.max_row + 1):
            barcode_val = sheet.cell(row=row_idx, column=barkod_col).value
            if barcode_val is not None:
                if isinstance(barcode_val, float):
                    barcode_str = str(int(barcode_val)).strip()
                else:
                    barcode_str = str(barcode_val).strip()
                    
                if not barcode_str:
                    continue
                    
                kod_val = ""
                if kod_col:
                    k_val = sheet.cell(row=row_idx, column=kod_col).value
                    if k_val is not None:
                        if isinstance(k_val, float):
                            kod_val = str(int(k_val)).strip()
                        else:
                            kod_val = str(k_val).strip()

                brend_val = str(sheet.cell(row=row_idx, column=brend_col).value or "Naməlum Brend").strip()
                adi_val = ""
                if adi_col:
                    adi_val = str(sheet.cell(row=row_idx, column=adi_col).value or "").strip()
                elif brend_val and " - " in brend_val:
                    parts = brend_val.split(" - ", 1)
                    brend_val = parts[0].strip()
                    adi_val = parts[1].strip()
                    
                qaliq_val = sheet.cell(row=row_idx, column=anbar_qaligi_col).value
                qaliq = parse_num(qaliq_val)
                    
                qiymet_val = sheet.cell(row=row_idx, column=qiymet_col).value if qiymet_col else 0.0
                qiymet = parse_num(qiymet_val)
                    
                yeni_col = headers.get('yeni sayim') or headers.get('real say')
                yeni = None
                if yeni_col:
                    yeni_val = sheet.cell(row=row_idx, column=yeni_col).value
                    if yeni_val is not None and not str(yeni_val).startswith('='):
                        try:
                            yeni = float(yeni_val)
                        except ValueError:
                            yeni = None
                            
                operator_col = headers.get('operator') or headers.get('sayimci') or headers.get('user')
                operator = ""
                if operator_col:
                    operator_val = str(sheet.cell(row=row_idx, column=operator_col).value or "").strip()
                    clean_parts = []
                    for part in operator_val.split(","):
                        clean_parts.append(part.split(":")[0].strip())
                    operator = ", ".join(clean_parts)
                    
                products_to_insert.append((
                    barcode_str, kod_val, brend_val, adi_val, qaliq, qiymet, yeni, operator, 0, row_idx
                ))
                
        # Bulk insert
        conn = get_db_connection()
        cursor = conn.cursor()
        try:
            insert_query = "INSERT INTO products (barcode, kod, brend, adi, qaliq, qiymet, yeni, operator, order_num, row_idx) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)"
            if db_type == "sqlite":
                insert_query = insert_query.replace("%s", "?")
                
            seen_barcodes = set()
            filtered_products = []
            for p in products_to_insert:
                if p[0] not in seen_barcodes:
                    seen_barcodes.add(p[0])
                    filtered_products.append(p)
                    
            cursor.executemany(insert_query, filtered_products)
            conn.commit()
        finally:
            cursor.close()
            if db_type == "postgres":
                _return_connection(conn)
            
        invalidate_products_cache()
        return jsonify({"status": "success", "message": f"Excel uğurla yükləndi: {len(filtered_products)} məhsul daxil edildi."})
    except Exception as e:
        return jsonify({"status": "error", "message": f"Fayl oxunarkən xəta baş verdi: {str(e)}"}), 500


def is_summary_row(barkod_val, kod_val, brend_val, adi_val):
    combined = (str(barkod_val) + " " + str(kod_val) + " " + str(brend_val) + " " + str(adi_val)).lower()
    keywords = ['yekun', 'cemi', 'cem', 'total', 'итого', 'всего', 'summary', 'subtotal', 'cəmi', 'cəm']
    for kw in keywords:
        if kw in combined:
            return True
    return False


def merge_excel_files_data(files_list):
    """
    files_list: list of tuples (filename, file_bytes)
    Returns openpyxl.Workbook matching Picture 3 structure
    """
    merged_products = {}
    
    barkod_synonyms = ['barkod', 'barkkod', 'barcode', 'strixkod', 'strix-kod', 'strix_kod', 'ean', 'shtrihkod', 'штрихкод', 'bar_code', 'sh_kod', 'bar_kod', 'barkod_no', 'barkkod_no']
    kod_synonyms = ['kod', 'kodu', 'product_code', 'mal_kodu', 'mehsul_kodu', 'kod_mehsul', 'kod_mal', 'код', 'артикул', 'artikul', 'item_code', 'item_id']
    brend_synonyms = ['brend', 'brand', 'marka', 'istehsalci', 'производитель', 'фирма', 'firma', 'vendor', 'manufacturer', 'malin_markasi', 'taminatci', 'supplier']
    adi_synonyms = ['mehsulun adi', 'mehsulun_adi', 'malin adi', 'malin_adi', 'adi', 'name', 'mehsul', 'description', 'nomenklatura', 'наименование', 'название', 'товар', 'urun_adi', 'tam_adi', 'mehsul_adi', 'item_name', 'aciglama']
    qaliq_synonyms = ['sistem qaligi', 'sistem_qaligi', 'anbar qaligi', 'anbar_qaligi', 'qaliq', 'stock', 'qaliq_miqdari', 'остаток', 'stok_miktari', 'stok', 'balance', 'qty', 'quantity', 'son_qaliq', 'tek_qaliq', 'mevcut']
    sayim_synonyms = ['yeni sayim', 'yeni_sayim', 'sayim', 'sayilan', 'sayim miqdari', 'sayim_miqdari', 'sayilan_miqdar', 'sayim_miqdar', 'faktiki sayim', 'faktiki_sayim', 'miqdar', 'sayi', 'say', 'количество', 'fakt', 'yeni_sayim_miqdari', 'real say']
    qiymet_synonyms = ['qiymet', 'qiymeti', 'mehsulun qiymeti', 'mehsulun_qiymeti', 'price', 'satis_qiymeti', 'satis qiymeti', 'satis_qiymet', 'satis qiymet', 'satiş qiyməti', 'maya_qiymeti', 'maya qiymeti', 'цена', 'стоимость', 'fiyat', 'cost', 'unit_price', 'qiymat', 'mebleg', 'perakende']
    operator_synonyms = ['operator', 'sayimci', 'user', 'istifadəçi', 'operatorlar', 'sayımçı']

    for filename, file_bytes in files_list:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb.active
        except Exception:
            continue
            
        headers = {}
        for col in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col).value
            if val is not None:
                headers[normalize_header(val)] = col

        def get_col_idx(synonyms):
            for syn in synonyms:
                norm_syn = normalize_header(syn)
                for h_key, col_idx in headers.items():
                    if norm_syn == h_key or norm_syn in h_key or h_key in norm_syn:
                        return col_idx
            return None

        barkod_col = get_col_idx(barkod_synonyms)
        kod_col = get_col_idx(kod_synonyms)
        brend_col = get_col_idx(brend_synonyms)
        adi_col = get_col_idx(adi_synonyms)
        qaliq_col = get_col_idx(qaliq_synonyms)
        sayim_col = get_col_idx(sayim_synonyms)
        qiymet_col = get_col_idx(qiymet_synonyms)
        operator_col = get_col_idx(operator_synonyms)

        if not sayim_col:
            sayim_col = qaliq_col

        for row in range(2, sheet.max_row + 1):
            barkod_val = str(sheet.cell(row=row, column=barkod_col).value or "").strip() if barkod_col else ""
            if barkod_val.endswith(".0"): barkod_val = barkod_val[:-2]

            kod_val = str(sheet.cell(row=row, column=kod_col).value or "").strip() if kod_col else ""
            if kod_val.endswith(".0"): kod_val = kod_val[:-2]

            brend_val = str(sheet.cell(row=row, column=brend_col).value or "").strip() if brend_col else ""
            adi_val = str(sheet.cell(row=row, column=adi_col).value or "").strip() if adi_col else ""
            operator_val = str(sheet.cell(row=row, column=operator_col).value or "").strip() if operator_col else ""

            if is_summary_row(barkod_val, kod_val, brend_val, adi_val):
                continue

            raw_count = sheet.cell(row=row, column=sayim_col).value if sayim_col else 0
            try:
                count_num = float(raw_count) if raw_count is not None else 0.0
            except (ValueError, TypeError):
                count_num = 0.0

            raw_qaliq = sheet.cell(row=row, column=qaliq_col).value if qaliq_col else 0
            try:
                qaliq_num = float(raw_qaliq) if raw_qaliq is not None else 0.0
            except (ValueError, TypeError):
                qaliq_num = 0.0

            raw_qiymet = sheet.cell(row=row, column=qiymet_col).value if qiymet_col else 0
            try:
                qiymet_num = float(raw_qiymet) if raw_qiymet is not None else 0.0
            except (ValueError, TypeError):
                qiymet_num = 0.0

            product_key = barkod_val or kod_val or normalize_header(adi_val)
            if not product_key:
                continue

            clean_filename = os.path.basename(filename)

            if product_key not in merged_products:
                merged_products[product_key] = {
                    'barkod': barkod_val,
                    'kod': kod_val,
                    'brend': brend_val,
                    'adi': adi_val,
                    'anbar_qaligi': qaliq_num,
                    'yeni_sayim': count_num,
                    'qiymet': qiymet_num,
                    'operators': [operator_val] if operator_val else [],
                    'sources': [clean_filename]
                }
            else:
                existing = merged_products[product_key]
                existing['yeni_sayim'] += count_num
                if not existing['anbar_qaligi'] and qaliq_num:
                    existing['anbar_qaligi'] = qaliq_num
                if not existing['qiymet'] and qiymet_num:
                    existing['qiymet'] = qiymet_num
                if not existing['brend'] and brend_val:
                    existing['brend'] = brend_val
                if not existing['adi'] and adi_val:
                    existing['adi'] = adi_val
                if operator_val and operator_val not in existing['operators']:
                    existing['operators'].append(operator_val)
                if clean_filename not in existing['sources']:
                    existing['sources'].append(clean_filename)

    out_wb = openpyxl.Workbook()
    out_sheet = out_wb.active
    out_sheet.title = "Yekun Sayım"

    out_sheet.views.sheetView[0].showGridLines = True

    # Header matching Picture 3
    headers_list = [
        "No", "Kod", "Barkod", "Brend", "Məhsulun Adı",
        "Qiymət", "Sistem Qalıq", "Yeni Sayım", "Say Fərqi",
        "Qiymət Fərqi", "Operator", "Keçdiyi Fayllar"
    ]
    out_sheet.append(headers_list)
    out_sheet.row_dimensions[1].height = 28

    sorted_products = sorted(
        merged_products.values(),
        key=lambda x: (x['brend'].lower(), x['adi'].lower(), x['barkod'])
    )

    total_qaliq = 0
    total_sayim = 0
    total_ferq = 0
    total_mebleg_ferqi = 0.0

    for idx, prod in enumerate(sorted_products, 1):
        r = idx + 1
        qaliq = prod['anbar_qaligi']
        sayim = prod['yeni_sayim']
        qiymet = prod['qiymet']
        operators_str = ", ".join(prod['operators']) if prod['operators'] else ""
        sources_str = ", ".join(prod['sources'])

        q_disp = int(qaliq) if qaliq.is_integer() else round(qaliq, 2)
        s_disp = int(sayim) if sayim.is_integer() else round(sayim, 2)

        # Dynamic Excel Formulas as requested by user:
        # Say Fərqi = Yeni Sayım - Sistem Qalıq (=H{r}-G{r})
        # Qiymət Fərqi = Say Fərqi * Qiymət (=I{r}*F{r})
        say_ferqi_formula = f"=H{r}-G{r}"
        qiymet_ferqi_formula = f"=I{r}*F{r}"

        out_sheet.append([
            idx,
            prod['kod'],
            prod['barkod'],
            prod['brend'],
            prod['adi'],
            round(qiymet, 2),
            q_disp,
            s_disp,
            say_ferqi_formula,
            qiymet_ferqi_formula,
            operators_str,
            sources_str
        ])

    last_row = len(sorted_products) + 2
    data_last_row = last_row - 1

    total_row = [
        "YEKUN CƏM", "", "", "", "",
        "",
        f"=SUM(G2:G{data_last_row})",
        f"=SUM(H2:H{data_last_row})",
        f"=SUM(I2:I{data_last_row})",
        f"=SUM(J2:J{data_last_row})",
        "",
        f"Cəmi {len(sorted_products)} çeşit məhsul"
    ]
    out_sheet.append(total_row)

    # Enable AutoFilter for header row
    last_col_letter = get_column_letter(len(headers_list))
    out_sheet.auto_filter.ref = f"A1:{last_col_letter}{out_sheet.max_row}"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_wrap_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col in range(1, len(headers_list) + 1):
        cell = out_sheet.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    
    green_font = Font(name="Segoe UI", size=10, bold=True, color="15803D")
    red_font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
    data_font = Font(name="Segoe UI", size=10, bold=False, color="1E293B")
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for r in range(2, last_row):
        out_sheet.row_dimensions[r].height = 22
        is_even = (r % 2 == 0)
        row_bg = zebra_fill if is_even else white_fill

        for c in range(1, len(headers_list) + 1):
            cell = out_sheet.cell(row=r, column=c)
            cell.font = data_font
            cell.border = thin_border
            cell.fill = row_bg
            
            if c == 1:
                cell.alignment = center_align
                cell.number_format = '0'
            elif c in [2, 3]:
                cell.alignment = center_align
                cell.number_format = '@'
            elif c in [4, 5, 11]:
                cell.alignment = left_wrap_align
            elif c in [6, 10]:
                cell.alignment = right_align
                cell.number_format = '#,##0.00'
            elif c in [7, 8, 9]:
                cell.alignment = right_align
                cell.number_format = '#,##0' if isinstance(cell.value, int) else '#,##0.00'
            elif c == 12:
                cell.alignment = left_wrap_align

        idx_prod = r - 2
        prod_data = sorted_products[idx_prod]
        ferq_calc = prod_data['yeni_sayim'] - prod_data['anbar_qaligi']
        cell_f = out_sheet.cell(row=r, column=9)
        cell_m = out_sheet.cell(row=r, column=10)
        
        if ferq_calc > 0:
            cell_f.fill = green_fill
            cell_f.font = green_font
            cell_m.fill = green_fill
            cell_m.font = green_font
        elif ferq_calc < 0:
            cell_f.fill = red_fill
            cell_f.font = red_font
            cell_m.fill = red_fill
            cell_m.font = red_font

    out_sheet.row_dimensions[last_row].height = 26
    summary_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    summary_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    
    for c in range(1, len(headers_list) + 1):
        cell = out_sheet.cell(row=last_row, column=c)
        cell.fill = summary_fill
        cell.font = summary_font
        cell.border = thin_border
        if c in [7, 8, 9]:
            cell.alignment = right_align
            cell.number_format = '#,##0'
        elif c in [6, 10]:
            cell.alignment = right_align
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = center_align if c == 1 else left_align

    min_widths = {
        1: 10,   # No
        2: 16,   # Kod
        3: 20,   # Barkod
        4: 22,   # Brend
        5: 45,   # Məhsulun Adı
        6: 14,   # Qiymət
        7: 18,   # Sistem Qalıq
        8: 18,   # Yeni Sayım
        9: 18,   # Say Fərqi
        10: 20,  # Qiymət Fərqi
        11: 20,  # Operator
        12: 40   # Keçdiyi Fayllar
    }

    for col_idx in range(1, len(headers_list) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in out_sheet[col_letter]:
            val_str = str(cell.value or '')
            if cell.row == 1:
                max_len = max(max_len, len(val_str) + 6)
            else:
                max_len = max(max_len, len(val_str) + 3)
        
        calculated_width = max(max_len, min_widths.get(col_idx, 15))
        out_sheet.column_dimensions[col_letter].width = calculated_width

    return out_wb


@app.route("/merge_excel_files", methods=["POST"])
@admin_required
def merge_excel_files():
    uploaded_files = request.files.getlist('files')
    if not uploaded_files:
        if 'file' in request.files:
            uploaded_files = [request.files['file']]
            
    if not uploaded_files or len(uploaded_files) == 0:
        return jsonify({"status": "error", "message": "Zəhmət olmasa ən azı 1 Excel faylı yükləyin!"}), 400

    files_list = []
    for f in uploaded_files:
        if f and f.filename and f.filename.endswith('.xlsx'):
            files_list.append((f.filename, f.read()))

    if len(files_list) == 0:
        return jsonify({"status": "error", "message": "Heç bir keçərli .xlsx faylı tapılmadı!"}), 400

    try:
        out_wb = merge_excel_files_data(files_list)
        out_stream = io.BytesIO()
        out_wb.save(out_stream)
        out_stream.seek(0)
        
        return send_file(
            out_stream,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="Yekun_Birlesdirilmis_Sayim.xlsx"
        )
    except Exception as e:
        return jsonify({"status": "error", "message": f"Fayllar birləşdirilərkən xəta baş verdi: {str(e)}"}), 500


@app.route("/import_counted_excel", methods=["POST"])
@admin_required
def import_counted_excel():
    uploaded_files = request.files.getlist('files')
    if not uploaded_files:
        return jsonify({"status": "error", "message": "Fayl seçilməyib."}), 400

    merged_data = {}

    barkod_synonyms = ['barkod', 'barkkod', 'barcode', 'strixkod', 'strix-kod', 'strix_kod', 'ean', 'shtrihkod', 'штрихкод']
    kod_synonyms = ['kod', 'kodu', 'product_code', 'mal_kodu', 'mehsul_kodu', 'kod_mehsul', 'kod_mal']
    brend_synonyms = ['brend', 'brand', 'marka', 'istehsalci', 'firma']
    adi_synonyms = ['mehsulun adi', 'mehsulun_adi', 'malin adi', 'malin_adi', 'adi', 'name', 'mehsul']
    qaliq_synonyms = ['sistem qaligi', 'sistem_qaligi', 'anbar qaligi', 'qaliq', 'stock']
    sayim_synonyms = ['yeni sayim', 'yeni_sayim', 'sayim', 'sayilan', 'sayim miqdari', 'sayilan_miqdar', 'miqdar', 'say']
    qiymet_synonyms = ['qiymet', 'qiymeti', 'mehsulun qiymeti', 'price', 'satis_qiymeti']
    operator_synonyms = ['operator', 'sayimci', 'user', 'istifadəçi']

    try:
        for file in uploaded_files:
            if not file.filename.lower().endswith('.xlsx'):
                continue

            file_bytes = file.read()
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                sheet = wb.active
            except Exception:
                continue

            headers = {}
            for col in range(1, sheet.max_column + 1):
                val = sheet.cell(row=1, column=col).value
                if val is not None:
                    headers[normalize_header(val)] = col

            def get_col_idx(synonyms):
                for syn in synonyms:
                    norm_syn = normalize_header(syn)
                    for h_key, col_idx in headers.items():
                        if norm_syn == h_key or norm_syn in h_key or h_key in norm_syn:
                            return col_idx
                return None

            barkod_col = get_col_idx(barkod_synonyms)
            kod_col = get_col_idx(kod_synonyms)
            brend_col = get_col_idx(brend_synonyms)
            adi_col = get_col_idx(adi_synonyms)
            qaliq_col = get_col_idx(qaliq_synonyms)
            sayim_col = get_col_idx(sayim_synonyms)
            qiymet_col = get_col_idx(qiymet_synonyms)
            operator_col = get_col_idx(operator_synonyms)

            if not sayim_col:
                sayim_col = qaliq_col

            for row in range(2, sheet.max_row + 1):
                barkod_val = str(sheet.cell(row=row, column=barkod_col).value or "").strip() if barkod_col else ""
                if barkod_val.endswith(".0"): barkod_val = barkod_val[:-2]

                kod_val = str(sheet.cell(row=row, column=kod_col).value or "").strip() if kod_col else ""
                if kod_val.endswith(".0"): kod_val = kod_val[:-2]

                brend_val = str(sheet.cell(row=row, column=brend_col).value or "").strip() if brend_col else ""
                adi_val = str(sheet.cell(row=row, column=adi_col).value or "").strip() if adi_col else ""
                operator_val = str(sheet.cell(row=row, column=operator_col).value or "").strip() if operator_col else "Excel Geri Yükləmə"

                if is_summary_row(barkod_val, kod_val, brend_val, adi_val):
                    continue

                raw_count = sheet.cell(row=row, column=sayim_col).value if sayim_col else 0
                try:
                    count_num = float(raw_count) if raw_count is not None else 0.0
                except (ValueError, TypeError):
                    count_num = 0.0

                raw_qaliq = sheet.cell(row=row, column=qaliq_col).value if qaliq_col else None
                qaliq_num = None
                if raw_qaliq is not None:
                    try:
                        qaliq_num = float(raw_qaliq)
                    except (ValueError, TypeError):
                        qaliq_num = None

                raw_qiymet = sheet.cell(row=row, column=qiymet_col).value if qiymet_col else None
                qiymet_num = None
                if raw_qiymet is not None:
                    try:
                        qiymet_num = float(raw_qiymet)
                    except (ValueError, TypeError):
                        qiymet_num = None

                key = barkod_val or kod_val or normalize_header(adi_val)
                if not key:
                    continue

                if key not in merged_data:
                    merged_data[key] = {
                        'barkod': barkod_val,
                        'kod': kod_val,
                        'brend': brend_val,
                        'adi': adi_val,
                        'qaliq': qaliq_num,
                        'qiymet': qiymet_num,
                        'yeni': count_num,
                        'operators': [operator_val] if operator_val else []
                    }
                else:
                    item = merged_data[key]
                    item['yeni'] += count_num
                    if item['qaliq'] is None and qaliq_num is not None:
                        item['qaliq'] = qaliq_num
                    if item['qiymet'] is None and qiymet_num is not None:
                        item['qiymet'] = qiymet_num
                    if not item['brend'] and brend_val:
                        item['brend'] = brend_val
                    if not item['adi'] and adi_val:
                        item['adi'] = adi_val
                    if operator_val and operator_val not in item['operators']:
                        item['operators'].append(operator_val)

        conn = get_db_connection()
        cursor = conn.cursor()
        imported_count = 0
        updated_items = 0

        for key, item in merged_data.items():
            b_val = item['barkod'] or item['kod']
            k_val = item['kod'] or item['barkod']
            op_str = ", ".join(item['operators']) if item['operators'] else "Excel Geri Yükləmə"
            qaliq_val = item['qaliq'] if item['qaliq'] is not None else 0.0
            qiymet_val = item['qiymet'] if item['qiymet'] is not None else 0.0
            yeni_val = item['yeni']

            if get_db_type() == "postgres":
                cursor.execute("SELECT barcode FROM products WHERE barcode = %s OR kod = %s LIMIT 1", (b_val, k_val))
            else:
                cursor.execute("SELECT barcode FROM products WHERE barcode = ? OR kod = ? LIMIT 1", (b_val, k_val))

            row_db = cursor.fetchone()
            if row_db:
                db_barcode = row_db[0]
                if item['qaliq'] is not None and item['qiymet'] is not None:
                    if get_db_type() == "postgres":
                        cursor.execute("UPDATE products SET yeni = %s, qaliq = %s, qiymet = %s, operator = %s WHERE barcode = %s", (yeni_val, qaliq_val, qiymet_val, op_str, db_barcode))
                    else:
                        cursor.execute("UPDATE products SET yeni = ?, qaliq = ?, qiymet = ?, operator = ? WHERE barcode = ?", (yeni_val, qaliq_val, qiymet_val, op_str, db_barcode))
                elif item['qaliq'] is not None:
                    if get_db_type() == "postgres":
                        cursor.execute("UPDATE products SET yeni = %s, qaliq = %s, operator = %s WHERE barcode = %s", (yeni_val, qaliq_val, op_str, db_barcode))
                    else:
                        cursor.execute("UPDATE products SET yeni = ?, qaliq = ?, operator = ? WHERE barcode = ?", (yeni_val, qaliq_val, op_str, db_barcode))
                else:
                    if get_db_type() == "postgres":
                        cursor.execute("UPDATE products SET yeni = %s, operator = %s WHERE barcode = %s", (yeni_val, op_str, db_barcode))
                    else:
                        cursor.execute("UPDATE products SET yeni = ?, operator = ? WHERE barcode = ?", (yeni_val, op_str, db_barcode))
                updated_items += 1
            else:
                final_barcode = b_val or f"BAR_{int(time.time()*1000)}"
                final_code = k_val
                if get_db_type() == "postgres":
                    cursor.execute(
                        "INSERT INTO products (barcode, kod, brend, adi, qaliq, qiymet, yeni, operator) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
                        (final_barcode, final_code, item['brend'], item['adi'], qaliq_val, qiymet_val, yeni_val, op_str)
                    )
                else:
                    cursor.execute(
                        "INSERT INTO products (barcode, kod, brend, adi, qaliq, qiymet, yeni, operator) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                        (final_barcode, final_code, item['brend'], item['adi'], qaliq_val, qiymet_val, yeni_val, op_str)
                    )
                imported_count += 1

        conn.commit()
    except Exception as e:
        return jsonify({"status": "error", "message": f"Yükləmə xətası: {str(e)}"}), 500
    finally:
        cursor.close()
        if get_db_type() == "postgres":
            _return_connection(conn)

    invalidate_products_cache()
    return jsonify({
        "status": "success",
        "message": f"Sayılmış məhsullar sistemə yükləndi: {updated_items} məhsulun sayımı və qalığı yeniləndi, {imported_count} yeni məhsul daxil edildi."
    })


def apply_excel_styling(sheet, text_columns=None, label_col=None, summary_cols=None):
    """Apply common Excel styling, AutoFilter and total row to a sheet."""
    text_columns = text_columns or []
    label_col = label_col or 2
    summary_cols = summary_cols or []
    
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    # AutoFilter
    if sheet.max_row > 0 and sheet.max_column > 0:
        last_col_letter = get_column_letter(sheet.max_column)
        sheet.auto_filter.ref = f"A1:{last_col_letter}{sheet.max_row}"
    
    # Design tokens
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Header styling
    sheet.row_dimensions[1].height = 28
    for col_idx in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = cell_border
    
    # Data rows styling
    for row_idx in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = 20
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.border = cell_border
            cell.font = Font(name="Segoe UI", size=10)
            if col_idx in text_columns:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Total row
    total_row = sheet.max_row + 1
    sheet.cell(row=total_row, column=label_col, value="YEKUN CƏMİ:")
    
    for col_idx in summary_cols:
        col_letter = get_column_letter(col_idx)
        sheet.cell(row=total_row, column=col_idx, value=f"=SUM({col_letter}2:{col_letter}{total_row-1})")
    
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    sheet.row_dimensions[total_row].height = 24
    for col_idx in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=total_row, column=col_idx)
        cell.fill = total_fill
        cell.border = cell_border
        if col_idx == label_col:
            cell.font = Font(name="Segoe UI", size=10, bold=True)
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx in summary_cols:
            cell.font = Font(name="Segoe UI", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Auto-adjust column widths
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if val_str.startswith('='):
                val_str = "100.00 AZN"
            if len(val_str) > max_len:
                max_len = len(val_str)
        sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
    
    sheet.views.sheetView[0].showGridLines = True


@app.route("/download_excel", methods=["GET"])
@admin_required
def download_excel():
    try:
        # Load template
        res = execute_query("SELECT filename, file_bytes FROM excel_template WHERE id = 1", fetch=True)
        if not res:
            return "Sistemdə hələ heç bir Excel yüklənməyib. Zəhmət olmasa əvvəlcə Excel yükləyin.", 404
            
        filename = res[0]['filename']
        file_bytes = res[0]['file_bytes']
        
        # Load workbook in memory
        wb = openpyxl.load_workbook(io.BytesIO(bytes(file_bytes)), data_only=False)
        sheet = wb.active
        
        # Map headers
        headers = {}
        for col in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col).value
            if val:
                headers[normalize_header(val)] = col
                
        # Calculation columns check
        yeni_sayim_col = headers.get('yeni sayim') or headers.get('real say')
        say_ferqi_col = headers.get('say ferqi')
        qiymet_ferqi_col = headers.get('qiymet ferqi')
        operator_col = headers.get('operator') or headers.get('sayimci') or headers.get('user')
        
        max_col = sheet.max_column
        if not yeni_sayim_col:
            max_col += 1
            sheet.cell(row=1, column=max_col, value="Yeni Sayim")
            yeni_sayim_col = max_col
        if not say_ferqi_col:
            max_col += 1
            sheet.cell(row=1, column=max_col, value="Say ferqi")
            say_ferqi_col = max_col
        if not qiymet_ferqi_col:
            max_col += 1
            sheet.cell(row=1, column=max_col, value="Qiymet ferqi")
            qiymet_ferqi_col = max_col
        if not operator_col:
            max_col += 1
            sheet.cell(row=1, column=max_col, value="Operator")
            operator_col = max_col
            
        qiymet_col = find_col_index(headers, qiymet_synonyms)
        if not qiymet_col:
            max_col += 1
            sheet.cell(row=1, column=max_col, value="Qiyməti")
            qiymet_col = max_col

        yeni_letter = get_column_letter(yeni_sayim_col)
        qaliq_letter = get_column_letter(find_col_index(headers, qaliq_synonyms) or 1)
        say_ferqi_letter = get_column_letter(say_ferqi_col)
        qiymet_letter = get_column_letter(qiymet_col)
        
        # Get products
        all_db_products = execute_query("SELECT * FROM products ORDER BY row_idx ASC, barcode ASC", fetch=True)
        
        # Write back existing rows
        for p in all_db_products:
            row = p['row_idx']
            if row:
                yeni_val = float(p['yeni']) if p['yeni'] is not None else None
                sheet.cell(row=row, column=yeni_sayim_col, value=yeni_val)
                sheet.cell(row=row, column=say_ferqi_col, value=f"={yeni_letter}{row}-{qaliq_letter}{row}")
                sheet.cell(row=row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{row}*{qiymet_letter}{row}")
                sheet.cell(row=row, column=operator_col, value=clean_operator_string(p['operator']) if yeni_val is not None else None)
                
        # Append manual products
        barkod_col = find_col_index(headers, barkod_synonyms)
        kod_col = find_col_index(headers, kod_synonyms, exclude_indices={barkod_col} if barkod_col else set())
        if not barkod_col and kod_col:
            barkod_col = kod_col
            kod_col = None

        brend_col = find_col_index(headers, brend_synonyms)
        adi_col = find_col_index(headers, adi_synonyms, exclude_indices={brend_col} if brend_col else set())
        anbar_qaligi_col = find_col_index(headers, qaliq_synonyms) or 1
        
        for p in all_db_products:
            if not p['row_idx']:  # Added manually
                new_row = sheet.max_row + 1
                sheet.cell(row=new_row, column=barkod_col, value=p['barcode'])
                if kod_col and p.get('kod'):
                    sheet.cell(row=new_row, column=kod_col, value=p['kod'])
                sheet.cell(row=new_row, column=brend_col, value=p['brend'])
                if adi_col:
                    sheet.cell(row=new_row, column=adi_col, value=p['adi'])
                sheet.cell(row=new_row, column=anbar_qaligi_col, value=float(p['qaliq']))
                sheet.cell(row=new_row, column=qiymet_col, value=float(p['qiymet']))
                yeni_val = float(p['yeni']) if p['yeni'] is not None else None
                sheet.cell(row=new_row, column=yeni_sayim_col, value=yeni_val)
                sheet.cell(row=new_row, column=say_ferqi_col, value=f"={yeni_letter}{new_row}-{qaliq_letter}{new_row}")
                sheet.cell(row=new_row, column=qiymet_ferqi_col, value=f"={say_ferqi_letter}{new_row}*{qiymet_letter}{new_row}")
                sheet.cell(row=new_row, column=operator_col, value=clean_operator_string(p['operator']))
        
        # Styling
        text_columns = [brend_col, adi_col] if adi_col else [brend_col]
        apply_excel_styling(
            sheet,
            text_columns=text_columns,
            label_col=brend_col or 2,
            summary_cols=[yeni_sayim_col, say_ferqi_col, qiymet_ferqi_col]
        )
        
        # Save in memory buffer
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        wb.close()
        
        out_filename = f"yekun_sayim_neticesi_{filename}" if not filename.startswith("yekun_sayim_neticesi") else filename
        
        return send_file(
            out,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=out_filename
        )
    except Exception as e:
        return f"Hesabat yaradılarkən xəta: {str(e)}", 500

@app.route("/download_discrepancies_excel", methods=["GET"])
@admin_required
def download_discrepancies_excel():
    try:
        # Get products
        all_db_products = execute_query("SELECT * FROM products ORDER BY row_idx ASC, barcode ASC", fetch=True)
        
        # Filter products with differences (counted and yeni != qaliq)
        diff_products = []
        for p in all_db_products:
            yeni_val = float(p['yeni']) if p['yeni'] is not None else None
            qaliq_val = float(p['qaliq'])
            if yeni_val is not None and yeni_val != qaliq_val:
                diff_products.append(p)
                
        # Create a new workbook
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Ferq Hesabati"
        
        # Headers
        headers = ["No", "Kod", "Barkod", "Brend", "Məhsulun Adı", "Qiymət", "Sistem Qalığı", "Yeni Sayım", "Say Fərqi", "Qiymət Fərqi", "Operator"]
        sheet.append(headers)
        
        # Fill data
        row_num = 1
        for p in diff_products:
            row_num += 1
            yeni_val = float(p['yeni'])
            qaliq_val = float(p['qaliq'])
            diff_qty = yeni_val - qaliq_val
            diff_price = diff_qty * float(p['qiymet'])
            
            row_data = [
                row_num - 1,
                p.get('kod') or '',
                p['barcode'],
                p['brend'],
                p['adi'],
                float(p['qiymet']),
                qaliq_val,
                yeni_val,
                diff_qty,
                diff_price,
                clean_operator_string(p['operator'])
            ]
            sheet.append(row_data)
            
        # Add formulas and styling
        for r_idx in range(2, sheet.max_row + 1):
            sheet.cell(row=r_idx, column=9, value=f"=H{r_idx}-G{r_idx}")
            sheet.cell(row=r_idx, column=10, value=f"=I{r_idx}*F{r_idx}")
        
        apply_excel_styling(
            sheet,
            text_columns=[4, 5],
            label_col=5,
            summary_cols=[7, 8, 9, 10]
        )
        
        # Save in memory buffer
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        wb.close()
        
        return send_file(
            out,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="ferqler_hesabati.xlsx"
        )
    except Exception as e:
        return f"Hesabat yaradılarkən xəta: {str(e)}", 500

@app.route("/download_counted_excel", methods=["GET"])
@admin_required
def download_counted_excel():
    try:
        # Get products
        all_db_products = execute_query("SELECT * FROM products ORDER BY row_idx ASC, barcode ASC", fetch=True)
        
        # Filter products that are counted (yeni is not None)
        counted_products = []
        for p in all_db_products:
            yeni_val = float(p['yeni']) if p['yeni'] is not None else None
            if yeni_val is not None:
                counted_products.append(p)
                
        # Create a new workbook
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Sayilan Mehsullar"
        
        # Headers
        headers = ["No", "Kod", "Barkod", "Brend", "Məhsulun Adı", "Qiymət", "Sistem Qalığı", "Yeni Sayım", "Say Fərqi", "Qiymət Fərqi", "Operator"]
        sheet.append(headers)
        
        # Fill data
        row_num = 1
        for p in counted_products:
            row_num += 1
            yeni_val = float(p['yeni'])
            qaliq_val = float(p['qaliq'])
            diff_qty = yeni_val - qaliq_val
            diff_price = diff_qty * float(p['qiymet'])
            
            row_data = [
                row_num - 1,
                p.get('kod') or '',
                p['barcode'],
                p['brend'],
                p['adi'],
                float(p['qiymet']),
                qaliq_val,
                yeni_val,
                diff_qty,
                diff_price,
                clean_operator_string(p['operator'])
            ]
            sheet.append(row_data)
            
        # Add formulas and styling
        for r_idx in range(2, sheet.max_row + 1):
            sheet.cell(row=r_idx, column=9, value=f"=H{r_idx}-G{r_idx}")
            sheet.cell(row=r_idx, column=10, value=f"=I{r_idx}*F{r_idx}")
        
        apply_excel_styling(
            sheet,
            text_columns=[4, 5],
            label_col=5,
            summary_cols=[7, 8, 9, 10]
        )
        
        # Save in memory buffer
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        wb.close()
        
        return send_file(
            out,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="sayilanlar_hesabati.xlsx"
        )
    except Exception as e:
        return f"Hesabat yaradılarkən xəta: {str(e)}", 500

@app.route("/reset_inventory", methods=["POST"])
@admin_required
def reset_inventory():
    try:
        execute_query("UPDATE products SET yeni = NULL, operator = '', order_num = 0")
        execute_query("DELETE FROM scan_logs")
        execute_query("DELETE FROM operator_counts")
        invalidate_products_cache()
        return jsonify({"status": "success", "message": "Bütün sayımlar sıfırlandı!"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

if __name__ == "__main__":
    # Hugging Face Spaces dynamically binds the port, but must listen on 7860
    port = int(os.environ.get("PORT", 7860))
    app.run(host="0.0.0.0", port=port, debug=False)
